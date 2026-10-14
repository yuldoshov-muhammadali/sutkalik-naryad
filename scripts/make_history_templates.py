"""data/tarix_shablon.xlsx (bo'sh) va data/tarix_namuna.xlsx (18-sentabr) fayllarini yaratadi.

Ishga tushirish: python scripts/make_history_templates.py
"""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from scripts.demo import SEPT_18  # noqa: E402

from naryad.engine.config import POSTS_BY_CODE  # noqa: E402
from naryad.engine.models import ROLE_TITLES  # noqa: E402

HEADER = ["sana", "joy", "lavozim", "kursant", "guruh"]
WIDTHS = [14, 30, 24, 28, 12]
POST_NAMES = [p.name for p in POSTS_BY_CODE.values()]
ROLE_NAMES = [t.lower() for t in ROLE_TITLES.values()]

INSTRUCTIONS = [
    "Qog‘ozdagi o‘tgan sutkalarni kiritish shabloni",
    "",
    "Har qator — bitta o‘rin. Har sutka uchun aynan 20 qator (17 o‘g‘il + 3 qiz).",
    "sana — sutka boshlangan kun (18:00): 2026-09-18 yoki 18.09.2026.",
    "joy — " + "; ".join(POST_NAMES) + ".",
    "lavozim — navbatchi / navbatchi yordamchisi / posbon (patrulda hammasi posbon).",
    "kursant — qisqa (Sa’dullayev B) yoki to‘liq F.I.Sh. Rasmiy ro‘yxatdagidek yozing.",
    "guruh — YuP-125, BI-225 … BI-625. Bir xil qisqa ismlarni ajratish uchun kerak (masalan, Aslonov A).",
    "",
    "Tekshirish: python manage.py import_history fayl.xlsx --dry-run",
    "Yozish:     python manage.py import_history fayl.xlsx",
]


def build(path: Path, rows: list[list[object]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "tarix"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, width in enumerate(WIDTHS):
        ws.column_dimensions[chr(ord("A") + i)].width = width
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    post_dv = DataValidation(type="list", formula1='"' + ",".join(POST_NAMES) + '"', allow_blank=True)
    role_dv = DataValidation(type="list", formula1='"' + ",".join(ROLE_NAMES) + '"', allow_blank=True)
    ws.add_data_validation(post_dv)
    ws.add_data_validation(role_dv)
    post_dv.add("B2:B2000")
    role_dv.add("C2:C2000")

    info = wb.create_sheet("yo‘riqnoma")
    info.column_dimensions["A"].width = 110
    for line in INSTRUCTIONS:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    wb.save(path)


def sample_rows() -> list[list[object]]:
    day = dt.date(2026, 9, 18)
    rows = []
    for post_code, position, group, last, initial in SEPT_18:
        post = POSTS_BY_CODE[post_code]
        role = ROLE_TITLES[post.slots[position - 1].role].lower()
        rows.append([day, post.name, role, f"{last} {initial}", group])
    return rows


def main() -> None:
    build(ROOT / "data" / "tarix_shablon.xlsx", [])
    build(ROOT / "data" / "tarix_namuna.xlsx", sample_rows())
    print("Yaratildi: data/tarix_shablon.xlsx, data/tarix_namuna.xlsx")


if __name__ == "__main__":
    main()
