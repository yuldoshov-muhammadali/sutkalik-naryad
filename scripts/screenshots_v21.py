"""20-bosqich (v1.11) skrinshotlari: komandirning zalyot kiritish formasi va tasdiqlash oynasi (telefon), komandirning
zalyotlar ro'yxati, admin banneri “yangi zalyot”, “Zalyot hisobi”, Excel varag'i, “Sinov ma'lumotlarini tozalash”.
Natija: docs/skrinshotlar/v1.11/.

    BASE=http://127.0.0.1:8139 python scripts/screenshots_v21.py docs/skrinshotlar/v1.11
"""
import html
import io
import os
import sys
import urllib.request

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

OUT = sys.argv[1]
BASE = os.environ.get("BASE", "http://127.0.0.1:8139")
PASSWORD = os.environ.get("PASSWORD", "Demo-parol-2026")
PHONE = {"viewport": {"width": 360, "height": 780}, "device_scale_factor": 2, "is_mobile": True, "has_touch": True}
PC = {"viewport": {"width": 1280, "height": 860}, "device_scale_factor": 1}


def login(browser, user, device):
    page = browser.new_context(locale="uz-UZ", **device).new_page()
    page.goto(BASE + "/kirish/")
    page.fill("#id_username", user)
    page.fill("#id_password", PASSWORD)
    page.click("form:has(#id_password) button[type=submit]")
    page.wait_for_url(BASE + "/")
    return page


with sync_playwright() as p:
    os.makedirs(OUT, exist_ok=True)
    b = p.chromium.launch(executable_path=os.environ.get("CHROMIUM_PATH") or None)
    page = login(b, "komandir.bi525", PHONE)
    page.goto(BASE + "/mening-guruhim/zalyot/")
    page.select_option("#id_cadet", label=page.locator("#id_cadet option").nth(3).inner_text())
    page.select_option("#id_reason_type", "tartib")
    page.fill("#id_note", "Dars paytida telefondan foydalandi")
    page.fill("#id_count", "2")
    page.screenshot(path=f"{OUT}/komandir-zalyot-formasi-telefon.png", full_page=True)
    page.click("button[type=submit]:has-text('Davom etish')")
    page.wait_for_selector("[data-testid=penalty-confirm]")
    page.screenshot(path=f"{OUT}/komandir-zalyot-tasdiqlash-telefon.png", full_page=True)
    page.goto(BASE + "/mening-guruhim/zalyotlar/")
    page.screenshot(path=f"{OUT}/komandir-zalyotlar-royxati-telefon.png", full_page=True)
    page.context.close()

    page = login(b, "kursboshligi", PC)
    page.wait_for_timeout(500)
    page.locator("[data-testid=new-penalties-banner]").screenshot(path=f"{OUT}/admin-banner-yangi-zalyot.png")
    page.screenshot(path=f"{OUT}/admin-bosh-sahifa-banner.png")
    page.goto(BASE + "/zalyotchiklar/")
    page.screenshot(path=f"{OUT}/admin-zalyotchiklar-yangi.png", full_page=True)
    page.goto(BASE + "/zalyotchiklar/hisob/")
    page.screenshot(path=f"{OUT}/admin-zalyot-hisobi.png", full_page=True)
    # Excel varag'i — yuklab olinadi va jadval ko'rinishida chiziladi
    cookies = "; ".join(f"{c['name']}={c['value']}" for c in page.context.cookies())
    req = urllib.request.Request(BASE + "/zalyotchiklar/hisob/excel/", headers={"Cookie": cookies})
    wb = load_workbook(io.BytesIO(urllib.request.urlopen(req).read()))
    parts = []
    for ws in wb.worksheets:
        rows = "".join("<tr>" + "".join(f"<td>{html.escape(str(v if v is not None else ''))}</td>" for v in r) + "</tr>"
                       for r in ws.iter_rows(values_only=True))
        parts.append(f"<h2>Varaq: {html.escape(ws.title)}</h2><table>{rows}</table>")
    excel = page.context.new_page()
    excel.set_content("<style>body{font:14px Calibri,Arial;margin:16px}table{border-collapse:collapse}td{border:1px solid #bbb;"
                      "padding:3px 8px}tr:nth-child(2) td{font-weight:bold;background:#eef}h2{font-size:15px;margin:14px 0 6px}</style>"
                      + "".join(parts))
    excel.screenshot(path=f"{OUT}/excel-zalyot-hisobi.png", full_page=True)
    excel.close()
    page.goto(BASE + "/sozlamalar/tozalash/")
    page.screenshot(path=f"{OUT}/tozalash-korib-chiqish.png", full_page=True)
    page.fill("input[name=word_1]", "TOZALASH")
    page.fill("input[name=word_2]", "HAMMASI NOL")
    page.locator("[data-testid=clear-confirm]").screenshot(path=f"{OUT}/tozalash-tasdiq.png")
    page.context.close()
    b.close()
print("ok")
