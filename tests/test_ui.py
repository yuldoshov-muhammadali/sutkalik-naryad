"""5-bosqich: kirish, rollar, ruxsatlar matritsasi, bosh sahifa, akkauntlar."""

from __future__ import annotations

import datetime as dt
import re
from io import StringIO
from pathlib import Path

import pytest
from django.contrib.auth.models import AnonymousUser
from django.core.management import call_command
from django.http import HttpResponse
from django.test import RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from naryad import permissions as perms
from naryad.accounts import base_username, latin_slug
from naryad.models import Assignment, Cadet, DutyDay, LoginThrottle, User

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "tarix_namuna.xlsx"
PASSWORD = "Kuchli-parol-2026"

pytestmark = pytest.mark.django_db


# ---------------------------------------------------------------------------
# Foydalanuvchilar
# ---------------------------------------------------------------------------


@pytest.fixture
def users(db_cadets):
    girl = next(c for c in db_cadets if c.gender == "F")
    return {
        "admin": User.objects.create_user("kursboshligi", password=PASSWORD, role="admin"),
        "superuser": User.objects.create_superuser("super", password=PASSWORD),
        "cadet": User.objects.create_user("kursant", password=PASSWORD, role="cadet", cadet=girl),
    }


ROLES = ["anon", "admin", "superuser", "cadet"]


def as_role(client, users, role):
    client.logout()
    if role != "anon":
        client.force_login(users[role])
    return client


# ---------------------------------------------------------------------------
# Ruxsatlar matritsasi: har amal × har rol (role_required orqali)
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("action", sorted(perms.PERMISSIONS))
@pytest.mark.parametrize("role", ROLES)
def test_permission_matrix_decorator(action, role, users):
    view = perms.role_required(action)(lambda request: HttpResponse("ok"))
    request = RequestFactory().get("/sinov/")
    request.user = AnonymousUser() if role == "anon" else users[role]
    effective = "admin" if role == "superuser" else role
    if role == "anon":
        resp = view(request)
        assert resp.status_code == 302 and "/kirish/" in resp["Location"]
    elif effective in perms.PERMISSIONS[action][1]:
        assert view(request).status_code == 200
    else:
        from django.core.exceptions import PermissionDenied

        with pytest.raises(PermissionDenied):
            view(request)


# Mavjud sahifalar × rollar: kutilgan natija
PAGES = {
    # nom: (url, {rol: kutilgan})  "ok" = 200, "login" = login sahifasiga, "admin_login" = admin login'ga
    "home": (lambda: reverse("home"), {"anon": "login", "admin": "ok", "superuser": "ok", "cadet": "ok"}),
    "password_change": (
        lambda: reverse("password_change"),
        {"anon": "login", "admin": "ok", "superuser": "ok", "cadet": "ok"},
    ),
    "admin_index": (
        lambda: reverse("admin:index"),
        {"anon": "admin_login", "admin": "ok", "superuser": "ok", "cadet": "admin_login"},
    ),
    "admin_week_add": (
        lambda: reverse("admin:naryad_dutyweek_add"),
        {"anon": "admin_login", "admin": "ok", "superuser": "ok", "cadet": "admin_login"},
    ),
    "login": (lambda: reverse("login"), {"anon": "ok", "admin": "home", "superuser": "home", "cadet": "home"}),
}


@pytest.mark.parametrize("page", sorted(PAGES))
@pytest.mark.parametrize("role", ROLES)
def test_page_access_matrix(client, users, page, role):
    url_fn, expected = PAGES[page]
    resp = as_role(client, users, role).get(url_fn())
    exp = expected[role]
    if exp == "ok":
        assert resp.status_code == 200
    elif exp == "login":
        assert resp.status_code == 302 and resp["Location"].startswith(reverse("login"))
    elif exp == "admin_login":
        assert resp.status_code == 302 and resp["Location"].startswith(reverse("admin:login"))
    elif exp == "home":
        assert resp.status_code == 302 and resp["Location"] == reverse("home")


def test_logout_requires_post_and_works(client, users):
    as_role(client, users, "cadet")
    resp = client.post(reverse("logout"))
    assert resp.status_code == 302 and resp["Location"] == reverse("login")
    assert client.get(reverse("home")).status_code == 302


def test_roles_doc_matches_permission_table():
    text = (ROOT / "docs" / "05-rollar.md").read_text(encoding="utf-8")
    for action, (_, roles) in perms.PERMISSIONS.items():
        row = next((line for line in text.splitlines() if f"`{action}`" in line), None)
        assert row is not None, f"docs/05-rollar.md da {action} yo‘q"
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        marks = cells[-2:]  # admin, kursant
        assert marks == ["✅" if r in roles else "—" for r in (perms.ADMIN, perms.CADET)], action


def test_menu_depends_on_role(client, users):
    body = as_role(client, users, "admin").get(reverse("home")).content.decode()
    assert "Boshqaruv paneli" in body and "Haftalar" in body and "Sababli holatlar" in body
    body = as_role(client, users, "cadet").get(reverse("home")).content.decode()
    assert "Boshqaruv paneli" not in body and "Haftalar" not in body and "Mening naryadlarim" in body
    assert reverse("approved_days") not in body and "Tasdiqlangan naryadlar" not in body  # 14-bosqich


# ---------------------------------------------------------------------------
# Birinchi kirishda parolni almashtirish
# ---------------------------------------------------------------------------


def test_first_login_forces_password_change(client, users):
    user = users["cadet"]
    user.must_change_password = True
    user.save()
    resp = client.post(reverse("login"), {"username": "kursant", "password": PASSWORD})
    assert resp.status_code == 302
    resp = client.get(reverse("home"))
    assert resp.status_code == 302 and resp["Location"] == reverse("password_change")
    assert client.get(reverse("admin:index"))["Location"] == reverse("password_change")

    resp = client.post(
        reverse("password_change"),
        {"old_password": PASSWORD, "new_password1": "Yangi-maxfiy-parol-7", "new_password2": "Yangi-maxfiy-parol-7"},
    )
    assert resp.status_code == 302 and resp["Location"] == reverse("home")
    user.refresh_from_db()
    assert user.must_change_password is False
    assert client.get(reverse("home")).status_code == 200  # sessiya saqlanib qoldi


def test_password_change_errors_in_uzbek(client, users):
    as_role(client, users, "cadet")
    resp = client.post(reverse("password_change"), {"old_password": PASSWORD, "new_password1": "12345678", "new_password2": "12345678"})
    body = resp.content.decode()
    assert resp.status_code == 200
    assert "Parol faqat raqamlardan iborat." in body or "Bu parol juda keng tarqalgan." in body


# ---------------------------------------------------------------------------
# Login bloklanishi
# ---------------------------------------------------------------------------


def _login(client, username, password, url="login"):
    return client.post(reverse(url), {"username": username, "password": password, "next": "/"})


def test_login_blocked_after_5_failures(client, users):
    for i in range(4):
        resp = _login(client, "kursant", "xato")
        assert "Login yoki parol noto‘g‘ri" in resp.content.decode()
    resp = _login(client, "kursant", "xato")  # 5-xato → blok
    assert "15 daqiqaga bloklandi" in resp.content.decode()
    resp = _login(client, "kursant", PASSWORD)  # to'g'ri parol ham o'tmaydi
    assert resp.status_code == 200 and "bloklandi" in resp.content.decode()
    assert not resp.wsgi_request.user.is_authenticated
    # boshqa login bloklanmagan
    assert _login(client, "kursboshligi", PASSWORD).status_code == 302

    # 15 daqiqa o'tdi
    LoginThrottle.objects.filter(key="kursant").update(blocked_until=timezone.now() - dt.timedelta(seconds=1))
    client.logout()
    assert _login(client, "kursant", PASSWORD).status_code == 302
    assert not LoginThrottle.objects.filter(key="kursant").exists()  # muvaffaqiyatli kirish hisobni tozalaydi


def test_admin_login_is_throttled_too(client, users):
    for _ in range(5):
        _login(client, "kursboshligi", "xato", url="admin:login")
    resp = _login(client, "kursboshligi", PASSWORD, url="admin:login")
    assert resp.status_code == 200 and "bloklandi" in resp.content.decode()


@override_settings(LOGIN_MAX_FAILURES=2)
def test_old_failures_expire(client, users):
    _login(client, "kursant", "xato")
    LoginThrottle.objects.filter(key="kursant").update(last_failure_at=timezone.now() - dt.timedelta(minutes=20))
    resp = _login(client, "kursant", "xato")  # eski xato hisobga olinmaydi → hali blok yo'q
    assert "bloklandi" not in resp.content.decode()


# ---------------------------------------------------------------------------
# Bosh sahifa
# ---------------------------------------------------------------------------


@pytest.fixture
def imported(db_cadets):
    call_command("import_history", str(SAMPLE), stdout=StringIO())


def _now(day: dt.date, hour: int):
    return timezone.make_aware(dt.datetime.combine(day, dt.time(hour, 0)))


def test_staff_home_shows_week_and_next_step(client, users, imported, monkeypatch):
    monkeypatch.setattr("naryad.dashboard.timezone.now", lambda: _now(dt.date(2026, 9, 19), 10))  # 18-sentabr sutkasi davom etmoqda
    body = as_role(client, users, "admin").get(reverse("home")).content.decode()
    assert "13.09.2026" in body and "joriy hafta" in body
    assert body.count("Tasdiqlangan") >= 1 and body.count("Tuzilmagan") == 6
    assert "19-sentabr sutkasi qoralamasini tuzish" in body
    assert reverse("day_detail", args=["2026-09-19"]) in body  # sutka sahifasiga havola
    assert "hozir" in body


def test_admin_home_without_week_offers_opening(client, users):
    body = as_role(client, users, "admin").get(reverse("home")).content.decode()
    assert "2-kurs haftasini ochish" in body and reverse("week_list") in body
    body = as_role(client, users, "cadet").get(reverse("home")).content.decode()
    assert "Mening naryadlarim" in body and "2-kurs haftasini ochish" not in body


def test_cadet_home_shows_next_and_recent_duties(client, imported, monkeypatch):
    cadet = Cadet.objects.get(last_name="Panjiyeva")
    user = User.objects.create_user("panjiyeva", password=PASSWORD, role="cadet", cadet=cadet)
    client.force_login(user)

    monkeypatch.setattr("naryad.dashboard.timezone.now", lambda: _now(dt.date(2026, 9, 18), 9))  # 17-sentabr sutkasi
    body = client.get(reverse("home")).content.decode()
    assert "Mening navbatdagi naryadim" in body and "18-sentabr, juma 18:00" in body and "19-sentabr 18:00" in body
    assert "Qizlar yotoqxonasi bo‘yicha" in body and "Navbatchi" in body
    assert "Naryadga 9 soat 0 daqiqa qoldi" in body

    monkeypatch.setattr("naryad.dashboard.timezone.now", lambda: _now(dt.date(2026, 9, 25), 9))
    body = client.get(reverse("home")).content.decode()
    assert "Yaqin kunlarda tasdiqlangan naryadingiz yo‘q" in body
    assert reverse("my_card") in body  # tarix — “Mening tarixim” sahifasida


def test_cadet_sees_only_approved(client, db_cadets, users):
    from naryad import services
    from naryad.models import DutyWeek

    from .conftest import INITIAL_ROTATION

    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13), initial_rotation=INITIAL_ROTATION)
    girl_assignment = Assignment.objects.filter(duty_day=day, post_slot__post__code="qizlar").first()
    user = User.objects.create_user("q", password=PASSWORD, role="cadet", cadet=girl_assignment.cadet)
    client.force_login(user)
    body = client.get(reverse("home")).content.decode()
    assert "Yaqin kunlarda tasdiqlangan naryadingiz yo‘q" in body  # qoralama ko'rinmaydi
    assert DutyDay.objects.get().status == "draft"


# ---------------------------------------------------------------------------
# create_cadet_accounts
# ---------------------------------------------------------------------------


def test_latin_slug_and_usernames(db_groups):
    assert latin_slug("Sa’dullayev") == "sadullayev"
    assert latin_slug("G‘ayratov") == "gayratov"
    assert latin_slug("Dástan") == "dastan"
    c = Cadet(group=db_groups["BI-625"], last_name="Tursunmurodov", first_name="Bahodir Xon", gender="M")
    assert base_username(c) == "tursunmurodov.bahodir"


def test_create_cadet_accounts_is_idempotent(db_cadets, tmp_path):
    out = StringIO()
    call_command("create_cadet_accounts", "--output-dir", str(tmp_path), stdout=out)
    assert "Yangi akkauntlar: 100 ta" in out.getvalue()
    assert User.objects.filter(role="cadet").count() == 100
    assert User.objects.filter(role="cadet", must_change_password=True).count() == 100
    files = list(tmp_path.glob("akkauntlar_*.xlsx"))
    assert len(files) == 1

    ws = load_workbook(files[0]).active
    rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    assert len(rows) == 100
    login, password = next((r[3], r[4]) for r in rows if r[2].startswith("Sa’dullayev"))
    assert login == "sadullayev.behroz"
    assert re.fullmatch(r"[a-z0-9]{10}", password)
    assert User.objects.get(username=login).check_password(password)

    out = StringIO()
    call_command("create_cadet_accounts", "--output-dir", str(tmp_path), stdout=out)
    assert "Yangi akkauntlar: 0 ta" in out.getvalue() and "100 ta" in out.getvalue()
    assert User.objects.filter(role="cadet").count() == 100
    assert len(list(tmp_path.glob("akkauntlar_*.xlsx"))) == 1  # yangi fayl yozilmadi
    usernames = list(User.objects.values_list("username", flat=True))
    assert len(usernames) == len(set(usernames))


def test_duplicate_names_get_number_suffix(db_groups, tmp_path):
    for i in range(3):
        Cadet.objects.create(group=db_groups["BI-525"], last_name="Karimov", first_name="Ali", patronymic=f"P{i}", gender="M")
    User.objects.create_user("Karimov.Ali2", password=PASSWORD, role="admin")  # band login (katta-kichik harfsiz)
    call_command("create_cadet_accounts", "--output-dir", str(tmp_path), stdout=StringIO())
    names = sorted(User.objects.filter(cadet__isnull=False).values_list("username", flat=True))
    assert names == ["karimov.ali", "karimov.ali3", "karimov.ali4"]


def test_inactive_cadets_get_no_account(db_cadets, tmp_path):
    c = db_cadets[0]
    c.is_active = False
    c.deactivation_reason = "test"
    c.save()
    call_command("create_cadet_accounts", "--output-dir", str(tmp_path), stdout=StringIO())
    assert not User.objects.filter(cadet=c).exists()
    assert User.objects.filter(role="cadet").count() == 99


# ---------------------------------------------------------------------------
# Tarjimalar va import ogohlantirishi
# ---------------------------------------------------------------------------


def test_admin_login_page_in_uzbek(client):
    body = client.get(reverse("admin:login")).content.decode()
    assert "Kirish" in body and "Log in" not in body


def test_admin_buttons_in_uzbek(client, users):
    as_role(client, users, "admin")
    body = client.get(reverse("admin:naryad_absence_add")).content.decode()
    assert "Saqlash" in body and "Saqlash va yana qo‘shish" in body
    body = client.get(reverse("admin:naryad_cadet_changelist")).content.decode()
    assert "Qidirish" in body and "Filtr" in body


def test_import_warns_about_consecutive_days(db_cadets, tmp_path):
    from .test_commands import sample_rows, write_csv

    rows = sample_rows()
    day2 = [[dt.date(2026, 9, 19), *r[1:]] for r in rows]  # xuddi shu 20 kishi ertasiga ham
    path = write_csv(tmp_path / "tarix.csv", rows + day2)
    out = StringIO()
    call_command("import_history", str(path), stdout=out)
    text = out.getvalue()
    assert "Import muvaffaqiyatli" in text
    assert text.count("ketma-ket ikki sutka turgan") == 20
    assert "18.09.2026 va 19.09.2026" in text
    assert DutyDay.objects.count() == 2


def test_admin_can_unblock_login_and_force_password_change(client, users):
    for _ in range(5):
        _login(client, "kursant", "xato")
    throttle = LoginThrottle.objects.get(key="kursant")
    as_role(client, users, "admin")
    assert client.get(reverse("admin:naryad_loginthrottle_changelist")).status_code == 200
    resp = client.post(reverse("admin:naryad_loginthrottle_delete", args=[throttle.pk]), {"post": "yes"})
    assert resp.status_code == 302 and not LoginThrottle.objects.exists()
    body = client.get(reverse("admin:naryad_user_change", args=[users["cadet"].pk])).content.decode()
    assert "must_change_password" in body


def test_only_two_roles_exist():
    from naryad.models import UserRole

    assert set(UserRole.values) == {"admin", "cadet"}
    assert set(perms.ROLE_TITLES) == {"admin", "cadet"}
    for action, (_, roles) in perms.PERMISSIONS.items():
        assert roles <= {"admin", "cadet"}, action
