"""15-bosqich: zalyotchiklar — navbatdan tashqari jazo naryadi (CLAUDE.md 10-bo'lim, 61-band)."""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json

import pytest
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from naryad import days, penalties, services, stats
from naryad.engine import generate_day, pick_cadets
from naryad.engine.allocator import EXCLUDED_OFFENDER, History
from naryad.engine.models import KIND_PENALTY, Absence, Offender
from naryad.models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    Gender,
    InitialRotation,
    Penalty,
    Post,
    PostSlot,
    Setting,
    StudyGroup,
    User,
)
from naryad.services import ServiceError

from .conftest import INITIAL_ROTATION, load_roster, simulate

SUNDAY = dt.date(2026, 9, 13)
MONDAY = SUNDAY + dt.timedelta(days=1)
PASSWORD = "Kuchli-parol-2026"
# v1.5.0 dagi generatorning 84 sutkalik natijasi (zalyotchiksiz) — bayt-bayt o'zgarmasligi kerak
GOLDEN_84 = "8a3ebd425dc8ffa87ea9e082d2f211ffeccdb13f403677eadac6b2eba38a6fc7"


# ---------------------------------------------------------------------------
# Algoritm (toza Python)
# ---------------------------------------------------------------------------


def test_generator_without_offenders_is_byte_identical():
    results, _ = simulate(load_roster(), SUNDAY, 84)
    blob = json.dumps(
        [[str(r.day), sorted(r.group_by_post.items()),
          [[a.slot.post, a.slot.position, a.cadet_id, a.group] for a in r.assignments], r.warnings] for r in results],
        ensure_ascii=False,
    )
    assert hashlib.sha256(blob.encode()).hexdigest() == GOLDEN_84
    assert all(a.kind == "normal" for r in results for a in r.assignments)


@pytest.fixture(scope="module")
def roster():
    cadets = load_roster()
    return cadets, {c.id: c for c in cadets}


def male_of(cadets, group, index=0):
    return [c for c in cadets if c.group == group and c.gender == "M"][index]


def penalty_rows(result):
    return [a for a in result.assignments if a.kind == KIND_PENALTY]


def test_offender_from_markaz_or_patrul_group_stands_as_guard_at_noj1_or_yotoq1(roster):
    cadets, by_id = roster
    a = male_of(cadets, INITIAL_ROTATION["markaz"])  # guruhi bugun Markazda
    b = male_of(cadets, INITIAL_ROTATION["patrul"])  # guruhi bugun Patrulda
    res = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION,
                       offenders=[Offender(a.id, SUNDAY, 2), Offender(b.id, SUNDAY, 3)])
    pen = {x.cadet_id: x for x in penalty_rows(res)}
    assert set(pen) == {a.id, b.id}
    assert {x.slot.post for x in pen.values()} == {"noj1", "yotoq1"}
    assert all(x.slot.role == "posbon" for x in pen.values())
    # jazo o'rni egallangan joy guruhi bitta kam beradi, qolganlari — odatdagidek o'sha guruhdan
    for x in pen.values():
        group = res.group_by_post[x.slot.post]
        others = [y for y in res.assignments if y.slot.post == x.slot.post and y.kind == "normal"]
        assert len(others) == {"noj1": 3, "yotoq1": 2}[x.slot.post]  # 4−1 va 3−1
        assert all(by_id[y.cadet_id].group == group for y in others)
    assert not res.empty_slots


def test_at_most_one_per_post_and_two_per_day_third_waits(roster):
    cadets, _by_id = roster
    trio = [male_of(cadets, g) for g in ("BI-625", "YuP-125", "BI-325")]
    offs = [Offender(trio[0].id, SUNDAY, 2), Offender(trio[1].id, SUNDAY, 2), Offender(trio[2].id, SUNDAY + dt.timedelta(days=1), 2)]
    res = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION, offenders=offs)
    placed = [x.cadet_id for x in penalty_rows(res)]
    assert sorted(placed) == sorted([trio[0].id, trio[1].id])  # tanbeh sanasi eskisi birinchi
    assert any(trio[2].display_name in w and "keyingi sutkaga qoldi" in w for w in res.warnings)
    # ertasi kuni — uchinchisi qo'yiladi (ikkalasi kecha turgani uchun qo'yilmaydi)
    nxt = generate_day(SUNDAY + dt.timedelta(days=1), cadets, res.to_records(), (), INITIAL_ROTATION, offenders=offs)
    assert [x.cadet_id for x in penalty_rows(nxt)] == [trio[2].id]
    assert sum(1 for w in nxt.warnings if "kecha turgani uchun" in w) == 2


def test_per_post_setting_two(roster):
    cadets, _ = roster
    trio = [male_of(cadets, g) for g in ("BI-625", "YuP-125", "BI-325")]
    res = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION,
                       offenders=[Offender(c.id, SUNDAY, 1) for c in trio], penalty_per_post=2)
    rows = penalty_rows(res)
    assert len(rows) == 3 and [x.slot.post for x in rows].count("noj1") == 2


def test_yesterday_and_absent_offenders_are_not_placed(roster):
    cadets, by_id = roster
    saturday = SUNDAY - dt.timedelta(days=1)
    prev = generate_day(saturday, cadets, [], (), INITIAL_ROTATION)
    history = prev.to_records()
    a = by_id[next(x.cadet_id for x in prev.filled if x.slot.post == "markaz")]  # kecha Markazda turgan
    b = next(c for c in cadets if c.gender == "M" and c.id not in {r.cadet_id for r in history})
    absences = [Absence(b.id, SUNDAY, SUNDAY)]
    res = generate_day(SUNDAY, cadets, history, absences, None,
                       offenders=[Offender(a.id, SUNDAY, 2), Offender(b.id, SUNDAY, 2)])
    assert not penalty_rows(res)
    assert f"Zalyotchik {a.display_name} kecha turgani uchun bugun qo‘yilmadi." in res.warnings
    assert f"Zalyotchik {b.display_name} sababli holatda bo‘lgani uchun bugun qo‘yilmadi." in res.warnings


def test_post_alternates(roster):
    cadets, _ = roster
    a = male_of(cadets, "BI-625")
    posts = []
    for last in [None, "noj1", "yotoq1"]:
        res = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION, offenders=[Offender(a.id, SUNDAY, 3, last)])
        posts.append(penalty_rows(res)[0].slot.post)
    assert posts == ["noj1", "yotoq1", "noj1"]  # o'tgan safar 1-NO‘J — endi 1-Yotoqxona va aksincha


def test_penalty_duty_does_not_count_for_fairness_or_rotation(roster):
    cadets, _ = roster
    a = male_of(cadets, "BI-625")
    first = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION, offenders=[Offender(a.id, SUNDAY, 1)])
    history = first.to_records()
    assert any(r.kind == KIND_PENALTY for r in history)
    hist = History(history, before=MONDAY + dt.timedelta(days=1))
    assert hist.of(a.id) == [] and any(r.cadet_id == a.id for r in hist.all_records)  # adolatda yo'q
    # guruhlar aylanishi jazo yozuvi tufayli buzilmaydi
    nxt = generate_day(MONDAY, cadets, history, (), INITIAL_ROTATION)
    plain = generate_day(MONDAY, cadets, [r for r in history if r.kind == "normal"], (), INITIAL_ROTATION)
    assert nxt.group_by_post == plain.group_by_post


def test_active_offender_not_in_regular_turn_elsewhere_but_ok_at_noj1(roster):
    cadets, _ = roster
    old1, old2 = male_of(cadets, "BI-225", 1), male_of(cadets, "BI-425", 1)
    x = male_of(cadets, INITIAL_ROTATION["markaz"], 0)  # guruhi Markazda; jazo o'rni qolmaydi
    y = male_of(cadets, INITIAL_ROTATION["noj1"], 0)  # guruhi 1-NO‘J da
    offs = [Offender(old1.id, SUNDAY - dt.timedelta(days=5), 1), Offender(old2.id, SUNDAY - dt.timedelta(days=5), 1),
            Offender(x.id, SUNDAY, 1), Offender(y.id, SUNDAY, 1)]
    res = generate_day(SUNDAY, cadets, [], (), INITIAL_ROTATION, offenders=offs)
    assert res.picks["markaz"].excluded.get(x.id) == EXCLUDED_OFFENDER
    assert y.id not in res.picks["noj1"].excluded  # o'z guruhi bilan oddiy navbatda tura oladi
    assert any(y.id == s.cadet_id for s in res.picks["noj1"].ranking)
    # pick_cadets: offenders parametri berilmasa — avvalgidek
    hist = History([], before=SUNDAY)
    pool = [c for c in cadets if c.group == x.group and c.gender == "M"]
    assert x.id not in pick_cadets(pool, "markaz", len(pool), SUNDAY, hist).excluded
    assert pick_cadets(pool, "markaz", 1, SUNDAY, hist, offenders=[x.id]).excluded[x.id] == EXCLUDED_OFFENDER


# ---------------------------------------------------------------------------
# Baza: zalyot, tasdiqlash, bekor qilish
# ---------------------------------------------------------------------------


@pytest.fixture
def admin(db_cadets):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=SUNDAY)
    Setting.load()
    return User.objects.create_user("kursboshligi", password=PASSWORD, role="admin")


def boy(group, index=0) -> Cadet:
    return Cadet.objects.filter(group__code=group, gender=Gender.MALE, duty_exempt=False, is_active=True).order_by("list_number")[index]


def approve(day, admin):
    duty_day = DutyDay.objects.get(date=day)
    return days.approve(day, admin, duty_day.version)


def test_cannot_make_girl_exempt_or_inactive_an_offender(admin):
    girl = Cadet.objects.filter(gender=Gender.FEMALE).first()
    exempt = Cadet.objects.filter(duty_exempt=True).first() or boy("BI-225", 2)
    if not exempt.duty_exempt:
        exempt.duty_exempt, exempt.duty_position = True, "Guruh komandiri"
        exempt.save()
    inactive = boy("BI-325", 3)
    inactive.is_active = False
    inactive.deactivation_reason = "Akademik ta’til"
    inactive.save()
    for cadet in (girl, exempt, inactive):
        with pytest.raises(ServiceError):
            penalties.add(cadet, "tartib", "x", SUNDAY, 2, admin)
    assert not Penalty.objects.exists()
    with pytest.raises(ServiceError):
        penalties.add(boy("BI-625"), "tartib", "   ", SUNDAY, 2, admin)  # izoh majburiy
    with pytest.raises(ServiceError):
        penalties.add(boy("BI-625"), "tartib", "x", SUNDAY, 11, admin)  # 1–10


def test_second_penalty_adds_to_active_with_separate_log(admin):
    cadet = boy("BI-625")
    p1 = penalties.add(cadet, "tartib", "Kechikdi", SUNDAY, 2, admin)
    p2 = penalties.add(cadet, "uxlash", "Postda uxladi", MONDAY, 3, admin)
    assert p1.pk == p2.pk and Penalty.objects.get().count == 5
    logs = ChangeLog.objects.filter(cadet=cadet, action="penalty_add")
    assert logs.count() == 2 and any("Naryadda uxlab qolish: Postda uxladi" == e.reason for e in logs)


def test_served_counts_only_on_approval_and_reverts_on_reopen(admin):
    offender = boy(INITIAL_ROTATION["markaz"])
    p = penalties.add(offender, "tartib", "Kechikdi", SUNDAY, 1, admin)
    day = services.generate_draft(SUNDAY, admin)
    a = Assignment.objects.get(duty_day=day, cadet=offender)
    assert a.kind == "penalty" and a.penalty_id == p.pk and a.post_slot.post.code in ("noj1", "yotoq1")
    assert a.post_slot.role == "posbon"
    assert penalties.with_served(Penalty.objects.filter(pk=p.pk)).get().served == 0  # qoralama — hisobga kirmaydi
    approve(SUNDAY, admin)
    p.refresh_from_db()
    assert p.status == "done" and p.served == 1  # soni tugadi — avtomatik “bajarildi”
    days.reopen(SUNDAY, "tekshirish", admin, DutyDay.objects.get(date=SUNDAY).version)
    p.refresh_from_db()
    assert p.status == "active" and p.served == 0  # hisob qaytdi
    actions = list(ChangeLog.objects.filter(cadet=offender).values_list("action", flat=True))
    assert {"penalty_add", "penalty_done", "penalty_resume"} <= set(actions)


def test_cancel_requires_reason_and_logs(admin):
    p = penalties.add(boy("BI-625"), "boshqa", "Izoh", SUNDAY, 2, admin)
    with pytest.raises(ServiceError):
        penalties.cancel(p, " ", admin)
    penalties.cancel(p, "Buyruq bekor qilindi", admin)
    p.refresh_from_db()
    assert p.status == "cancelled" and ChangeLog.objects.filter(action="penalty_cancel", reason="Buyruq bekor qilindi").exists()
    assert penalties.engine_offenders(SUNDAY) == []


def test_replacement_rules_for_penalty_slot(admin):
    a = boy(INITIAL_ROTATION["markaz"])
    penalties.add(a, "tartib", "x", SUNDAY, 2, admin)
    day = services.generate_draft(SUNDAY, admin)
    pen = Assignment.objects.get(duty_day=day, kind="penalty")
    busy = set(day.assignments.values_list("cadet_id", flat=True))
    other = next(c for c in Cadet.objects.filter(gender=Gender.MALE, is_active=True, duty_exempt=False, group__code=INITIAL_ROTATION["patrul"])
                 .order_by("list_number") if c.pk not in busy)
    penalties.add(other, "tartib", "y", MONDAY, 1, admin)  # tuzilgandan keyin kiritildi
    cands = days.replacement_candidates(day, pen.post_slot)
    assert cands[0].offender and cands[0].cadet == other and not cands[0].excluded
    # boshqa guruhning oddiy kursanti — yo'q (guruh qoidasi)
    stranger = boy(INITIAL_ROTATION["oquv"], 2)
    assert days.replacement_error(day, pen.post_slot, stranger)
    # faol zalyotchikni boshqa joyga (Markaz) qo'lda qo'yib bo'lmaydi
    markaz_slot = PostSlot.objects.get(post__code="markaz")
    assert "faol zalyotchik" in days.replacement_error(day, markaz_slot, other)
    # jazo o'rniga boshqa zalyotchik — jazo naryadi bo'lib qoladi
    days.replace_cadet(SUNDAY, pen.post_slot_id, other.pk, "boshqa", "almashtirildi", admin, day.version)
    new = Assignment.objects.get(duty_day=day, post_slot=pen.post_slot)
    assert new.kind == "penalty" and new.penalty.cadet == other
    # shu joy guruhi kursanti qo'yilsa — oddiy naryad
    day.refresh_from_db()
    group_cadet = next(c.cadet for c in days.replacement_candidates(day, pen.post_slot) if not c.excluded and not c.offender)
    days.replace_cadet(SUNDAY, pen.post_slot_id, group_cadet.pk, "boshqa", "guruhdan", admin, day.version)
    assert Assignment.objects.get(duty_day=day, post_slot=pen.post_slot).kind == "normal"


def test_penalty_role_cannot_be_swapped(admin):
    penalties.add(boy(INITIAL_ROTATION["markaz"]), "tartib", "x", SUNDAY, 2, admin)
    day = services.generate_draft(SUNDAY, admin)
    pen = Assignment.objects.get(duty_day=day, kind="penalty")
    duty = Assignment.objects.get(duty_day=day, post_slot__post=pen.post_slot.post, post_slot__position=1)
    with pytest.raises(ServiceError, match="posbon"):
        days.swap_slots(SUNDAY, pen.post_slot_id, duty.post_slot_id, "boshqa", "x", admin, day.version)


# ---------------------------------------------------------------------------
# Statistika, Excel, PDF, sahifalar
# ---------------------------------------------------------------------------


def test_stats_exclude_penalty_duties_and_excel_has_column(client, admin):
    offender = boy(INITIAL_ROTATION["markaz"])
    penalties.add(offender, "tartib", "x", SUNDAY, 1, admin)
    services.generate_draft(SUNDAY, admin)
    approve(SUNDAY, admin)
    rows = {r.cadet.pk: r for r in stats.cadet_rows(offender.group.code, stats.Period())}
    assert rows[offender.pk].total == 0 and rows[offender.pk].penalty == 1
    fair = {f.pool: f for f in stats.fairness_all(stats.Period())}[offender.group.code]
    boys_normal = Assignment.objects.filter(group=offender.group, cadet__gender="M", kind="normal", duty_day__status="approved")
    assert fair.total == boys_normal.count()  # jazo naryadi adolat hisobiga kirmaydi
    client.force_login(admin)
    assert 'data-testid="penalty-count"' in client.get(reverse("stats")).content.decode()
    wb = load_workbook(io.BytesIO(client.get(reverse("stats_excel")).content))
    ws = wb[f"Guruh {offender.group.code}"]
    header = [c.value for c in ws[3]]
    col = header.index("Jazo naryadlari (navbatdan tashqari)")
    values = {r[0].value: r[col].value for r in ws.iter_rows(min_row=4)}
    assert values[offender.full_name] == 1
    assert "Zalyotchiklar" in wb.sheetnames


def test_pdf_has_no_zalyot_word(client, admin):
    penalties.add(boy(INITIAL_ROTATION["markaz"]), "tartib", "x", SUNDAY, 1, admin)
    services.generate_draft(SUNDAY, admin)
    approve(SUNDAY, admin)
    client.force_login(admin)
    pdf = client.get(reverse("day_pdf", args=[SUNDAY.isoformat()])).content
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages).lower()
    assert "zalyot" not in text and "jazo" not in text and "navbatdan tashqari" not in text
    assert "posbon" in text


def test_day_page_marks_penalty_slot(client, admin):
    penalties.add(boy(INITIAL_ROTATION["markaz"]), "tartib", "x", SUNDAY, 1, admin)
    services.generate_draft(SUNDAY, admin)
    client.force_login(admin)
    body = client.get(reverse("day_detail", args=[SUNDAY.isoformat()])).content.decode()
    assert body.count('data-testid="penalty-row"') == 1 and ">Zalyot<" in body


def test_penalty_page_admin_only_and_works(client, admin):
    cadet_user = User.objects.create_user("k", password=PASSWORD, role="cadet", cadet=boy("BI-525"))
    client.force_login(cadet_user)
    assert client.get(reverse("penalty_list")).status_code == 403
    assert "Zalyotchiklar" not in client.get(reverse("home")).content.decode()
    client.force_login(admin)
    target = boy("BI-625")
    resp = client.post(reverse("penalty_list"), {
        "cadet": target.pk, "reason_type": "uxlash", "order_date": "2026-09-12", "count": 3, "note": "Postda uxladi",
    })
    assert resp.status_code == 302
    body = client.get(reverse("penalty_list")).content.decode()
    assert target.short_name in body and "/ 3" in body and "Postda uxladi" in body
    p = Penalty.objects.get()
    assert client.post(reverse("penalty_cancel", args=[p.pk]), {"reason": "Bekor"}).status_code == 302
    p.refresh_from_db()
    assert p.status == "cancelled"
    card = client.get(reverse("cadet_detail", args=[target.pk])).content.decode()
    assert 'data-testid="penalty-history"' in card and "Postda uxladi" in card


def test_cadet_home_shows_extra_duty_and_progress(client, admin, monkeypatch):
    from django.utils import timezone

    offender = boy(INITIAL_ROTATION["markaz"])
    penalties.add(offender, "tartib", "x", SUNDAY - dt.timedelta(days=1), 3, admin)
    services.generate_draft(SUNDAY, admin)
    approve(SUNDAY, admin)
    monkeypatch.setattr("django.utils.timezone.now", lambda: timezone.make_aware(dt.datetime(2026, 9, 13, 10, 0)))
    user = User.objects.create_user("z", password=PASSWORD, role="cadet", cadet=offender)
    client.force_login(user)
    body = client.get(reverse("home")).content.decode()
    assert "Navbatdan tashqari naryad" in body and 'data-testid="duty-penalty"' in body
    assert "Jazo naryadlari: 3 tadan 1 tasi o‘tildi" in body


def test_post_alternates_through_database(admin):
    offender = boy(INITIAL_ROTATION["markaz"])
    penalties.add(offender, "tartib", "x", SUNDAY, 3, admin)
    posts = []
    for i in range(5):
        day = SUNDAY + dt.timedelta(days=i)
        services.generate_draft(day, admin)
        approve(day, admin)
        a = Assignment.objects.filter(duty_day__date=day, cadet=offender, kind="penalty").select_related("post_slot__post").first()
        if a:
            posts.append(a.post_slot.post.code)
    assert posts == ["noj1", "yotoq1", "noj1"]  # kun ora (kecha turgan qo'yilmaydi), joy almashadi
    assert Penalty.objects.get().status == "done"


# ---------------------------------------------------------------------------
# 1.6.0 ochiq savollari bo'yicha qarorlar (CLAUDE.md 62-band)
# ---------------------------------------------------------------------------


def test_active_offender_cannot_be_exempted(client, admin):
    from naryad import cadets as cadet_service

    offender = boy("BI-625")
    p = penalties.add(offender, "tartib", "x", SUNDAY, 2, admin)
    with pytest.raises(ValueError):
        cadet_service.set_exempt(offender, "Guruh komandiri", "buyruq", admin)
    client.force_login(admin)
    assert 'data-testid="exempt-blocked"' in client.get(reverse("cadet_detail", args=[offender.pk])).content.decode()
    client.post(reverse("cadet_exempt", args=[offender.pk]), {"position": "Guruh komandiri", "reason": "buyruq"})
    offender.refresh_from_db()
    assert not offender.duty_exempt
    penalties.cancel(p, "Bekor", admin)  # zalyot yakunlangach — ozod qilish mumkin
    client.post(reverse("cadet_exempt", args=[offender.pk]), {"position": "Guruh komandiri", "reason": "buyruq"})
    offender.refresh_from_db()
    assert offender.duty_exempt


def test_cancelled_penalty_in_draft_blocks_approval(admin):
    offender = boy(INITIAL_ROTATION["markaz"])
    p = penalties.add(offender, "tartib", "x", SUNDAY, 2, admin)
    day = services.generate_draft(SUNDAY, admin)
    penalties.cancel(p, "Buyruq bekor qilindi", admin)
    blocking = [i for i in days.day_issues(day) if i.blocks_approval and i.kind == "penalty"]
    assert len(blocking) == 1
    with pytest.raises(ServiceError, match="zalyoti bekor qilingan"):
        days.approve(SUNDAY, admin, day.version)
    # jazo o'rniga shu guruh kursanti qo'yilsa — tasdiqlanadi
    pen = Assignment.objects.get(duty_day=day, kind="penalty")
    group_cadet = next(c.cadet for c in days.replacement_candidates(day, pen.post_slot) if not c.excluded and not c.offender)
    days.replace_cadet(SUNDAY, pen.post_slot_id, group_cadet.pk, "boshqa", "zalyot bekor qilindi", admin, day.version)
    day.refresh_from_db()
    days.approve(SUNDAY, admin, day.version)
