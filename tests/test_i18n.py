"""13-bosqich: uch til (o‘zbek, rus, ingliz) — CLAUDE.md 10-bo‘lim, 59-band."""

# ruff: noqa: F811 — `world` fixture test_security dan import qilinadi va parametr sifatida keladi

from __future__ import annotations

import datetime as dt
import html as htmllib
import io
import re
from pathlib import Path

import polib
import pytest
from django.urls import reverse
from django.utils import translation
from openpyxl import load_workbook
from pypdf import PdfReader

from naryad import i18n
from naryad.models import Absence, Cadet, ChangeLog, DutyWeek, Penalty, StudyGroup, User

from .test_security import POST_ONLY, SUNDAY, urls, world  # noqa: F401 — umumiy fixture

ROOT = Path(__file__).resolve().parents[1]
LANGS = ("uz", "ru", "en")
pytestmark = pytest.mark.django_db


def as_user(client, user, lang):
    user.language = lang
    user.save(update_fields=["language"])
    client.force_login(user)
    return client


# ---------------------------------------------------------------------------
# Har bir sahifa uch tilda
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("lang", LANGS)
@pytest.mark.parametrize("who", ["admin", "cadet"])
def test_every_page_opens_in_every_language(client, world, lang, who):
    as_user(client, world[who], lang)
    for key, url in urls(world).items():
        if key in POST_ONLY or key == "login":
            continue
        response = client.get(url)
        assert response.status_code < 500, (lang, who, key)
        if response.status_code == 200 and response["Content-Type"].startswith("text/html"):
            body = response.content.decode()
            if "<html" in body:
                assert f'<html lang="{lang}"' in body, (lang, who, key)


@pytest.mark.parametrize("lang", LANGS)
def test_login_and_error_pages_in_every_language(client, db, lang):
    client.cookies[i18n.settings.LANGUAGE_COOKIE_NAME] = lang
    body = client.get(reverse("login")).content.decode()
    assert f'<html lang="{lang}"' in body and 'data-testid="language-switch"' in body
    assert client.get("/mavjud-emas/").status_code == 404


# ---------------------------------------------------------------------------
# Ruscha va inglizcha sahifada o'zbekcha qolib ketmagan
# ---------------------------------------------------------------------------

UZBEK_WORDS = {
    "va", "yoki", "uchun", "bilan", "bo‘yicha", "sutka", "sutkalar", "sutkasi", "naryad", "naryadlar", "naryadi",
    "kursantlar", "guruh", "guruhi", "hafta", "haftalar", "haftasi", "tasdiqlash", "tasdiqlangan", "qoralama", "tuzish",
    "almashtirish", "sabab", "jami", "yo‘q", "kirish", "chiqish", "sahifa", "keyingi", "joriy", "holat", "holatlar",
    "ogohlantirish", "statistika", "kursant", "lavozim", "joy", "joylar", "sana", "bosh", "parol", "faol", "emas",
    "ozod", "qo‘shish", "tahrirlash", "saqlash", "bekor", "qilish", "ko‘rish", "yuklab", "olish", "chop", "etish",
}
LATIN_ALLOWED_RU = {"pdf", "excel", "xlsx", "uz", "ru", "en", "min", "max", "csv", "htmx", "id"}


def visible_text(html: str) -> str:
    html = re.sub(r"(?s)<(script|style)\b.*?</\1>", " ", html)
    html = re.sub(r"(?s)<!--.*?-->", " ", html)
    # atributlardagi matnlar ham foydalanuvchiga ko'rinadi (title, aria-label, placeholder, alt)
    attrs = " ".join(re.findall(r'(?:title|aria-label|placeholder|alt)="([^"]*)"', html))
    text = re.sub(r"(?s)<[^>]+>", " ", html) + " " + attrs
    return htmllib.unescape(text)


def data_strings() -> list[str]:
    """Ma'lumotlar (tarjima qilinmaydi): F.I.Sh., guruh kodlari, login, admin yozgan sabab/izohlar, jurnal matnlari."""
    out = []
    for c in Cadet.objects.all():
        out += [c.full_name, c.short_name, c.last_name, c.first_name, c.patronymic, c.deactivation_reason, c.duty_position]
    out += list(StudyGroup.objects.values_list("code", flat=True))
    out += list(User.objects.values_list("username", flat=True))
    for e in ChangeLog.objects.all():
        out += [e.details, e.reason]
    out += list(Absence.objects.values_list("note", flat=True))
    out += list(DutyWeek.objects.values_list("note", flat=True))
    for note, reason in Penalty.objects.values_list("note", "cancel_reason"):  # zalyot izohi — admin yozgan matn
        out += [note, reason]
    return sorted({s for s in out if s}, key=len, reverse=True)


# til tanlagichdagi har tilning o'z nomi va lotincha ism namunalari (ismlar lotincha saqlanadi) — tarjima qilinmaydi
LANGUAGE_NAMES = ["O‘zbekcha", "Русский", "English", "Sh, Ch, G‘, O‘", "Shavkatov Sh, G‘ayratov G‘",
                  "Bahodir o‘g‘li, Oybek qizi, Davronovich", "CLAUDE.md"]


def leftovers(text: str, lang: str) -> list[str]:
    for s in data_strings() + LANGUAGE_NAMES:
        text = text.replace(s, " ")
    words = re.findall(r"[A-Za-zʻ‘’'-]+", text)
    found = []
    for w in words:
        lw = w.lower().strip("-’'")
        if "o‘" in lw or "g‘" in lw or lang == "en" and lw in UZBEK_WORDS or lang == "ru" and len(lw) >= 3 and lw not in LATIN_ALLOWED_RU and not re.fullmatch(r"[a-z]{1,3}\d*", lw):
            found.append(w)
    return found


MAIN_PAGES = ["home", "week_list", "day_detail:approved", "day_detail:draft", "day_log", "cadet_list", "cadet_detail:other",
              "absence_list", "penalty_list", "stats", "approved_days", "password_change", "cadet_create", "cadet_import"]


@pytest.mark.parametrize("lang", ["ru", "en"])
def test_no_uzbek_left_on_main_pages(client, world, lang):
    as_user(client, world["admin"], lang)
    all_urls = urls(world)
    problems = {}
    for key in MAIN_PAGES:
        found = leftovers(visible_text(client.get(all_urls[key]).content.decode()), lang)
        if found:
            problems[key] = sorted(set(found))
    as_user(client, world["cadet"], lang)
    for key in ["home", "home_status", "cadet_detail:own"]:  # sutka sahifalari kursantga yopiq (14-bosqich)
        found = leftovers(visible_text(client.get(all_urls[key]).content.decode()), lang)
        if found:
            problems["kursant:" + key] = sorted(set(found))
    client.logout()
    client.cookies[i18n.settings.LANGUAGE_COOKIE_NAME] = lang
    found = leftovers(visible_text(client.get(reverse("login")).content.decode()), lang)
    if found:
        problems["login"] = sorted(set(found))
    assert not problems, problems


# ---------------------------------------------------------------------------
# .po fayllar to'liq
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("lang", ["ru", "en"])
def test_po_files_complete(lang):
    po = polib.pofile(str(ROOT / "locale" / lang / "LC_MESSAGES" / "django.po"))
    assert len(po) > 300
    assert not po.untranslated_entries(), [e.msgid for e in po.untranslated_entries()][:10]
    assert not po.fuzzy_entries(), [e.msgid for e in po.fuzzy_entries()][:10]
    for e in po:
        if e.msgid_plural:
            assert all(v.strip() for v in e.msgstr_plural.values()), e.msgid
        # o'rinbosarlar (%(name)s) tarjimada ham bo'lishi kerak
        for ph in re.findall(r"%\(\w+\)[sd]", e.msgid):
            text = e.msgstr or " ".join(e.msgstr_plural.values())
            assert ph in text, (lang, e.msgid, ph)
    mo = ROOT / "locale" / lang / "LC_MESSAGES" / "django.mo"
    assert mo.exists() and mo.stat().st_mtime >= (ROOT / "locale" / lang / "LC_MESSAGES" / "django.po").stat().st_mtime - 5


# ---------------------------------------------------------------------------
# Til tanlash qoidalari
# ---------------------------------------------------------------------------


def test_default_is_uzbek_even_if_browser_says_russian(client, db):
    body = client.get(reverse("login"), HTTP_ACCEPT_LANGUAGE="ru-RU,ru;q=0.9,en;q=0.8").content.decode()
    assert '<html lang="uz"' in body and "Tizimga kirish" in body


def test_admin_panel_ignores_browser_language(client, world):
    client.force_login(world["admin"])
    body = client.get("/admin/", HTTP_ACCEPT_LANGUAGE="ru").content.decode()
    assert "Sutkalik naryad" in body and "Администрирование" not in body


def test_choice_persists_across_logout_and_devices(client, world):
    from django.test import Client

    admin = world["admin"]
    admin.set_password("Kuchli-parol-2026")
    admin.save()
    # kirish sahifasida ruscha tanlandi (cookie)
    client.post(reverse("set_language"), {"language": "ru", "next": reverse("login")})
    assert '<html lang="ru"' in client.get(reverse("login")).content.decode()
    client.post(reverse("login"), {"username": admin.username, "password": "Kuchli-parol-2026"})
    client.get(reverse("home"))
    admin.refresh_from_db()
    assert admin.language == "ru"
    # inglizchaga o'tkazildi, chiqildi — kirish sahifasi inglizcha qoladi
    client.post(reverse("set_language"), {"language": "en", "next": reverse("home")})
    admin.refresh_from_db()
    assert admin.language == "en"
    client.post(reverse("logout"))
    assert '<html lang="en"' in client.get(reverse("login")).content.decode()
    # boshqa qurilma (cookie yo'q): kirgach — inglizcha
    other = Client()
    other.post(reverse("login"), {"username": admin.username, "password": "Kuchli-parol-2026"})
    assert '<html lang="en"' in other.get(reverse("home")).content.decode()


def test_choice_on_login_page_overrides_saved_language(client, world):
    """Kirish sahifasida tugma bilan tanlangan til saqlangan tildan ustun (va saqlanadi)."""
    admin = world["admin"]
    admin.set_password("Kuchli-parol-2026")
    admin.language = "en"
    admin.save()
    client.cookies[i18n.settings.LANGUAGE_COOKIE_NAME] = "en"  # eski cookie
    client.post(reverse("set_language"), {"language": "uz", "next": reverse("login")})
    client.post(reverse("login"), {"username": admin.username, "password": "Kuchli-parol-2026"})
    assert '<html lang="uz"' in client.get(reverse("home")).content.decode()
    admin.refresh_from_db()
    assert admin.language == "uz"
    # keyingi so'rovlarda ham (sessiya belgisi bir marta ishlatiladi)
    assert '<html lang="uz"' in client.get(reverse("home")).content.decode()


def test_old_cookie_does_not_override_saved_language(client, world):
    """Oldingi tashrifdan qolgan cookie (tugma bosilmagan) saqlangan tilni o'zgartirmaydi."""
    admin = world["admin"]
    admin.language = "ru"
    admin.save()
    client.cookies[i18n.settings.LANGUAGE_COOKIE_NAME] = "en"
    client.force_login(admin)
    assert '<html lang="ru"' in client.get(reverse("home")).content.decode()


def test_set_language_rules(client, db):
    assert client.get(reverse("set_language")).status_code == 405
    r = client.post(reverse("set_language"), {"language": "ru", "next": "https://yomon.example.com/"})
    assert r.status_code == 302 and r["Location"] == "/"
    assert '<html lang="ru"' in client.get(reverse("login")).content.decode()
    r = client.post(reverse("set_language"), {"language": "fr", "next": "/kirish/"})  # noma'lum til — e'tiborsiz
    assert r.status_code == 302 and i18n.settings.LANGUAGE_COOKIE_NAME not in r.cookies
    assert '<html lang="ru"' in client.get(reverse("login")).content.decode()


def test_language_switch_is_csrf_protected(db):
    from django.test import Client

    strict = Client(enforce_csrf_checks=True)
    assert strict.post(reverse("set_language"), {"language": "ru"}).status_code == 403


# ---------------------------------------------------------------------------
# Rasmiy hujjatlar doim o'zbekcha; sanalar har tilda
# ---------------------------------------------------------------------------


def test_pdf_stays_uzbek_in_english_session(client, world):
    as_user(client, world["admin"], "en")
    pdf = client.get(reverse("day_pdf", args=[SUNDAY.isoformat()]))
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "kunidan" in text and "kuniga" in text and "navbatchi" in text
    assert "Duty" not in text and "duty" not in text


def test_excel_stays_uzbek_in_russian_session(client, world):
    as_user(client, world["admin"], "ru")
    wb = load_workbook(io.BytesIO(client.get(reverse("stats_excel")).content))
    assert "Adolat" in wb.sheetnames and wb["Adolat"]["A1"].value == "Davr: Butun tarix"
    assert wb["Adolat"]["A3"].value == "Hovuz"


@pytest.mark.parametrize(
    "lang, expected",
    [("uz", "14-oktabr 2026"), ("ru", "14 октября 2026"), ("en", "14 October 2026")],
)
def test_long_date_formats(lang, expected):
    with translation.override(lang):
        assert i18n.long_date(dt.date(2026, 10, 14)) == expected


def test_clock_texts_come_from_server_in_selected_language(client, world):
    as_user(client, world["admin"], "ru")
    body = client.get(reverse("home")).content.decode()
    assert 'data-months="января,' in body and "октября" in body
    assert 'data-weekdays="Воскресенье,Понедельник' in body
    assert "{h}" in body and "{m}" in body  # “Smena almashishiga …” shablonlari JS uchun
    assert "<script>" not in body
