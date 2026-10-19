"""19-bosqich: guruh komandiri roli va “Mening guruhim” paneli — faqat ko'rish (CLAUDE.md 67-band)."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from django.core.exceptions import ValidationError
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from naryad import days, penalties, services
from naryad.models import (
    Absence,
    Assignment,
    Cadet,
    DutyDay,
    DutyWeek,
    Gender,
    InitialRotation,
    Post,
    SickDocument,
    StudyGroup,
    User,
)

from .conftest import INITIAL_ROTATION

SUNDAY = dt.date(2026, 9, 13)
MONDAY, TUESDAY = SUNDAY + dt.timedelta(days=1), SUNDAY + dt.timedelta(days=2)
PASSWORD = "Kuchli-parol-2026"
GROUPS = ["YuP-125", "BI-225", "BI-325", "BI-425", "BI-525", "BI-625"]
ILLNESS = "Gripp, yuqori harorat"
pytestmark = pytest.mark.django_db


def at(day: dt.date, hour: int = 20) -> dt.datetime:
    return timezone.make_aware(dt.datetime.combine(day, dt.time(hour, 0)))


@pytest.fixture
def admin(db_cadets, monkeypatch):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=SUNDAY)
    monkeypatch.setattr("django.utils.timezone.now", lambda: at(MONDAY, 9))  # dushanba ertalab — yakshanba sutkasi davom etmoqda
    return User.objects.create_user("kursboshligi", password=PASSWORD, role="admin")


def commander_cadet(group: str) -> Cadet:
    return Cadet.objects.get(group__code=group, duty_position=f"{group} guruh komandiri")


def make_commander(group: str) -> User:
    user = User(username=f"komandir.{group.lower()}", role="commander", cadet=commander_cadet(group))
    user.set_password(PASSWORD)
    user.full_clean()
    user.save()
    return user


def generate_and_approve(day: dt.date, admin) -> DutyDay:
    draft = services.generate_draft(day, admin)
    return days.approve(day, admin, draft.version)


def page(client, name: str, *args, **params) -> str:
    response = client.get(reverse(name, args=list(args)), params)
    assert response.status_code == 200, (name, response.status_code)
    return response.content.decode()


# ---------------------------------------------------------------------------
# 1. Rol
# ---------------------------------------------------------------------------


def test_six_group_commanders_can_get_role(admin):
    for group in GROUPS:
        user = make_commander(group)
        assert user.is_commander and user.cadet.group.code == group and not user.is_staff


@pytest.mark.parametrize("position", ["Kurs starshinasi", "Kurs pisari", "Kaptyorchik", ""])
def test_commander_role_only_for_group_commander_position(admin, position):
    cadet = Cadet.objects.filter(duty_position=position).first() if position else Cadet.objects.filter(duty_exempt=False).first()
    user = User(username="x", role="commander", cadet=cadet)
    with pytest.raises(ValidationError, match="guruh komandiri lavozimida emas"):
        user.full_clean(exclude=["password"])


def test_deputy_is_not_commander_and_one_commander_per_group(admin):
    from naryad import cadets as cadet_service

    deputy = Cadet.objects.filter(group__code="BI-525", duty_exempt=False, gender=Gender.MALE).first()
    cadet_service.set_exempt(deputy, "BI-525 guruh komandiri o‘rinbosari", "buyruq", admin)
    assert not deputy.is_group_commander
    with pytest.raises(ValidationError, match="guruh komandiri lavozimida emas"):
        User(username="orinbosar", role="commander", cadet=deputy).full_clean(exclude=["password"])
    make_commander("BI-525")
    with pytest.raises(ValidationError):  # bir kursantga ikkinchi akkaunt ham bo'lmaydi
        User(username="y", role="commander", cadet=commander_cadet("BI-525")).full_clean(exclude=["password"])


def test_second_active_commander_in_group_rejected(admin):
    from naryad import cadets as cadet_service

    make_commander("BI-325")
    extra = Cadet.objects.filter(group__code="BI-325", duty_exempt=False, gender=Gender.MALE).first()
    cadet_service.set_exempt(extra, "BI-325 guruh komandiri", "ikkinchi buyruq", admin)
    user = User(username="ikkinchi", role="commander", cadet=extra)
    with pytest.raises(ValidationError, match="faol komandir akkaunti allaqachon bor"):
        user.full_clean(exclude=["password"])


def test_admin_form_rejects_wrong_cadet(client, admin):
    admin.is_superuser = True
    admin.save()
    client.force_login(admin)
    starshina = Cadet.objects.get(duty_position="Kurs starshinasi")
    response = client.post(reverse("admin:naryad_user_add"), {
        "username": "notogri", "password1": "Juda-kuchli-parol-77", "password2": "Juda-kuchli-parol-77",
        "role": "commander", "cadet": starshina.pk, "usable_password": "true",
    })
    assert response.status_code == 200 and "guruh komandiri lavozimida emas" in response.content.decode()
    assert not User.objects.filter(username="notogri").exists()
    body = client.get(reverse("admin:naryad_user_add")).content.decode()
    assert 'value="commander"' in body  # rol tanlash


def test_accounts_command_and_admin_actions_work_for_commanders(client, admin, tmp_path):
    from naryad.accounts import create_cadet_accounts

    report = create_cadet_accounts(tmp_path)
    commanders = User.objects.filter(role="commander")
    assert commanders.count() == 6 and {u.cadet.group.code for u in commanders} == set(GROUPS)
    rows = list(load_workbook(report.file).active.iter_rows(min_row=5, values_only=True))
    assert sum(1 for r in rows if r[5] == "Guruh komandiri") == 6
    # Parolni tiklash (boshqaruv panelidagi amal) — komandirga ham
    admin.is_superuser = True
    admin.save()
    client.force_login(admin)
    boss = commanders.get(cadet__group__code="BI-225")
    response = client.post(reverse("admin:naryad_user_changelist"), {"action": "reset_temp_passwords", "_selected_action": [boss.pk]})
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    row = next(r for r in load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=5, values_only=True) if r[0])
    boss.refresh_from_db()
    assert row[3] == boss.username and boss.check_password(row[4]) and boss.must_change_password
    # Boshqaruv panelidagi “Akkaunt yaratish (Excel)” — akkaunti yo'q kursantlar uchun (komandirga rol avtomatik)
    boss.delete()
    target = commander_cadet("BI-225")
    response = client.post(reverse("admin:naryad_cadet_changelist"), {"action": "create_accounts", "_selected_action": [target.pk]})
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    boss = User.objects.get(cadet=target)
    assert boss.role == "commander" and boss.must_change_password
    row = next(r for r in load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=5, values_only=True) if r[0])
    assert row[3] == boss.username and row[5] == "Guruh komandiri" and boss.check_password(row[4])
    # Majburiy parol almashtirish komandirga ham amal qiladi
    client.force_login(boss)
    assert client.get(reverse("group_overview"))["Location"] == reverse("password_change")


# ---------------------------------------------------------------------------
# 2. Mening guruhim — faqat o'z guruhi
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("group", GROUPS)
def test_commander_sees_only_own_group(client, admin, group):
    generate_and_approve(SUNDAY, admin)
    client.force_login(make_commander(group))
    html = page(client, "group_overview")
    own = Cadet.objects.filter(group__code=group, is_active=True)
    for c in own:
        assert c.full_name in html
    for c in Cadet.objects.exclude(group__code=group)[:40]:
        assert f'data-testid="row-{c.pk}"' not in html
        assert client.get(reverse("group_cadet", args=[c.pk])).status_code == 404  # IDOR
    assert client.get(reverse("group_cadet", args=[999999])).status_code == 404  # mavjud emasligi bilan bir xil
    assert client.get(reverse("group_cadet", args=[own.exclude(duty_exempt=True).first().pk])).status_code == 200


def test_group_parameter_in_url_is_ignored(client, admin):
    client.force_login(make_commander("BI-225"))
    html = page(client, "group_overview", guruh="BI-525", group="BI-525")
    assert "BI-225" in html and all(c.full_name not in html for c in Cadet.objects.filter(group__code="BI-525"))


def test_next_last_on_duty_now_and_sick_dates(client, admin, monkeypatch):
    sunday = generate_and_approve(SUNDAY, admin)
    group = INITIAL_ROTATION["noj1"]
    on_duty = sunday.assignments.filter(group__code=group).select_related("cadet").first()
    monkeypatch.setattr("django.utils.timezone.now", lambda: at(SUNDAY - dt.timedelta(days=1)))  # shanba: yakshanba — navbatdagi
    generate_and_approve(MONDAY, admin)
    on_monday = DutyDay.objects.get(date=MONDAY).assignments.filter(group__code=group).select_related("cadet", "post_slot__post").first()
    sick = Cadet.objects.filter(group__code=group, duty_exempt=False, gender=Gender.MALE).exclude(
        pk__in=Assignment.objects.values("cadet_id")).first()
    Absence.objects.create(cadet=sick, start_date=SUNDAY, end_date=SUNDAY + dt.timedelta(days=5), reason_type="kasal", illness=ILLNESS)
    commander = make_commander(group)
    client.force_login(commander)

    html = page(client, "group_overview")  # shanba: hali hech kim naryadda emas, yakshanba — navbatdagi
    assert f"{SUNDAY:%d.%m.%Y}" in html and 'data-testid="on-duty-now"' not in html

    monkeypatch.setattr("django.utils.timezone.now", lambda: at(MONDAY, 9))  # dushanba ertalab — yakshanba sutkasi
    client.force_login(commander)  # sessiya muddati (12 soat) — vaqt sakrashida qayta kirish
    html = page(client, "group_overview")
    row = html[html.index(f'data-testid="row-{on_duty.cadet_id}"'):]
    row = row[:row.index("</a>")]
    assert 'data-testid="on-duty-now"' in row and "Hozir naryadda" in row
    count = Assignment.objects.filter(duty_day__date=SUNDAY, group__code=group).count()
    assert f'data-testid="on-duty-count">{count} kishi' in html and 'data-testid="sick-count">1 kishi' in html
    assert "Kasal: 13.09–18.09" in html and ILLNESS not in html
    nxt = html[html.index(f'data-testid="row-{on_monday.cadet_id}"'):]
    nxt = nxt[:nxt.index("</a>")]
    assert f"{MONDAY:%d.%m.%Y}" in nxt  # navbatdagi tasdiqlangan naryadi

    monkeypatch.setattr("django.utils.timezone.now", lambda: at(TUESDAY + dt.timedelta(days=1)))  # chorshanba
    client.force_login(commander)
    html = page(client, "group_overview")
    row = html[html.index(f'data-testid="row-{on_duty.cadet_id}"'):]
    row = row[:row.index("</a>")]
    assert "Yaqin kunlarda naryadi yo‘q" in row and f"{SUNDAY:%d.%m.%Y}" in row  # oxirgi naryadi


def test_draft_never_visible_to_commander(client, admin):
    draft = services.generate_draft(SUNDAY, admin)
    group = INITIAL_ROTATION["noj1"]
    in_draft = draft.assignments.filter(group__code=group).select_related("cadet").first()
    client.force_login(make_commander(group))
    html = page(client, "group_overview")
    assert 'data-testid="on-duty-now"' not in html and f"{SUNDAY:%d.%m.%Y}" not in html
    html = page(client, "group_cadet", in_draft.cadet_id)
    assert f"{SUNDAY:%d.%m.%Y}" not in html
    week = page(client, "group_week")
    assert 'bg-brand-50 font-semibold' not in week  # hech bir katak to'ldirilmagan
    assert client.get(reverse("day_detail", args=[SUNDAY.isoformat()])).status_code == 403


def test_weekly_view_current_and_next_week(client, admin):
    generate_and_approve(SUNDAY, admin)
    generate_and_approve(MONDAY, admin)
    group = INITIAL_ROTATION["markaz"]
    client.force_login(make_commander(group))
    html = page(client, "group_week")
    assert 'data-testid="week-1"' in html and 'data-testid="week-2"' in html
    assert "13-sentabr" in html and "20-sentabr" in html
    filled = html.count("bg-brand-50 font-semibold")
    expected = Assignment.objects.filter(group__code=group, duty_day__status="approved").count()
    assert filled == expected


def test_hidden_information_not_in_commander_html(client, admin):
    from naryad import documents as doc_service

    generate_and_approve(SUNDAY, admin)
    group = "BI-325"
    boys = list(Cadet.objects.filter(group__code=group, duty_exempt=False, gender=Gender.MALE))
    absence = Absence.objects.create(cadet=boys[0], start_date=SUNDAY, end_date=TUESDAY, reason_type="kasal", illness=ILLNESS)
    full, thumb, w, h = doc_service.process_image(__import__("tests.test_documents", fromlist=["jpeg_bytes"]).jpeg_bytes())
    doc = SickDocument.objects.create(absence=absence, image=full, thumb=thumb, width=w, height=h, size=len(full))
    Absence.objects.create(cadet=boys[1], start_date=SUNDAY, end_date=TUESDAY, reason_type="tatil", note="Oilaviy sabab — to‘y")
    penalties.add(boys[2], "tartib", "Kechikib keldi — maxfiy izoh", SUNDAY, 2, admin)
    commander = make_commander(group)
    client.force_login(commander)
    pages = [page(client, "group_overview"), page(client, "group_week"), page(client, "home"),
             page(client, "group_cadet", boys[0].pk), page(client, "group_cadet", boys[2].pk), page(client, "group_cadet", commander.cadet_id)]
    for html in pages:
        # 20-bosqichdan o'z guruhi zalyoti ko'rinadi (zalyot izohi “Guruh zalyotlari” sahifasida), qolganlari — yo'q
        for secret in (ILLNESS, str(doc.pk), "Oilaviy sabab", "maxfiy izoh", "buyruq", commander.cadet.duty_position, "Hujjat"):
            assert secret not in html, secret
    assert "Kasal: 13.09–15.09" in pages[0]
    assert client.get(reverse("sick_doc", args=[absence.pk, doc.pk])).status_code == 403  # hujjat rasmi


def test_commander_menu_has_no_admin_sections_in_three_languages(client, admin):
    commander = make_commander("BI-425")
    expected = {"uz": ("Mening guruhim", ["Haftalar", "Kasallar ro‘yxati", "Zalyotchiklar", "Statistika", "Boshqaruv paneli", "Sababli holatlar", "Kursantlar<"]),
                "ru": ("Моя группа", ["Недели", "Список больных", "Залётчики", "Статистика", "Панель управления", "Уважительные причины"]),
                "en": ("My group", ["Weeks", "Sick list", "Offenders", "Statistics", "Admin panel", "Excused absences"])}
    for lang, (label, forbidden) in expected.items():
        commander.language = lang
        commander.save()
        client.force_login(commander)
        html = page(client, "home")
        nav = html[html.index("<nav"):html.index("</header>")]
        assert label in nav and reverse("group_overview") in nav, lang
        for word in forbidden:
            assert word not in nav, (lang, word)
        for name in ("group_overview", "group_week"):
            assert client.get(reverse(name)).status_code == 200


def test_commander_keeps_cadet_abilities(client, admin):
    generate_and_approve(SUNDAY, admin)
    commander = make_commander("BI-625")
    client.force_login(commander)
    assert client.get(reverse("approved_days")).status_code == 200
    assert client.get(reverse("day_detail", args=[SUNDAY.isoformat()])).status_code == 200
    pdf = client.get(reverse("day_pdf", args=[SUNDAY.isoformat()]))
    assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
    home = page(client, "home")
    assert 'data-testid="my-group"' in home and 'data-testid="my-duty"' not in home  # “Mening naryadim” o'rnida
    assert 'data-testid="duty-reason-notice"' in home  # sababli holat eslatmasi
    assert client.get(reverse("my_card"), follow=True).status_code == 200
    assert client.get(reverse("password_change")).status_code == 200


def test_commander_without_position_loses_group_access(client, admin):
    from naryad import cadets as cadet_service

    commander = make_commander("YuP-125")
    cadet_service.clear_exempt(commander.cadet, "lavozimdan bo‘shadi", admin)
    client.force_login(commander)
    assert client.get(reverse("group_overview")).status_code == 403
    assert client.get(reverse("approved_days")).status_code == 200  # kursant sifatida qoladi


@pytest.mark.parametrize("lang", ["uz", "ru", "en"])
def test_group_pages_open_in_three_languages(client, admin, lang):
    generate_and_approve(SUNDAY, admin)
    commander = make_commander("BI-225")
    commander.language = lang
    commander.save()
    client.force_login(commander)
    cadet = Cadet.objects.filter(group__code="BI-225", duty_exempt=False).first()
    for name, args in [("group_overview", []), ("group_week", []), ("group_status", []), ("group_cadet", [cadet.pk])]:
        assert client.get(reverse(name, args=args)).status_code == 200, (lang, name)
    title = {"uz": "Mening guruhim", "ru": "Моя группа", "en": "My group"}[lang]
    assert title in page(client, "group_overview")
