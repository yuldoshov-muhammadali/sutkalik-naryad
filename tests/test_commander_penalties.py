"""20-bosqich: komandir zalyot kiritadi, admin ko'rib chiqadi va tahrirlaydi, zalyot hisobi (CLAUDE.md 68-band)."""

from __future__ import annotations

import datetime as dt
import io
import re

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from naryad import days, penalties, services
from naryad.models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    Gender,
    InitialRotation,
    Penalty,
    PenaltyEntry,
    Post,
    Setting,
    SlotRole,
    StudyGroup,
    User,
)

from .conftest import INITIAL_ROTATION

SUNDAY = dt.date(2026, 9, 13)
MONDAY = SUNDAY + dt.timedelta(days=1)
PASSWORD = "Kuchli-parol-2026"
GROUP = INITIAL_ROTATION["markaz"]  # bu guruh 1-NO‘J / 1-Yotoqxonada emas — zalyotchisi jazo o'rnida turadi
pytestmark = pytest.mark.django_db


def at(day: dt.date, hour: int = 9) -> dt.datetime:
    return timezone.make_aware(dt.datetime.combine(day, dt.time(hour, 0)))


@pytest.fixture
def admin(db_cadets, monkeypatch):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=SUNDAY)
    Setting.load()
    monkeypatch.setattr("django.utils.timezone.now", lambda: at(MONDAY))  # bugun — 14.09
    return User.objects.create_user("kursboshligi", password=PASSWORD, role="admin", first_name="Xasan", last_name="Turobov")


def make_commander(group: str) -> User:
    cadet = Cadet.objects.get(group__code=group, duty_position=f"{group} guruh komandiri")
    user = User(username=f"komandir.{group.lower()}", role="commander", cadet=cadet, first_name=cadet.first_name, last_name=cadet.last_name)
    user.set_password(PASSWORD)
    user.full_clean()
    user.save()
    return user


@pytest.fixture
def commander(admin):
    return make_commander(GROUP)


def boy(group: str, index: int = 0) -> Cadet:
    return Cadet.objects.filter(group__code=group, gender=Gender.MALE, duty_exempt=False, is_active=True).order_by("list_number")[index]


def form_data(cadet, count=2, order_date=SUNDAY, reason="tartib", note="Kechikib keldi"):
    return {"cadet": cadet.pk, "reason_type": reason, "note": note, "order_date": order_date.isoformat(), "count": count}


def enter(client, cadet, **kw):
    """Komandir: forma → tasdiqlash → saqlash."""
    data = form_data(cadet, **kw)
    first = client.post(reverse("group_penalty_add"), data)
    return first, client.post(reverse("group_penalty_add"), {**data, "confirm": "1"})


# ---------------------------------------------------------------------------
# 1. Kiritish
# ---------------------------------------------------------------------------


def test_commander_enters_penalty_with_confirmation(client, commander):
    client.force_login(commander)
    target = boy(GROUP)
    body = client.post(reverse("group_penalty_add"), form_data(target)).content.decode()
    assert 'data-testid="penalty-confirm"' in body
    assert f"{target.full_name} ({GROUP}) ga 2 ta jazo naryadi kiritiladi. Davom etasizmi?" in body
    assert not Penalty.objects.exists()  # tasdiqlanmaguncha yaratilmaydi
    second = client.post(reverse("group_penalty_add"), {**form_data(target), "confirm": "1"})
    assert second.status_code == 302 and second["Location"] == reverse("group_penalties")
    penalty = Penalty.objects.get(cadet=target)
    assert penalty.is_active and penalty.count == 2 and penalty.created_by == commander and penalty.is_new
    entry = PenaltyEntry.objects.get(penalty=penalty)
    assert entry.by_commander and entry.created_by == commander
    log = ChangeLog.objects.get(action="penalty_add", cadet=target)
    assert log.user == commander and "guruh komandiri" in log.details


def test_only_own_group_in_form_and_other_group_id_404(client, commander):
    client.force_login(commander)
    html = client.get(reverse("group_penalty_add")).content.decode()
    options = re.findall(r'<option value="(\d+)"', html)
    assert options and all(Cadet.objects.get(pk=int(pk)).group.code == GROUP for pk in options)
    stranger = boy("BI-225" if GROUP != "BI-225" else "BI-325")
    assert str(stranger.pk) not in options
    for confirm in ("", "1"):
        response = client.post(reverse("group_penalty_add"), {**form_data(stranger), "confirm": confirm})
        assert response.status_code == 404
    assert not Penalty.objects.exists()
    # Service darajasida ham: boshqa guruh — rad
    from django.core.exceptions import PermissionDenied

    with pytest.raises(PermissionDenied):
        penalties.add(stranger, "tartib", "x", SUNDAY, 1, commander)


def test_girl_exempt_inactive_rejected(client, commander, admin):
    client.force_login(commander)
    girl = Cadet.objects.filter(group__code=GROUP, gender=Gender.FEMALE).first()
    inactive = boy(GROUP, 1)
    Cadet.objects.filter(pk=inactive.pk).update(is_active=False, deactivation_reason="sinov")
    inactive.refresh_from_db()
    for cadet in [c for c in (girl, commander.cadet, inactive) if c is not None]:
        response = client.post(reverse("group_penalty_add"), {**form_data(cadet), "confirm": "1"})
        # ro'yxatda yo'q: faolsiz — doiradan tashqari (404), qiz va lavozimdagi — forma xatosi
        assert response.status_code == 404 or 'data-testid="commander-penalty-form"' in response.content.decode(), cadet
        with pytest.raises(services.ServiceError):  # service darajasida ham — tushunarli xato
            penalties.add(cadet, "tartib", "x", SUNDAY, 1, commander)
    assert not Penalty.objects.exists()


@pytest.mark.parametrize("kw, message", [
    ({"count": 4}, "3"),
    ({"order_date": MONDAY + dt.timedelta(days=1)}, "Tanbeh sanasi kelajakda bo‘lmasin."),
    ({"order_date": MONDAY - dt.timedelta(days=31)}, "Tanbeh sanasi 30 kundan eski bo‘lmasin."),
])
def test_limits_for_commander(client, commander, kw, message):
    client.force_login(commander)
    _, response = enter(client, boy(GROUP), **kw)
    assert response.status_code == 200 and message in response.content.decode()
    assert not Penalty.objects.exists()


def test_limit_changes_in_settings_and_admin_keeps_1_to_10(client, commander, admin):
    Setting.objects.update(commander_penalty_max=5)
    client.force_login(commander)
    enter(client, boy(GROUP), count=5)
    assert Penalty.objects.get().count == 5
    penalties.add(boy("BI-325"), "tartib", "x", SUNDAY - dt.timedelta(days=60), 10, admin)  # admin: 1–10, eski sana ham
    assert Penalty.objects.get(cadet=boy("BI-325")).count == 10


def test_second_penalty_adds_count_with_separate_log(client, commander):
    client.force_login(commander)
    target = boy(GROUP)
    enter(client, target, count=2)
    enter(client, target, count=1, reason="uxlash", order_date=MONDAY)
    penalty = Penalty.objects.get(cadet=target)
    assert penalty.count == 3 and penalty.entries.count() == 2
    assert ChangeLog.objects.filter(action="penalty_add", cadet=target).count() == 2


def test_double_submit_creates_one_entry(client, commander):
    client.force_login(commander)
    target = boy(GROUP)
    enter(client, target)
    _, again = enter(client, target)
    assert again.status_code == 302
    assert PenaltyEntry.objects.count() == 1 and Penalty.objects.get().count == 2


def test_daily_limit(client, commander):
    Setting.objects.update(commander_daily_limit=2)
    client.force_login(commander)
    enter(client, boy(GROUP, 0), order_date=SUNDAY)
    enter(client, boy(GROUP, 1), order_date=SUNDAY)
    _, third = enter(client, boy(GROUP, 2), order_date=SUNDAY)
    assert "Bir kunda ko‘pi bilan 2 ta zalyot kiritasiz." in third.content.decode()
    assert Penalty.objects.count() == 2


def test_commander_penalty_used_by_generator_immediately(client, commander, admin):
    client.force_login(commander)
    target = boy(GROUP)
    enter(client, target, count=1)
    draft = services.generate_draft(SUNDAY, admin)
    a = draft.assignments.select_related("post_slot__post").get(cadet=target)
    assert a.kind == Assignment.Kind.PENALTY and a.post_slot.post.code in ("noj1", "yotoq1") and a.post_slot.role == SlotRole.GUARD


# ---------------------------------------------------------------------------
# 2. Admin: ko'rib chiqish, tahrir, bekor qilish
# ---------------------------------------------------------------------------


def test_commander_cannot_edit_cancel_review(client, commander):
    client.force_login(commander)
    enter(client, boy(GROUP))
    penalty = Penalty.objects.get()
    for name, args, data in [("penalty_edit", [penalty.pk], {"reason_type": "boshqa", "note": "x", "order_date": "2026-09-13", "count": 1}),
                             ("penalty_cancel", [penalty.pk], {"reason": "x"}), ("penalty_review", [penalty.pk], {}),
                             ("penalty_review_all", [], {}), ("penalty_list", [], form_data(boy(GROUP, 1)))]:
        assert client.post(reverse(name, args=args), data).status_code == 403, name
    assert client.get(reverse("penalty_ledger")).status_code == 403
    penalty.refresh_from_db()
    assert penalty.is_active and penalty.count == 2 and penalty.is_new


def test_banner_and_review(client, commander, admin):
    c2 = Client()
    c2.force_login(commander)
    enter(c2, boy(GROUP, 0))
    enter(c2, boy(GROUP, 1))
    client.force_login(admin)
    home = client.get(reverse("home")).content.decode()
    assert "Komandirlar kiritgan 2 ta yangi zalyot ko‘rib chiqilmagan" in home
    listing = client.get(reverse("penalty_list")).content.decode()
    assert listing.count('data-testid="new-badge"') == 2 and "(guruh komandiri)" in listing
    first = Penalty.objects.order_by("pk").first()
    client.post(reverse("penalty_review", args=[first.pk]))
    assert "Komandirlar kiritgan 1 ta yangi zalyot" in client.get(reverse("home")).content.decode()
    client.post(reverse("penalty_review_all"), {"next": reverse("home")})
    assert 'data-testid="new-penalties-banner"' not in client.get(reverse("home")).content.decode()
    assert Penalty.objects.filter(reviewed_at__isnull=True).count() == 0
    assert all(p.count == 2 and p.is_active for p in Penalty.objects.all())  # zalyot o'zgarmadi


def test_admin_edits_with_log_and_cannot_go_below_served(client, commander, admin):
    c2 = Client()
    c2.force_login(commander)
    target = boy(GROUP)
    enter(c2, target, count=3)
    penalty = Penalty.objects.get()
    draft = services.generate_draft(SUNDAY, admin)
    days.approve(SUNDAY, admin, draft.version)  # 1 ta jazo naryadi o'tildi
    client.force_login(admin)
    data = {"reason_type": "uxlash", "note": "Naryadda uxlab qoldi", "order_date": "2026-09-12", "count": 0}
    assert client.post(reverse("penalty_edit", args=[penalty.pk]), data).status_code == 200  # 0 — forma xatosi
    response = client.post(reverse("penalty_edit", args=[penalty.pk]), {**data, "count": 1})
    assert response.status_code == 302  # 1 = o'tilgan → bajarildi
    penalty.refresh_from_db()
    assert penalty.count == 1 and penalty.status == Penalty.Status.DONE and penalty.reason_type == "uxlash"
    log = ChangeLog.objects.get(action="penalty_edit")
    assert "soni: 3 → 1" in log.details and "tanbeh sanasi: 13.09.2026 → 12.09.2026" in log.details and "sabab:" in log.details
    # bekor qilingan/bajarilgan — tahrir yo'q; o'tilganidan kam — xato (service)
    with pytest.raises(services.ServiceError):
        penalties.edit(penalty, "tartib", "x", SUNDAY, 1, admin)


def test_admin_cancel_with_reason_logged(client, commander, admin):
    c2 = Client()
    c2.force_login(commander)
    enter(c2, boy(GROUP))
    penalty = Penalty.objects.get()
    client.force_login(admin)
    client.post(reverse("penalty_cancel", args=[penalty.pk]), {"reason": "Xato kiritilgan"})
    penalty.refresh_from_db()
    assert penalty.status == Penalty.Status.CANCELLED
    assert ChangeLog.objects.filter(action="penalty_cancel", reason="Xato kiritilgan", user=admin).exists()


# ---------------------------------------------------------------------------
# 3. Hisob
# ---------------------------------------------------------------------------


def test_ledger_counts_approval_reopen_period_and_excel(client, commander, admin):
    target = boy(GROUP)
    penalties.add(target, "tartib", "x", SUNDAY, 3, admin)
    c2 = Client()
    c2.force_login(make_commander("BI-325"))
    other = boy("BI-325")
    enter(c2, other, count=2)
    data = penalties.ledger()
    row = next(r for r in data["groups"] if r.key.code == GROUP)
    assert (len(row.offenders), row.added, row.served, row.remaining) == (1, 3, 0, 3)
    day = services.generate_draft(SUNDAY, admin)
    day = days.approve(SUNDAY, admin, day.version)
    row = next(r for r in penalties.ledger()["groups"] if r.key.code == GROUP)
    assert (row.served, row.remaining) == (1, 2)
    days.reopen(SUNDAY, "tekshiruv", admin, DutyDay.objects.get(date=SUNDAY).version)
    row = next(r for r in penalties.ledger()["groups"] if r.key.code == GROUP)
    assert (row.served, row.remaining) == (0, 3)  # tasdiq bekor — hisob qaytdi
    # davr filtri: tanbeh sanasi
    assert penalties.ledger(date_from=MONDAY)["totals"].added == 0
    assert penalties.ledger(date_to=SUNDAY)["totals"].added == 5
    users = {r.key.username: (r.entries, r.added) for r in penalties.ledger()["users"]}
    assert users == {"kursboshligi": (1, 3), "komandir.bi-325": (1, 2)}
    client.force_login(admin)
    page = client.get(reverse("penalty_ledger"), {"guruh": GROUP}).content.decode()
    assert f'data-testid="lg-{GROUP}"' in page and 'data-testid="lg-BI-325"' not in page and "Turobov" in page
    response = client.get(reverse("penalty_ledger_excel"))
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Guruhlar", "Kursantlar", "Kim kiritdi"]
    assert wb["Guruhlar"]["A2"].value == "Guruh" and wb["Guruhlar"]["C2"].value == "Kiritilgan jazo naryadlari"
    assert any(r[0] == target.full_name and r[3] == 3 for r in wb["Kursantlar"].iter_rows(min_row=3, values_only=True))


def test_commander_sees_only_own_group_penalties_and_ledger(client, commander, admin):
    penalties.add(boy(GROUP), "tartib", "Bizning guruh izohi", SUNDAY, 2, admin)
    penalties.add(boy("BI-325"), "tartib", "Boshqa guruh maxfiy izohi", SUNDAY, 3, admin)
    client.force_login(commander)
    html = client.get(reverse("group_penalties")).content.decode()
    assert "Bizning guruh izohi" in html and "Boshqa guruh maxfiy izohi" not in html
    assert 'data-testid="ledger-added">2<' in html and 'data-testid="ledger-remaining">2<' in html
    overview = client.get(reverse("group_overview")).content.decode()
    assert "Zalyotchi: 2 tadan 0 tasi o‘tildi" in overview


def test_cadet_does_not_see_who_entered_and_others_see_nothing(client, commander, admin, monkeypatch):
    c2 = Client()
    c2.force_login(commander)
    target = boy(GROUP)
    enter(c2, target, count=1)
    draft = services.generate_draft(SUNDAY, admin)
    days.approve(SUNDAY, admin, draft.version)
    target_user = User.objects.create_user("k1", password=PASSWORD, role="cadet", cadet=target)
    monkeypatch.setattr("django.utils.timezone.now", lambda: at(SUNDAY - dt.timedelta(days=1), 20))
    client.force_login(target_user)
    home = client.get(reverse("home")).content.decode()
    assert "Navbatdan tashqari naryad" in home
    assert commander.cadet.short_name not in home and commander.cadet.last_name not in home and "komandir" not in home.lower()
    other = User.objects.create_user("k2", password=PASSWORD, role="cadet", cadet=boy("BI-525" if GROUP != "BI-525" else "BI-425", 3))
    client.force_login(other)
    for name, args in [("home", []), ("day_detail", [SUNDAY.isoformat()]), ("approved_days", [])]:
        page = client.get(reverse(name, args=args)).content.decode()
        assert "Zalyot" not in page and "Kechikib keldi" not in page, name


@pytest.mark.parametrize("lang", ["uz", "ru", "en"])
def test_pages_open_in_three_languages(client, commander, admin, lang):
    penalties.add(boy(GROUP), "tartib", "x", SUNDAY, 1, admin)
    for user, names in ((commander, ["group_penalties", "group_penalty_add", "group_overview"]),
                        (admin, ["penalty_list", "penalty_ledger", "home"])):
        user.language = lang
        user.save()
        client.force_login(user)
        for name in names:
            assert client.get(reverse(name)).status_code == 200, (lang, name)
    client.force_login(admin)
    assert client.get(reverse("penalty_edit", args=[Penalty.objects.get().pk])).status_code == 200
    title = {"uz": "Zalyot hisobi", "ru": "Учёт взысканий", "en": "Reprimand records"}[lang]
    assert title in client.get(reverse("penalty_ledger")).content.decode()

