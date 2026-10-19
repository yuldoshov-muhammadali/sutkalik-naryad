"""20-bosqich: sinov ma'lumotlarini to'liq tozalash — hammasi 0 dan (CLAUDE.md 69-band)."""

from __future__ import annotations

import datetime as dt
import io
from io import StringIO

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from django.urls import reverse
from openpyxl import load_workbook

from naryad import cleanup, days, penalties, services
from naryad.models import (
    Absence,
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    GroupPostAssignment,
    InitialRotation,
    Penalty,
    PenaltyEntry,
    Post,
    PublishedDuty,
    Setting,
    StudyGroup,
    User,
)

from .conftest import INITIAL_ROTATION

SUNDAY = dt.date(2026, 9, 13)
PASSWORD = "Kuchli-parol-2026"
pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def backup_dir(settings, tmp_path):
    settings.BACKUP_DIR = tmp_path / "zaxira"
    return settings.BACKUP_DIR


@pytest.fixture
def world(db_cadets):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    Setting.load()
    admin = User.objects.create_user("kursboshligi", password=PASSWORD, role="admin")
    # “import qilingan” tarix: 06.09 (oldingi hafta), keyin platformada 2 sutka
    old_week = DutyWeek.objects.create(start_date=SUNDAY - dt.timedelta(days=7), created_by=admin)
    old_day = services.generate_draft(SUNDAY - dt.timedelta(days=7), admin)
    days.approve(old_day.date, admin, old_day.version)
    ChangeLog.objects.create(action=ChangeLog.Action.IMPORT, duty_day=old_day, user=admin, details="qog‘ozdan import")
    DutyWeek.objects.create(start_date=SUNDAY, created_by=admin)
    boy = Cadet.objects.filter(gender="M", duty_exempt=False, group__code=INITIAL_ROTATION["markaz"]).first()
    penalties.add(boy, "tartib", "sinov", SUNDAY, 2, admin)
    for day in (SUNDAY, SUNDAY + dt.timedelta(days=1)):
        draft = services.generate_draft(day, admin)
        days.approve(day, admin, draft.version)
    services.generate_draft(SUNDAY + dt.timedelta(days=2), admin)  # qoralama
    sick = Absence.objects.create(cadet=Cadet.objects.last(), start_date=SUNDAY, end_date=SUNDAY + dt.timedelta(days=9),
                                  reason_type="kasal", illness="Gripp")
    cadet_user = User.objects.create_user("kursant", password=PASSWORD, role="cadet", cadet=boy)
    return {"admin": admin, "old_week": old_week, "sick": sick, "cadet": cadet_user}


def test_preview_changes_nothing_and_counts(world):
    before = (DutyDay.objects.count(), Assignment.objects.count(), ChangeLog.objects.count())
    p = cleanup.preview()
    assert (DutyDay.objects.count(), Assignment.objects.count(), ChangeLog.objects.count()) == before
    assert p.weeks == 2 and p.days == 4 and p.days_approved == 3 and p.days_draft == 1
    assert p.days_imported == 1 and p.days_platform == 3
    assert p.assignments == Assignment.objects.count() and p.assignments_penalty >= 1
    assert p.penalties == 1 and p.penalty_entries == 1 and p.changelog == ChangeLog.objects.count()
    by_cadet = {c.pk: (before_n, after) for c, before_n, after in p.cadets}
    assert len(by_cadet) == Cadet.objects.count() and all(after == 0 for _, after in by_cadet.values())
    assert sum(n for n, _ in by_cadet.values()) == Assignment.objects.filter(duty_day__status="approved", kind="normal").count()


def test_clear_removes_history_keeps_people_and_settings(world, backup_dir):
    users, cadets, absences = User.objects.count(), Cadet.objects.count(), Absence.objects.count()
    rotation = InitialRotation.as_mapping()
    p = cleanup.preview()
    result = cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", p.token)
    assert result.backup.exists() and result.backup.parent == backup_dir and result.backup.stat().st_size > 0
    for model in (DutyWeek, DutyDay, Assignment, GroupPostAssignment, Penalty, PenaltyEntry, PublishedDuty):
        assert not model.objects.exists(), model.__name__
    log = list(ChangeLog.objects.all())
    assert len(log) == 1 and log[0].action == "test_data_cleared" and log[0].user == world["admin"]
    assert "1 tasi import qilingan" in log[0].details
    assert (User.objects.count(), Cadet.objects.count(), Absence.objects.count()) == (users, cadets, absences)
    assert InitialRotation.as_mapping() == rotation and Setting.objects.exists()
    assert Absence.objects.filter(pk=world["sick"].pk, illness="Gripp").exists()


def test_after_clear_everything_is_zero_and_first_day_from_initial_rotation(client, world):
    cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    client.force_login(world["admin"])
    assert client.get(reverse("stats")).status_code == 200
    for c in Cadet.objects.all()[:100]:
        assert c.assignments.count() == 0
    wb = load_workbook(io.BytesIO(client.get(reverse("stats_excel")).content))
    for ws in wb.worksheets:
        if ws.title.startswith("Guruh ") or ws.title == "Qizlar":
            header = [c.value for c in ws[1]]
            if "Jami" in header:
                col = header.index("Jami")
                assert all((r[col] or 0) == 0 for r in ws.iter_rows(min_row=2, values_only=True)), ws.title
    # Birinchi haqiqiy sutka: bo'sh tarix to'smaydi, guruh joylari boshlang'ich taqsimot bo'yicha
    start = SUNDAY + dt.timedelta(days=14)
    days.open_week(start, world["admin"], "")
    assert days.generation_blocker(start) is None
    first = services.generate_draft(start, world["admin"])
    mapping = {g.post.code: g.group.code for g in first.group_posts.select_related("post", "group")}
    assert mapping == INITIAL_ROTATION


@pytest.mark.parametrize("w1, w2", [("", ""), ("TOZALASH", ""), ("tozalash", "HAMMASI NOL"), ("TOZALASH", "HAMMASI")])
def test_wrong_words_delete_nothing(world, w1, w2):
    n = DutyDay.objects.count()
    with pytest.raises(cleanup.CleanupError):
        cleanup.clear(world["admin"], w1, w2, cleanup.preview().token)
    assert DutyDay.objects.count() == n


def test_backup_failure_stops_deletion(world, monkeypatch):
    n = DutyDay.objects.count()

    def broken(*a, **k):
        raise OSError("disk to‘la")

    monkeypatch.setattr("naryad.cleanup.call_command", broken)
    with pytest.raises(cleanup.CleanupError, match="Zaxira olinmadi"):
        cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    assert DutyDay.objects.count() == n and ChangeLog.objects.filter(action="test_data_cleared").count() == 0


def test_error_inside_rolls_back(world, monkeypatch):
    n, logs = DutyDay.objects.count(), ChangeLog.objects.count()
    original = ChangeLog.objects.create

    def fail(**kwargs):
        if kwargs.get("action") == ChangeLog.Action.TEST_DATA_CLEARED:
            raise RuntimeError("kutilmagan xato")
        return original(**kwargs)

    monkeypatch.setattr(ChangeLog.objects, "create", fail)
    with pytest.raises(RuntimeError):
        cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    assert DutyDay.objects.count() == n and ChangeLog.objects.count() == logs and Penalty.objects.exists()


def test_second_run_does_not_delete_again(world):
    token = cleanup.preview().token
    cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", token)
    with pytest.raises(cleanup.CleanupError):
        cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", token)  # eski iz
    with pytest.raises(cleanup.CleanupError, match="allaqachon tozalangan"):
        cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    assert ChangeLog.objects.count() == 1


def test_page_flow_and_permissions(client, world):
    commander_cadet = Cadet.objects.get(duty_position=f"{INITIAL_ROTATION['noj1']} guruh komandiri")
    commander = User.objects.create_user("komandir", password=PASSWORD, role="commander", cadet=commander_cadet)
    for user in (world["cadet"], commander):
        client.force_login(user)
        assert client.get(reverse("clear_test_data")).status_code == 403
        assert client.post(reverse("clear_test_data_run"), {"word_1": "TOZALASH", "word_2": "HAMMASI NOL",
                                                             "token": cleanup.preview().token}).status_code == 403
    assert DutyDay.objects.exists()
    client.force_login(world["admin"])
    page = client.get(reverse("clear_test_data")).content.decode()
    assert 'data-testid="pv-days">4<' in page and "Qaytarib bo‘lmaydi" in page and "HAMMASI NOL" in page
    token = cleanup.preview().token
    response = client.post(reverse("clear_test_data_run"), {"word_1": "TOZALASH", "word_2": "XATO", "token": token}, follow=True)
    assert "noto‘g‘ri" in response.content.decode() and DutyDay.objects.exists()
    response = client.post(reverse("clear_test_data_run"), {"word_1": "TOZALASH", "word_2": "HAMMASI NOL", "token": token}, follow=True)
    body = response.content.decode()
    assert "Sinov ma’lumotlari tozalandi" in body and 'data-testid="nothing-to-delete"' in body
    assert not DutyDay.objects.exists()
    link = body.split('data-testid="backup-link"')[0].rsplit('href="', 1)[1].split('"')[0]
    download = client.get(link)
    assert download.status_code == 200 and download["Cache-Control"] == "private, no-store"
    assert client.get(reverse("backup_download", args=["..%2F..%2Fsettings.py"])).status_code == 404  # faqat zaxira nomi
    assert client.get(reverse("backup_download", args=["zaxira_1.json.gz"])).status_code == 404


def test_command_dry_run_and_confirmed(world):
    out = StringIO()
    call_command("clear_test_data", "--dry-run", stdout=out)
    assert "hech narsa o‘zgarmadi" in out.getvalue() and "sutkalar: 4" in out.getvalue() and DutyDay.objects.exists()
    with pytest.raises(CommandError):
        call_command("clear_test_data", "--confirm", "TOZALASH", "--confirm2", "yo‘q", stdout=StringIO())
    assert DutyDay.objects.exists()
    out = StringIO()
    call_command("clear_test_data", "--confirm", "TOZALASH", "--confirm2", "HAMMASI NOL", stdout=out)
    assert "Tozalandi" in out.getvalue() and not DutyDay.objects.exists()


@pytest.mark.django_db(transaction=True)
def test_backup_restores(world):
    days_before = set(DutyDay.objects.values_list("date", flat=True))
    result = cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    assert not DutyDay.objects.exists()
    call_command("flush", interactive=False, verbosity=0)
    call_command("loaddata", str(result.backup), verbosity=0)
    assert set(DutyDay.objects.values_list("date", flat=True)) == days_before
    assert Penalty.objects.count() == 1 and Cadet.objects.count() == 100 and User.objects.filter(username="kursboshligi").exists()


def test_penalty_ledger_works_after_clear(client, world):
    cleanup.clear(world["admin"], "TOZALASH", "HAMMASI NOL", cleanup.preview().token)
    assert penalties.ledger()["totals"].added == 0
    client.force_login(world["admin"])
    assert client.get(reverse("penalty_ledger")).status_code == 200
    boy = Cadet.objects.filter(gender="M", duty_exempt=False).first()
    penalties.add(boy, "tartib", "yangi", SUNDAY, 1, world["admin"])
    assert penalties.ledger()["totals"].added == 1
