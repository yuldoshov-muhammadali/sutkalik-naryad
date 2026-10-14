"""4-bosqich: load_roster, import_history, boshlang'ich taqsimot."""

from __future__ import annotations

import datetime as dt
from io import StringIO
from pathlib import Path

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from naryad import services
from naryad.engine.config import ROTATION_CHAIN
from naryad.models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    GroupPostAssignment,
    InitialRotation,
    Post,
    StudyGroup,
)

from .conftest import INITIAL_ROTATION

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "data" / "tarix_namuna.xlsx"
SEPT_18 = dt.date(2026, 9, 18)

pytestmark = pytest.mark.django_db


def run(*args) -> str:
    out = StringIO()
    call_command(*args, stdout=out, stderr=out)
    return out.getvalue()


@pytest.fixture
def roster(db):
    run("load_roster")


def sample_rows() -> list[list[object]]:
    ws = load_workbook(SAMPLE).worksheets[0]
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]


def write_csv(path: Path, rows: list[list[object]]) -> Path:
    lines = ["sana,joy,lavozim,kursant,guruh"]
    for r in rows:
        day = r[0].date() if isinstance(r[0], dt.datetime) else r[0]
        lines.append(",".join([day.isoformat() if isinstance(day, dt.date) else str(day), *[str(x) for x in r[1:]]]))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


# ---------------------------------------------------------------------------
# load_roster
# ---------------------------------------------------------------------------


def test_load_roster_is_idempotent(db):
    first = run("load_roster")
    assert "100 ta yangi" in first and "87/13 ✓" in first
    second = run("load_roster")
    assert "0 ta yangi" in second and "100 ta o‘zgarmadi" in second
    assert StudyGroup.objects.count() == 6
    assert Cadet.objects.count() == 100
    assert Cadet.objects.filter(gender="M").count() == 87
    assert Cadet.objects.filter(gender="F").count() == 13


def test_load_roster_updates_changed_fields(roster):
    cadet = Cadet.objects.get(group__code="BI-525", list_number=13)
    assert cadet.short_name == "Sa’dullayev B"
    cadet.last_name = "Xato"
    cadet.short_name = "Xato B"
    cadet.save()
    out = run("load_roster")
    assert "1 ta yangilandi" in out
    cadet.refresh_from_db()
    assert (cadet.last_name, cadet.short_name) == ("Sa’dullayev", "Sa’dullayev B")
    assert Cadet.objects.count() == 100


def test_load_roster_warns_on_composition_mismatch(roster):
    cadet = Cadet.objects.filter(group__code="BI-525").first()
    cadet.is_active = False
    cadet.deactivation_reason = "test"
    cadet.save()
    out = run("load_roster")
    assert "Ogohlantirish: BI-525: kutilgan 17 o‘g‘il / 0 qiz, bazada 16 / 0" in out


# ---------------------------------------------------------------------------
# import_history
# ---------------------------------------------------------------------------


def test_sample_import_then_generator_continues_from_sept_18(roster):
    out = run("import_history", str(SAMPLE))
    assert "Import muvaffaqiyatli" in out

    day = DutyDay.objects.get(date=SEPT_18)
    assert day.status == DutyDay.Status.APPROVED and day.approved_at is not None
    assert DutyWeek.objects.get().start_date == dt.date(2026, 9, 13)
    assert day.assignments.count() == 20
    assert set(day.assignments.values_list("source", flat=True)) == {Assignment.Source.MANUAL}
    log = ChangeLog.objects.get(duty_day=day)
    assert log.action == ChangeLog.Action.IMPORT and "qog‘ozdan import" in log.reason
    sept18 = {g.post.code: g.group.code for g in day.group_posts.all()}
    assert sept18 == INITIAL_ROTATION  # namunadagi taqsimot

    nxt = services.generate_draft(SEPT_18 + dt.timedelta(days=1))
    cur = {g.post.code: g.group.code for g in nxt.group_posts.all()}
    for i, code in enumerate(ROTATION_CHAIN):
        assert cur[ROTATION_CHAIN[(i + 1) % 6]] == sept18[code]
    assert cur["noj2"] == "BI-525"  # 18-sentabrda 1-NO‘J da edi
    yesterday = set(day.assignments.values_list("cadet_id", flat=True))
    assert not yesterday & set(nxt.assignments.values_list("cadet_id", flat=True))


def test_dry_run_writes_nothing(roster):
    out = run("import_history", str(SAMPLE), "--dry-run")
    assert "--dry-run" in out and "13.09.2026" in out
    assert not DutyDay.objects.exists()
    assert not DutyWeek.objects.exists()
    assert not Assignment.objects.exists()
    # load_roster faqat kursant yozuvlarini (naryaddan ozodlar) jurnalga yozadi; sutka yozuvi yo'q
    assert not ChangeLog.objects.exclude(action__in=ChangeLog.CADET_ACTIONS).exists()


def test_unknown_name_is_reported_with_line_and_nothing_written(roster, tmp_path):
    rows = sample_rows()
    rows[15][3] = "Mehmonov Sh"  # namunadagi yozuv; rasmiy ro'yxatda "Mexmonov"
    path = write_csv(tmp_path / "tarix.csv", rows)
    with pytest.raises(CommandError):
        run("import_history", str(path))
    with pytest.raises(Exception) as exc:
        from naryad.importers import import_history

        import_history(path)
    assert "17-qator: kursant topilmadi: «Mehmonov Sh» (YuP-125 guruhida)" in str(exc.value)
    assert not DutyDay.objects.exists()


def test_ambiguous_name_without_group_is_reported(roster, tmp_path):
    from naryad.importers import ImportFailed, import_history

    rows = sample_rows()
    rows[3][4] = ""  # Aslonov A — BI-425 va BI-525 da bor
    path = write_csv(tmp_path / "tarix.csv", rows)
    with pytest.raises(ImportFailed) as exc:
        import_history(path)
    assert "5-qator: bir nechta kursant mos keldi: «Aslonov A»" in str(exc.value)


def test_day_with_19_rows_is_rejected(roster, tmp_path):
    from naryad.importers import ImportFailed, import_history

    rows = sample_rows()[:-1]  # oxirgi qizni olib tashlaymiz
    path = write_csv(tmp_path / "tarix.csv", rows)
    with pytest.raises(ImportFailed) as exc:
        import_history(path)
    text = str(exc.value)
    assert "18.09.2026: 19 ta qator bor, 20 ta bo‘lishi kerak" in text
    assert "Qizlar yotoqxonasi bo‘yicha: lavozimlar mos emas" in text
    assert not DutyDay.objects.exists()


def test_rule_violations_are_reported(roster, tmp_path):
    from naryad.importers import ImportFailed, import_history

    rows = sample_rows()
    rows[2][3], rows[2][4] = "G‘iyozov I", "BI-425"  # 1-NO‘J ga boshqa guruhdan
    rows[19][3], rows[19][4] = "Panjiyeva M", "BI-425"  # qiz ikki marta
    rows[0][1] = "Qizlar yotoqxonasi bo‘yicha"  # o'g'il qizlar yotoqxonasida
    path = write_csv(tmp_path / "tarix.csv", rows)
    with pytest.raises(ImportFailed) as exc:
        import_history(path)
    text = str(exc.value)
    assert "1-NO‘J bo‘yicha: bir nechta guruhdan kursant" in text
    assert "Panjiyeva M 18.09.2026 sutkasida ikkinchi marta yozilgan" in text
    assert "2-qator: G‘ayratov G‘ o‘g‘il — Qizlar yotoqxonasiga qo‘yib bo‘lmaydi" in text


def test_apostrophe_variants_are_matched(roster, tmp_path):
    rows = sample_rows()
    rows[0][3] = "G'ayratov G'"  # oddiy apostrof
    rows[1][3] = "Sa’dullayev Behro‘z Bahodir o'g'li"  # to'liq F.I.Sh., boshqa belgilar
    path = write_csv(tmp_path / "tarix.csv", rows)
    run("import_history", str(path))
    assert DutyDay.objects.get(date=SEPT_18).assignments.count() == 20


def test_existing_day_is_not_imported_twice(roster):
    run("import_history", str(SAMPLE))
    with pytest.raises(CommandError):
        run("import_history", str(SAMPLE))
    assert DutyDay.objects.count() == 1


# ---------------------------------------------------------------------------
# Boshlang'ich taqsimot
# ---------------------------------------------------------------------------


def test_generator_stops_without_history_and_initial_rotation(roster):
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    with pytest.raises(services.ServiceError, match="Boshlang‘ich taqsimot"):
        services.generate_draft(dt.date(2026, 9, 13))


def test_generator_uses_initial_rotation_from_settings(roster):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13))
    assert {g.post.code: g.group.code for g in GroupPostAssignment.objects.filter(duty_day=day)} == INITIAL_ROTATION


def test_incomplete_initial_rotation_is_ignored(roster):
    InitialRotation.objects.create(post=Post.objects.get(code="noj1"), group=StudyGroup.objects.get(code="BI-525"))
    assert InitialRotation.as_mapping() is None


def test_settings_defaults():
    from naryad.models import Setting

    s = Setting.load()
    assert s.recipient_position.replace("\n", " ") == "Bojxona instituti Navbatchilik qismi boshlig‘i, podpolkovnik"
    assert (s.recipient_name, s.course_head_rank, s.course_head_name, s.course_number) == ("X.Sattorovga", "mayor", "X.Turobov", 2)
