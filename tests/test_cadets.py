"""6-bosqich: kursantlar bazasi — ro'yxat, qo'shish/tahrirlash, faolsizlantirish, Excel import, kartochka."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from naryad import services
from naryad.engine.names import short_name
from naryad.models import Assignment, Cadet, ChangeLog, DutyWeek, StudyGroup, User

from .conftest import INITIAL_ROTATION

PASSWORD = "Kuchli-parol-2026"
pytestmark = pytest.mark.django_db


@pytest.fixture
def admin_client(client, db_cadets):
    client.force_login(User.objects.create_user("kursboshligi", password=PASSWORD, role="admin"))
    return client


@pytest.fixture
def cadet_user(db_cadets):
    cadet = Cadet.objects.get(last_name="Sa’dullayev")
    return User.objects.create_user("sadullayev.behroz", password=PASSWORD, role="cadet", cadet=cadet)


def form_data(**overrides):
    data = {
        "last_name": "Karimov",
        "first_name": "Shoxrux",
        "patronymic": "Ali o‘g‘li",
        "group": StudyGroup.objects.get(code="BI-525").pk,
        "gender": "M",
        "short_name": "",
    }
    data.update(overrides)
    return data


# ---------------------------------------------------------------------------
# Ruxsatlar
# ---------------------------------------------------------------------------


ADMIN_ONLY_URLS = [
    ("cadet_list", []),
    ("cadet_create", []),
    ("cadet_import", []),
    ("cadet_import_template", []),
    ("cadet_edit", [1]),
]


@pytest.mark.parametrize("name,args", ADMIN_ONLY_URLS)
def test_cadet_role_cannot_open_admin_pages(client, cadet_user, name, args):
    client.force_login(cadet_user)
    assert client.get(reverse(name, args=args)).status_code == 403


@pytest.mark.parametrize("name,args", ADMIN_ONLY_URLS)
def test_anonymous_redirected_to_login(client, db_cadets, name, args):
    resp = client.get(reverse(name, args=args))
    assert resp.status_code == 302 and resp["Location"].startswith(reverse("login"))


def test_cadet_cannot_edit_or_deactivate_by_post(client, cadet_user):
    client.force_login(cadet_user)
    pk = cadet_user.cadet_id
    assert client.post(reverse("cadet_edit", args=[pk]), form_data()).status_code == 403
    assert client.post(reverse("cadet_deactivate", args=[pk]), {"reason": "Chetlatilgan"}).status_code == 403
    assert Cadet.objects.get(pk=pk).is_active


def test_cadet_sees_only_own_card(client, cadet_user):
    client.force_login(cadet_user)
    assert client.get(reverse("cadet_detail", args=[cadet_user.cadet_id])).status_code == 200
    other = Cadet.objects.exclude(pk=cadet_user.cadet_id).first()
    assert client.get(reverse("cadet_detail", args=[other.pk])).status_code == 403
    resp = client.get(reverse("my_card"))
    assert resp["Location"] == reverse("cadet_detail", args=[cadet_user.cadet_id])
    body = client.get(reverse("cadet_detail", args=[cadet_user.cadet_id])).content.decode()
    assert "Tahrirlash" not in body and "Faolsizlantirish" not in body


# ---------------------------------------------------------------------------
# Ro'yxat
# ---------------------------------------------------------------------------


def test_list_shows_summary_and_groups(admin_client):
    body = admin_client.get(reverse("cadet_list")).content.decode()
    assert "87/13 ✓" in body
    assert "BI-525" in body and "Sa’dullayev Behro‘z Bahodir o‘g‘li" in body
    assert "md:hidden" in body and "md:block" in body  # telefon: kartochkalar, kompyuter: jadval


def test_htmx_search_returns_partial(admin_client):
    resp = admin_client.get(reverse("cadet_list"), {"q": "sa’dul"}, HTTP_HX_REQUEST="true")
    body = resp.content.decode()
    assert "<html" not in body
    assert "1 ta kursant" in body and "Sa’dullayev" in body


def test_search_is_by_last_name_only(admin_client):
    body = admin_client.get(reverse("cadet_list"), {"q": "asadbek"}, HTTP_HX_REQUEST="true").content.decode()
    assert "Hech narsa topilmadi" in body  # ism bo'yicha emas
    body = admin_client.get(reverse("cadet_list"), {"q": "abduham"}, HTTP_HX_REQUEST="true").content.decode()
    assert "1 ta kursant" in body


def test_search_normalizes_apostrophes(admin_client):
    body = admin_client.get(reverse("cadet_list"), {"q": "Sa'dullayev"}, HTTP_HX_REQUEST="true").content.decode()
    assert "1 ta kursant" in body


def test_filters_group_gender_status(admin_client):
    body = admin_client.get(reverse("cadet_list"), {"guruh": "BI-625", "jins": "F"}, HTTP_HX_REQUEST="true").content.decode()
    assert "4 ta kursant" in body
    c = Cadet.objects.get(last_name="Sa’dullayev")
    c.is_active, c.deactivation_reason = False, "test"
    c.save()
    body = admin_client.get(reverse("cadet_list"), {"holat": "faol-emas"}, HTTP_HX_REQUEST="true").content.decode()
    assert "1 ta kursant" in body and "Sa’dullayev" in body
    body = admin_client.get(reverse("cadet_list"), {"holat": "faol"}, HTTP_HX_REQUEST="true").content.decode()
    assert "99 ta kursant" in body


def test_list_counts_only_approved_duties(admin_client):
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13), initial_rotation=INITIAL_ROTATION)
    cadet = Assignment.objects.filter(duty_day=day).first().cadet
    url = reverse("cadet_list")
    q = {"q": cadet.last_name, "guruh": cadet.group.code}
    assert ">0<" in admin_client.get(url, q, HTTP_HX_REQUEST="true").content.decode().replace(" ", "")
    day.status, day.approved_at = "approved", day.generated_at
    day.save()
    assert ">1<" in admin_client.get(url, q, HTTP_HX_REQUEST="true").content.decode().replace(" ", "")


# ---------------------------------------------------------------------------
# Qisqa ism va apostroflar
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "last, first, expected",
    [
        ("Shavkatov", "Shaxzod", "Shavkatov Sh"),
        ("Botirov", "Chingiz", "Botirov Ch"),
        ("G'ayratov", "G'iyosiddin", "G‘ayratov G‘"),
        ("Tursunov", "O`tkir", "Tursunov O‘"),
        ("Saʼdullayev", "Behroʻz", "Sa’dullayev B"),
    ],
)
def test_short_name_rules(last, first, expected):
    assert short_name(last, first) == expected


def test_create_normalizes_apostrophes_and_computes_short_name(admin_client):
    resp = admin_client.post(
        reverse("cadet_create"), form_data(last_name="G'ayratov", first_name="Shoxrux", patronymic="Ma`ruf o'g'li")
    )
    cadet = Cadet.objects.get(last_name="G‘ayratov", group__code="BI-525")
    assert resp.status_code == 302 and resp["Location"] == reverse("cadet_detail", args=[cadet.pk])
    assert (cadet.short_name, cadet.patronymic) == ("G‘ayratov Sh", "Ma’ruf o‘g‘li")
    assert cadet.list_number == 18  # BI-525 dagi 17 tadan keyingi
    log = ChangeLog.objects.get(cadet=cadet)
    assert log.action == "cadet_create" and log.duty_day is None


def test_manual_short_name_is_kept(admin_client):
    admin_client.post(reverse("cadet_create"), form_data(short_name="Karimov Shx"))
    cadet = Cadet.objects.get(last_name="Karimov", first_name="Shoxrux")
    assert cadet.short_name == "Karimov Shx"
    data = form_data(first_name="Bobur", short_name="Karimov Shx")
    admin_client.post(reverse("cadet_edit", args=[cadet.pk]), data)
    cadet.refresh_from_db()
    assert cadet.short_name == "Karimov Shx"  # qo'lda kiritilgani saqlanadi


def test_auto_short_name_follows_name_change(admin_client):
    cadet = Cadet.objects.get(last_name="Maxsudov")
    data = form_data(
        last_name="Maxsudov", first_name="Chori", patronymic=cadet.patronymic, short_name=cadet.short_name
    )
    admin_client.post(reverse("cadet_edit", args=[cadet.pk]), data)
    cadet.refresh_from_db()
    assert cadet.short_name == "Maxsudov Ch"


def test_short_name_preview_endpoint(admin_client):
    resp = admin_client.get(reverse("cadet_short_name"), {"last_name": "o'rmonjonov", "first_name": "o'tkir"})
    assert resp.content.decode() == "o‘rmonjonov O‘"


# ---------------------------------------------------------------------------
# Dublikat, tarix bor kursantni tahrirlash
# ---------------------------------------------------------------------------


def test_duplicate_in_same_group_is_rejected(admin_client):
    existing = Cadet.objects.get(last_name="Sa’dullayev")
    data = form_data(last_name="sa'dullayev", first_name="behro'z", patronymic="Bahodir o'g'li")
    resp = admin_client.post(reverse("cadet_create"), data)
    assert resp.status_code == 200
    assert "allaqachon bor" in resp.content.decode()
    assert Cadet.objects.filter(last_name__iexact="Sa’dullayev").count() == 1
    # boshqa guruhda bir xil F.I.Sh. — ruxsat
    data["group"] = StudyGroup.objects.get(code="BI-425").pk
    assert admin_client.post(reverse("cadet_create"), data).status_code == 302
    assert existing.group.code == "BI-525"


def test_group_change_with_history_requires_confirmation_and_is_logged(admin_client):
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13), initial_rotation=INITIAL_ROTATION)
    a = Assignment.objects.filter(duty_day=day, post_slot__post__code="noj1").first()
    cadet = a.cadet
    new_group = StudyGroup.objects.exclude(pk=cadet.group_id).first()
    data = form_data(
        last_name=cadet.last_name, first_name=cadet.first_name, patronymic=cadet.patronymic,
        group=new_group.pk, gender=cadet.gender, short_name=cadet.short_name,
    )
    resp = admin_client.post(reverse("cadet_edit", args=[cadet.pk]), data)
    assert resp.status_code == 200 and "naryad tarixi bor" in resp.content.decode()
    cadet.refresh_from_db()
    assert cadet.group_id != new_group.pk

    data["confirm_history_change"] = "on"
    resp = admin_client.post(reverse("cadet_edit", args=[cadet.pk]), data)
    assert resp.status_code == 302
    cadet.refresh_from_db()
    assert cadet.group_id == new_group.pk
    a.refresh_from_db()
    assert a.group_id != new_group.pk  # tarix o'chmadi: eski guruh saqlanib qoldi
    log = ChangeLog.objects.get(cadet=cadet, action="cadet_update")
    assert f"→ {new_group.code}" in log.details and "tarix o‘chmaydi" in log.details


def test_edit_without_history_needs_no_confirmation(admin_client):
    cadet = Cadet.objects.get(last_name="Maxsudov")
    data = form_data(
        last_name=cadet.last_name, first_name=cadet.first_name, patronymic=cadet.patronymic,
        group=StudyGroup.objects.get(code="BI-425").pk, gender="M",
    )
    assert admin_client.post(reverse("cadet_edit", args=[cadet.pk]), data).status_code == 302
    cadet.refresh_from_db()
    assert cadet.group.code == "BI-425"


# ---------------------------------------------------------------------------
# Faolsizlantirish
# ---------------------------------------------------------------------------


def test_deactivation_requires_reason_and_excludes_from_generator(admin_client):
    noj1_boys = Cadet.objects.filter(group__code=INITIAL_ROTATION["noj1"], gender="M").order_by("pk")
    target = noj1_boys.first()
    account = User.objects.create_user("t", password=PASSWORD, role="cadet", cadet=target)

    admin_client.post(reverse("cadet_deactivate", args=[target.pk]), {"reason": "Boshqa", "note": ""})
    target.refresh_from_db()
    assert target.is_active  # "Boshqa" izohsiz — rad etildi

    admin_client.post(reverse("cadet_deactivate", args=[target.pk]), {"reason": "Akademik ta’til", "note": "buyruq 12"})
    target.refresh_from_db()
    account.refresh_from_db()
    assert not target.is_active and target.deactivation_reason == "Akademik ta’til (buyruq 12)"
    assert not account.is_active
    assert ChangeLog.objects.filter(cadet=target, action="cadet_deactivate", reason__startswith="Akademik").exists()

    # generator faol emas kursantni olmaydi — bir necha sutka davomida ham
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    from django.utils import timezone

    for i in range(7):
        day = services.generate_draft(dt.date(2026, 9, 13) + dt.timedelta(days=i), initial_rotation=INITIAL_ROTATION)
        assert target.pk not in set(day.assignments.values_list("cadet_id", flat=True))
        day.status, day.approved_at = "approved", timezone.now()
        day.save()

    admin_client.post(reverse("cadet_activate", args=[target.pk]))
    target.refresh_from_db()
    account.refresh_from_db()
    assert target.is_active and target.deactivation_reason == "" and account.is_active
    assert ChangeLog.objects.filter(cadet=target, action="cadet_activate").exists()


def test_deactivated_cadet_keeps_history(admin_client):
    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13), initial_rotation=INITIAL_ROTATION)
    cadet = Assignment.objects.filter(duty_day=day).first().cadet
    admin_client.post(reverse("cadet_deactivate", args=[cadet.pk]), {"reason": "Chetlatilgan"})
    assert Assignment.objects.filter(cadet=cadet).exists()
    assert admin_client.get(reverse("cadet_detail", args=[cadet.pk])).status_code == 200


# ---------------------------------------------------------------------------
# Excel import
# ---------------------------------------------------------------------------


def xlsx(rows, header=("familiya", "ism", "otasining ismi", "guruh", "jins")):
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("royxat.xlsx", buf.getvalue())


def snapshot():
    return Cadet.objects.count(), ChangeLog.objects.count(), list(Cadet.objects.values_list("group_id", "gender", "patronymic"))


def test_import_preview_does_not_write(admin_client):
    before = snapshot()
    rows = [
        ("Yangiyev", "Shoxrux", "Ali o'g'li", "BI-525", "o'g'il"),  # yangi
        ("Sa’dullayev", "Behro‘z", "Bahodir o‘g‘li", "BI-525", "o‘g‘il"),  # o'zgarmaydi
        ("Maxsudov", "Hosilbek", "Baxtiyor o‘g‘li", "BI-425", "o‘g‘il"),  # guruh o'zgaradi
    ]
    resp = admin_client.post(reverse("cadet_import"), {"file": xlsx(rows)})
    body = resp.content.decode()
    assert resp.status_code == 200
    assert "Tasdiqlash va bazaga yozish" in body
    assert "Guruh: BI-525 → BI-425" in body
    assert "Bazada bor, faylda yo‘q (98)" in body
    assert snapshot() == before  # hech narsa yozilmadi

    resp = admin_client.post(reverse("cadet_import"), {"action": "confirm"})
    assert resp.status_code == 302
    assert Cadet.objects.count() == before[0] + 1
    assert Cadet.objects.get(last_name="Yangiyev").short_name == "Yangiyev Sh"
    assert Cadet.objects.get(last_name="Maxsudov").group.code == "BI-425"
    assert ChangeLog.objects.filter(action="cadet_import").count() == 2
    assert Cadet.objects.filter(is_active=True).count() == 101  # faylda yo'qlar o'chirilmadi


def test_import_with_error_row_is_blocked_entirely(admin_client):
    before = snapshot()
    rows = [
        ("Yangiyev", "Shoxrux", "", "BI-525", "o‘g‘il"),
        ("Xatoyev", "", "", "BI-999", "bola"),  # 3-qator: 3 ta xato
        ("Yangiyev", "Shoxrux", "", "BI-525", "o‘g‘il"),  # 4-qator: takror
    ]
    body = admin_client.post(reverse("cadet_import"), {"file": xlsx(rows)}).content.decode()
    assert "3-qator: ism bo‘sh" in body and "3-qator: guruh topilmadi: «BI-999»" in body
    assert "3-qator: jins noto‘g‘ri: «bola»" in body
    assert "4-qator: fayl ichida takrorlangan (2-qator bilan bir xil)" in body
    assert "Tasdiqlash va bazaga yozish" not in body
    # tugmani chetlab o'tib tasdiqlashga urinish ham rad etiladi
    resp = admin_client.post(reverse("cadet_import"), {"action": "confirm"})
    assert resp.status_code == 200 and "import tasdiqlanmadi" in resp.content.decode()
    assert snapshot() == before


def test_import_rejects_wrong_header_and_type(admin_client):
    resp = admin_client.post(reverse("cadet_import"), {"file": xlsx([], header=("ism", "familiya"))}, follow=True)
    assert "ustunlar yetishmaydi" in resp.content.decode()
    bad = SimpleUploadedFile("royxat.csv", b"a,b")
    resp = admin_client.post(reverse("cadet_import"), {"file": bad}, follow=True)
    assert "Faqat .xlsx" in resp.content.decode()


def test_import_template_download(admin_client):
    resp = admin_client.get(reverse("cadet_import_template"))
    assert resp.status_code == 200 and resp["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(resp.content)).worksheets[0]
    assert [c.value for c in ws[1]] == ["familiya", "ism", "otasining ismi", "guruh", "jins"]


# ---------------------------------------------------------------------------
# Kartochka
# ---------------------------------------------------------------------------


def test_card_shows_history_and_summaries(admin_client):
    from django.utils import timezone

    DutyWeek.objects.create(start_date=dt.date(2026, 9, 13))
    day = services.generate_draft(dt.date(2026, 9, 13), initial_rotation=INITIAL_ROTATION)
    day.status, day.approved_at = "approved", timezone.now()
    day.save()
    a = Assignment.objects.filter(duty_day=day, post_slot__post__code="noj1", post_slot__position=1).get()
    body = admin_client.get(reverse("cadet_detail", args=[a.cadet_id])).content.decode()
    assert "13.09.2026" in body and "1-NO‘J bo‘yicha" in body and "Navbatchi" in body
    assert "Joylar bo‘yicha" in body and "Lavozimlar bo‘yicha" in body
    assert "Tahrirlash" in body and "O‘zgarishlar jurnali" in body
