"""9-bosqich: statistika, adolat ko'rsatkichi, guruhlar kesimi, qo'lda tuzatishlar nazorati, Excel."""

from __future__ import annotations

import datetime as dt
import io

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from naryad import days, services, stats
from naryad.models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    Gender,
    GroupPostAssignment,
    InitialRotation,
    Post,
    PostSlot,
    StudyGroup,
    User,
)

from .conftest import INITIAL_ROTATION

SUNDAY = dt.date(2026, 9, 13)
PASSWORD = "Kuchli-parol-2026"
pytestmark = pytest.mark.django_db


@pytest.fixture
def admin(db_cadets):
    return User.objects.create_user("kursboshligi", password=PASSWORD, role="admin")


@pytest.fixture
def admin_client(client, admin):
    client.force_login(admin)
    return client


def boys(group: str) -> list[Cadet]:
    """Naryadga chiqadigan o'g'il bolalar (naryaddan ozodlarsiz)."""
    return list(Cadet.objects.filter(group__code=group, gender=Gender.MALE, duty_exempt=False).order_by("list_number"))


def slot(post_code: str, position: int = 1) -> PostSlot:
    return PostSlot.objects.get(post__code=post_code, position=position)


def make_day(date: dt.date, placements, status=DutyDay.Status.APPROVED) -> DutyDay:
    """Qo'lda tuzilgan ma'lum tarix: placements = [(kursant, post_code, position), ...]."""
    week, _ = DutyWeek.objects.get_or_create(start_date=date - dt.timedelta(days=(date.weekday() + 1) % 7))
    approved_at = timezone.now() if status == DutyDay.Status.APPROVED else None
    day = DutyDay.objects.create(week=week, date=date, status=status, approved_at=approved_at)
    for cadet, post_code, position in placements:
        Assignment.objects.create(duty_day=day, post_slot=slot(post_code, position), cadet=cadet, group=cadet.group)
    return day


def row_of(rows, cadet):
    return next(r for r in rows if r.cadet.pk == cadet.pk)


# ---------------------------------------------------------------------------
# Sonlar, qoralamalar, davr filtri
# ---------------------------------------------------------------------------


def test_counts_from_known_history(db_cadets):
    a, b = boys("BI-225")[:2]
    make_day(SUNDAY, [(a, "markaz", 1), (b, "noj1", 1)])
    make_day(SUNDAY + dt.timedelta(days=2), [(a, "patrul", 2), (b, "noj1", 3)])
    make_day(SUNDAY + dt.timedelta(days=4), [(a, "noj1", 1), (b, "oquv", 1)])

    rows = stats.cadet_rows("BI-225", stats.Period())
    ra, rb = row_of(rows, a), row_of(rows, b)
    assert ra.by_post == {"markaz": 1, "patrul": 1, "noj1": 1}
    assert ra.by_role == {"navbatchi_yordamchisi": 1, "posbon": 1, "navbatchi": 1}
    assert ra.total == 3 and ra.last_date == SUNDAY + dt.timedelta(days=4)
    assert rb.by_post == {"noj1": 2, "oquv": 1} and rb.total == 3
    # boshqa guruh o'g'illari va qizlar bu jadvalga kirmaydi
    assert all(r.cadet.group.code == "BI-225" and r.cadet.gender == Gender.MALE for r in rows)
    assert sum(r.total for r in rows) == 6


def test_drafts_are_excluded(db_cadets):
    a = boys("BI-325")[0]
    make_day(SUNDAY, [(a, "markaz", 1)])
    make_day(SUNDAY + dt.timedelta(days=2), [(a, "patrul", 1)], status=DutyDay.Status.DRAFT)
    row = row_of(stats.cadet_rows("BI-325", stats.Period()), a)
    assert row.total == 1 and row.by_post == {"markaz": 1}
    assert row.last_date == SUNDAY


def test_period_filter_range_and_week(db_cadets):
    a = boys("BI-425")[0]
    make_day(SUNDAY, [(a, "markaz", 1)])
    make_day(SUNDAY + dt.timedelta(days=3), [(a, "patrul", 1)])
    make_day(SUNDAY + dt.timedelta(days=8), [(a, "oquv", 1)])  # keyingi 2-kurs haftasi

    def total(params):
        return row_of(stats.cadet_rows("BI-425", stats.period_from_params(params)), a).total

    assert total({}) == 3
    assert total({"davr": "range", "dan": "2026-09-14", "gacha": "2026-09-30"}) == 2
    assert total({"davr": "range", "dan": "2026-09-16", "gacha": "2026-09-13"}) == 2  # teskari sanalar almashtiriladi
    assert total({"davr": "range", "gacha": "2026-09-13"}) == 1
    assert total({"davr": "week", "hafta": "2026-09-13"}) == 2
    assert total({"davr": "week", "hafta": "2026-09-20"}) == 1
    assert total({"davr": "week", "hafta": "2030-01-06"}) == 3  # bunday hafta yo'q → butun tarix

    week = stats.period_from_params({"davr": "week", "hafta": "2026-09-20"})
    assert week.label == "2-kurs haftasi: 20.09.2026 – 26.09.2026"
    row = row_of(stats.cadet_rows("BI-425", week), a)
    assert row.this_week == 1


def test_girls_pool_counts_roles(db_cadets):
    girls = list(Cadet.objects.filter(gender=Gender.FEMALE).order_by("last_name"))
    assert len(girls) == 13
    make_day(SUNDAY, [(girls[0], "qizlar", 1), (girls[1], "qizlar", 2), (girls[2], "qizlar", 3)])
    make_day(SUNDAY + dt.timedelta(days=2), [(girls[1], "qizlar", 1)])
    rows = stats.cadet_rows(stats.GIRLS, stats.Period())
    assert len(rows) == 13
    r1 = row_of(rows, girls[1])
    assert r1.total == 2 and r1.by_role == {"posbon": 1, "navbatchi": 1}


def test_group_sections_count_posts_and_per_boy(db_cadets):
    week = DutyWeek.objects.create(start_date=SUNDAY)
    groups = {g.code: g for g in StudyGroup.objects.all()}
    for i, date in enumerate([SUNDAY, SUNDAY + dt.timedelta(days=1)]):
        status = DutyDay.Status.APPROVED if i == 0 else DutyDay.Status.DRAFT
        day = DutyDay.objects.create(
            week=week, date=date, status=status, approved_at=timezone.now() if i == 0 else None
        )
        for post_code, group_code in INITIAL_ROTATION.items():
            GroupPostAssignment.objects.create(duty_day=day, post=Post.objects.get(code=post_code), group=groups[group_code])
    sections = {s.group.code: s for s in stats.group_sections(stats.Period())}
    assert sections["BI-625"].by_post["markaz"] == 1  # qoralama sanalmaydi
    assert sections["BI-625"].days == 1
    assert sections["YuP-125"].by_post == {"markaz": 0, "noj1": 0, "noj2": 0, "yotoq1": 0, "oquv": 0, "patrul": 1}
    boys_count = Cadet.objects.filter(group__code="BI-225", gender=Gender.MALE, is_active=True, duty_exempt=False).count()
    a = boys("BI-225")[0]
    make_day(SUNDAY + dt.timedelta(days=3), [(a, "markaz", 1)])
    assert stats.group_sections(stats.Period())[1].per_boy == round(1 / boys_count, 2)


# ---------------------------------------------------------------------------
# Adolat
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "diff, status", [(0, "green"), (1, "green"), (2, "yellow"), (3, "red"), (7, "red"), (None, "none")]
)
def test_fairness_thresholds(diff, status):
    assert stats.fairness_status(diff) == status


def rows_with_totals(totals):
    cadets = Cadet.objects.filter(group__code="BI-525", gender=Gender.MALE, duty_exempt=False).order_by("list_number")[: len(totals)]
    return [stats.CadetRow(cadet=c, total=t) for c, t in zip(cadets, totals)]


@pytest.mark.parametrize(
    "totals, status, diff",
    [([3, 3, 4, 4], "green", 1), ([3, 4, 5], "yellow", 2), ([2, 4, 5], "red", 3), ([0, 0], "none", None)],
)
def test_fairness_from_rows(db_cadets, totals, status, diff):
    f = stats.fairness_from_rows("BI-525", "BI-525", rows_with_totals(totals))
    assert f.status == status and f.diff == diff


def test_fairness_mean_and_within_one(db_cadets):
    f = stats.fairness_from_rows("BI-525", "BI-525", rows_with_totals([1, 4, 4, 4, 7]))
    assert (f.min, f.max, f.mean, f.diff, f.status) == (1, 7, 4.0, 6, "red")
    assert f.within_one_pct == 60.0


def test_inactive_cadet_not_in_fairness(db_cadets):
    group = boys("BI-525")
    for c in group[1:]:
        make_day(SUNDAY + dt.timedelta(days=group.index(c) * 2), [(c, "markaz", 1)])
    lazy = group[0]
    assert stats.fairness_all(stats.Period())[4].status == "green"  # faol, lekin naryadsiz: farq 1
    lazy.is_active = False
    lazy.deactivation_reason = "boshqa"
    lazy.save()
    f = next(x for x in stats.fairness_all(stats.Period()) if x.pool == "BI-525")
    assert f.min == 1 and f.diff == 0 and lazy.pk not in [r.cadet.pk for r in stats.cadet_rows("BI-525", stats.Period())]


def test_consecutive_same_post_detected(db_cadets):
    a, b = boys("BI-225")[:2]
    make_day(SUNDAY, [(a, "markaz", 1), (b, "noj1", 1)])
    make_day(SUNDAY + dt.timedelta(days=3), [(a, "markaz", 1), (b, "noj2", 1)])
    make_day(SUNDAY + dt.timedelta(days=6), [(a, "patrul", 1), (b, "noj1", 2)])
    make_day(SUNDAY + dt.timedelta(days=9), [(a, "patrul", 3)], status=DutyDay.Status.DRAFT)  # qoralama — sanalmaydi
    f = next(x for x in stats.fairness_all(stats.Period()) if x.pool == "BI-225")
    assert [(x["cadet"].pk, x["post"], x["first"], x["second"]) for x in f.repeats] == [
        (a.pk, "markaz", SUNDAY, SUNDAY + dt.timedelta(days=3))
    ]
    # b: noj1 → noj2 → noj1 — ketma-ket emas
    girls = next(x for x in stats.fairness_all(stats.Period()) if x.pool == stats.GIRLS)
    assert not girls.repeats_applicable


# ---------------------------------------------------------------------------
# Qo'lda tuzatishlar
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "placed, mean, alert",
    [(2, 0.2, True), (3, 1.5, True), (2, 1.5, False), (1, 0.05, False), (0, 0, False), (4, 2.5, False)],
)
def test_manual_outlier_rule(placed, mean, alert):
    assert stats.is_manual_outlier(placed, mean) is alert


def log_replace(day: DutyDay, old, new, code, user):
    return ChangeLog.objects.create(
        duty_day=day, action=ChangeLog.Action.REPLACE, old_cadet=old, new_cadet=new,
        reason=days.reason_text(code, "izoh" if code == "boshqa" else ""), reason_code=code, user=user,
    )


def test_manual_control_counts_and_alert(db_cadets, admin):
    group = boys("BI-325")
    favourite = group[0]
    d1 = make_day(SUNDAY, [(favourite, "markaz", 1)])
    d2 = make_day(SUNDAY + dt.timedelta(days=2), [(favourite, "noj1", 1)])
    d3 = make_day(SUNDAY + dt.timedelta(days=4), [(group[5], "noj1", 1)], status=DutyDay.Status.DRAFT)
    log_replace(d1, group[1], favourite, "kasal", admin)
    log_replace(d2, group[2], favourite, "sababli", admin)
    log_replace(d3, group[3], favourite, "tatil", admin)  # qoralama — sanalmaydi
    other = User.objects.create_user("ikkinchi", password=PASSWORD, role="admin")
    log_replace(d2, group[4], group[6], "boshqa", other)

    m = stats.manual_control(stats.Period())
    assert m["total"] == 3
    assert dict(m["by_reason"]) == {"Kasal": 1, "Ta’til": 0, "Sababli": 1, "Boshqa": 1}
    assert dict(m["by_user"]) == {"kursboshligi": 2, "ikkinchi": 1}
    rows = {r.cadet.pk: r for r in m["rows"]}
    assert (rows[favourite.pk].placed, rows[favourite.pk].removed) == (2, 0)
    assert (rows[group[1].pk].placed, rows[group[1].pk].removed) == (0, 1)
    assert [r.cadet.pk for r in m["alerts"]] == [favourite.pk]

    # davr filtri: faqat birinchi sutka
    m1 = stats.manual_control(stats.period_from_params({"davr": "range", "dan": "2026-09-13", "gacha": "2026-09-13"}))
    assert m1["total"] == 1 and not m1["alerts"]


def test_replacements_lost_by_regeneration_are_not_counted(db_cadets, admin):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=SUNDAY)
    draft = services.generate_draft(SUNDAY, admin)
    s = slot("noj1", 1)
    candidate = next(c.cadet for c in days.replacement_candidates(draft, s) if not c.excluded)
    days.replace_cadet(SUNDAY, s.pk, candidate.pk, "sababli", "", admin, draft.version)
    draft.refresh_from_db()
    services.generate_draft(SUNDAY, admin, expected_version=draft.version)  # qayta tuzish
    draft.refresh_from_db()
    days.approve(SUNDAY, admin, draft.version)
    assert ChangeLog.objects.filter(action=ChangeLog.Action.REPLACE).count() == 1
    assert stats.manual_control(stats.Period())["total"] == 0

    # tasdiqlangan sutkada saqlangan tuzatish sanaladi
    days.reopen(SUNDAY, "Xato tuzatiladi", admin, DutyDay.objects.get(date=SUNDAY).version)
    day = DutyDay.objects.get(date=SUNDAY)
    candidate = next(c.cadet for c in days.replacement_candidates(day, s) if not c.excluded)
    days.replace_cadet(SUNDAY, s.pk, candidate.pk, "kasal", "", admin, day.version)
    days.approve(SUNDAY, admin, DutyDay.objects.get(date=SUNDAY).version)
    m = stats.manual_control(stats.Period())
    assert m["total"] == 1 and dict(m["by_reason"])["Kasal"] == 1


# ---------------------------------------------------------------------------
# Sahifalar, ruxsatlar, Excel
# ---------------------------------------------------------------------------


def test_stats_page_renders_for_admin(admin_client, db_cadets):
    a = boys("BI-225")[0]
    make_day(SUNDAY, [(a, "markaz", 1)])
    response = admin_client.get(reverse("stats"), {"guruh": "BI-225"})
    assert response.status_code == 200
    html = response.content.decode()
    assert "Adolat ko‘rsatkichi" in html and "Guruhlar kesimi" in html and "Qo‘lda tuzatishlar nazorati" in html
    assert a.short_name in html
    assert 'data-pool="qizlar"' in html and html.count('data-status="') >= 7
    assert "kursanti kam guruhda bir o‘g‘ilga ko‘proq naryad" in html
    assert "position: sticky" not in html  # Tailwind klasslari bilan
    assert "sticky left-0" in html and "sticky top-0" in html
    # noto'g'ri parametrlar sahifani buzmaydi
    assert admin_client.get(reverse("stats"), {"davr": "range", "dan": "xato", "guruh": "yoq"}).status_code == 200


def test_stats_menu_link_and_home_badges(admin_client, db_cadets):
    group = boys("BI-425")
    make_day(SUNDAY, [(group[0], "markaz", 1)])
    for i in range(3):
        make_day(SUNDAY + dt.timedelta(days=2 + 2 * i), [(group[0], "patrul", 1)])
    html = admin_client.get(reverse("home")).content.decode()
    assert f'href="{reverse("stats")}"' in html
    assert 'data-testid="home-fairness"' in html
    section = html.split('data-testid="home-fairness"')[1].split("</section>")[0]
    assert section.count('data-pool="') == 7
    bi425 = section.split('data-pool="BI-425"')[1].split("</a>")[0]
    assert 'data-status="red"' in bi425 and "farq 4" in bi425


def test_cadet_cannot_open_stats(client, db_cadets):
    cadet = boys("BI-225")[0]
    user = User.objects.create_user("kursant", password=PASSWORD, role="cadet", cadet=cadet)
    client.force_login(user)
    assert client.get(reverse("stats")).status_code == 403
    assert client.get(reverse("stats_excel")).status_code == 403
    assert reverse("stats") not in client.get(reverse("home")).content.decode()


def test_cadet_card_compares_with_group_mean_without_names(client, db_cadets):
    group = boys("BI-225")
    me = group[0]
    make_day(SUNDAY, [(me, "markaz", 1), (group[1], "noj1", 1), (group[2], "noj1", 2)])
    make_day(SUNDAY + dt.timedelta(days=2), [(me, "patrul", 1), (group[3], "noj1", 1)])
    user = User.objects.create_user("kursant", password=PASSWORD, role="cadet", cadet=me)
    client.force_login(user)
    response = client.get(reverse("cadet_detail", args=[me.pk]))
    assert response.status_code == 200
    html = response.content.decode()
    comparison = response.context["comparison"]
    members = len(group)
    assert comparison["own"] == 2 and comparison["members"] == members
    assert comparison["mean"] == round(5 / members, 1)
    assert 'data-testid="comparison"' in html and "Guruh o‘rtachasi" in html
    others = Cadet.objects.exclude(pk=me.pk)
    leaked = [c.last_name for c in others if c.last_name in html and c.last_name not in me.full_name]
    assert leaked == []


def test_excel_export_has_all_sheets_and_period(admin_client, db_cadets, admin):
    group = boys("BI-225")
    d = make_day(SUNDAY, [(group[0], "markaz", 1)])
    make_day(SUNDAY + dt.timedelta(days=3), [(group[0], "markaz", 1)])
    log_replace(d, group[1], group[0], "kasal", admin)
    response = admin_client.get(reverse("stats_excel"), {"davr": "week", "hafta": "2026-09-13"})
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert 'filename="statistika_hafta_2026-09-13.xlsx"' in response["Content-Disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == [
        "Guruh YuP-125", "Guruh BI-225", "Guruh BI-325", "Guruh BI-425", "Guruh BI-525", "Guruh BI-625",
        "Qizlar", "Adolat", "Ketma-ket bir joy", "Guruhlar kesimi", "Qo‘lda tuzatishlar", "Qo‘lda — kursantlar",
    ]
    for ws in wb.worksheets:
        assert ws["A1"].value == "Davr: 2-kurs haftasi: 13.09.2026 – 19.09.2026"
    ws = wb["Guruh BI-225"]
    header = [c.value for c in ws[3]]
    row = next(r for r in ws.iter_rows(min_row=4, values_only=True) if r[0] == group[0].full_name)
    assert row[header.index("Jami")] == 2 and row[header.index("Markaz bo‘yicha")] == 2
    assert wb["Ketma-ket bir joy"].max_row == 4
    assert wb["Qo‘lda — kursantlar"]["A4"].value == group[0].full_name
    assert len(list(wb["Qizlar"].iter_rows(min_row=4))) == 13


def test_excel_filename_variants(db_cadets):
    assert stats.excel_filename(stats.Period()) == "statistika_butun_tarix.xlsx"
    p = stats.period_from_params({"davr": "range", "dan": "2026-09-01", "gacha": "2026-09-30"})
    assert stats.excel_filename(p) == "statistika_2026-09-01_2026-09-30.xlsx"


def test_heat_colors_darker_for_more():
    light, dark = stats.heat(1, 10), stats.heat(10, 10)
    assert stats.heat(0, 10) == "" and "0.20" in light and "0.90" in dark
    # jami ustuni: min…max oralig'ida — 16 och, 18 to'q
    assert "0.12" in stats.heat(16, 18, 16) and "0.90" in stats.heat(18, 18, 16)
    assert "0.51" in stats.heat(5, 5, 5)


def test_girls_consecutive_duty_detected(db_cadets):
    g = list(Cadet.objects.filter(gender=Gender.FEMALE).order_by("last_name"))
    make_day(SUNDAY, [(g[0], "qizlar", 1), (g[1], "qizlar", 2)])
    make_day(SUNDAY + dt.timedelta(days=3), [(g[0], "qizlar", 1), (g[1], "qizlar", 1 + 1)])
    make_day(SUNDAY + dt.timedelta(days=6), [(g[0], "qizlar", 2), (g[1], "qizlar", 1)])
    make_day(SUNDAY + dt.timedelta(days=9), [(g[0], "qizlar", 1)], status=DutyDay.Status.DRAFT)
    girls = next(x for x in stats.fairness_all(stats.Period()) if x.pool == stats.GIRLS)
    # g[0]: navbatchi → navbatchi (1 holat) → posbon; g[1]: posbon → posbon → navbatchi (navbatchi emas)
    assert [(x["cadet"].pk, x["first"], x["second"]) for x in girls.duty_repeats] == [
        (g[0].pk, SUNDAY, SUNDAY + dt.timedelta(days=3))
    ]
    boys_pool = next(x for x in stats.fairness_all(stats.Period()) if x.pool == "BI-225")
    assert boys_pool.duty_repeats == []


def test_girls_duty_repeats_on_page_and_excel(admin_client, db_cadets):
    g = Cadet.objects.filter(gender=Gender.FEMALE).order_by("last_name").first()
    make_day(SUNDAY, [(g, "qizlar", 1)])
    make_day(SUNDAY + dt.timedelta(days=4), [(g, "qizlar", 1)])
    html = admin_client.get(reverse("stats")).content.decode()
    assert "Ketma-ket 2 marta navbatchi" in html and "Ketma-ket ikki naryadda navbatchi bo‘lgan: 1 ta holat" in html
    wb = load_workbook(io.BytesIO(admin_client.get(reverse("stats_excel")).content))
    header = [c.value for c in wb["Adolat"][3]]
    girls_row = next(r for r in wb["Adolat"].iter_rows(min_row=4, values_only=True) if r[0] == "Qizlar (13)")
    assert girls_row[header.index("Ketma-ket 2 marta navbatchi (qizlar)")] == 1
    assert any(r[2] == "Navbatchi (ketma-ket 2 marta)" for r in wb["Ketma-ket bir joy"].iter_rows(min_row=4, values_only=True))
