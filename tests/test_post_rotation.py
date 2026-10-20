"""21-bosqich: joylar navbati va og'ir joylar (CLAUDE.md 70-band).

QOIDA A — haftada bitta og'ir joy (1-NO‘J, 1-Yotoqxona); QOIDA B — kecha turgan kursantni generator hech qachon
qo'ymaydi (jazo naryadi ham), faqat admin sabab va tasdiq bilan. Sozlama o'chiq bo'lsa — natija bayt-bayt avvalgidek.
"""

from __future__ import annotations

import datetime as dt
import io
import re
from collections import Counter, defaultdict
from itertools import pairwise

import pytest
from django.urls import reverse
from django.utils import translation
from openpyxl import load_workbook

from naryad import days, services
from naryad.engine import generate_day, pick_cadets
from naryad.engine.allocator import EXCLUDED_YESTERDAY, History, current_cycle, score_candidate, week_start
from naryad.engine.config import CYCLE_POSTS, PostRotation
from naryad.engine.models import KIND_PENALTY, MALE, ROLE_DUTY, ROLE_GUARD, Cadet, Record
from naryad.models import (
    Assignment,
    ChangeLog,
    DutyDay,
    DutyWeek,
    InitialRotation,
    Post,
    PostSlot,
    PrepNote,
    Setting,
    StudyGroup,
    User,
)
from naryad.models import Cadet as CadetModel
from naryad.services import ServiceError

from .conftest import INITIAL_ROTATION, exempt_words, load_roster, simulate

SUNDAY = dt.date(2026, 9, 13)
TUESDAY = SUNDAY + dt.timedelta(days=2)
PASSWORD = "Kuchli-parol-2026"
ROT = PostRotation()
HARD = ("noj1", "yotoq1")


# ---------------------------------------------------------------------------
# Algoritm (toza Python)
# ---------------------------------------------------------------------------


def cadet(i: int, group: str = "G1") -> Cadet:
    return Cadet(id=i, last_name=f"Kursant{i}", first_name="A", group=group, gender=MALE)


def rec(day: dt.date, cid: int, post: str, role: str = ROLE_GUARD, kind: str = "normal") -> Record:
    return Record(day=day, cadet_id=cid, post=post, role=role, group="G1", kind=kind)


def run(rotation, weeks: int = 12):
    """Haqiqiy guruh hajmlari (11/12/12/14/16/13), haftalar orasida 14 kun — scripts/simulate_weeks.py --joylar kabi."""
    cadets = load_roster()
    history, results = [], []
    for w in range(weeks):
        start = SUNDAY + dt.timedelta(days=14 * w)
        for i in range(7):
            res = generate_day(start + dt.timedelta(days=i), cadets, history, initial_rotation=INITIAL_ROTATION, rotation=rotation)
            results.append(res)
            history.extend(res.to_records())
    return cadets, results, history


@pytest.fixture(scope="module")
def sim_new():
    return run(ROT)


@pytest.fixture(scope="module")
def sim_old():
    return run(None)


def test_off_is_byte_identical_and_girls_unaffected(sim_old, sim_new):
    """rotation=None — eski kalit (oltin xesh tests/test_penalties.py da); qizlar ikkala holatda bir xil."""
    _, old, _ = sim_old
    _, new, _ = sim_new
    for a, b in zip(old, new, strict=True):
        assert [x for x in a.assignments if x.slot.post == "qizlar"] == [x for x in b.assignments if x.slot.post == "qizlar"]
    baseline, _ = simulate(load_roster(), SUNDAY, 7)
    for a, b in zip(baseline, old[:7], strict=True):
        assert a.assignments == b.assignments and a.warnings == b.warnings


@pytest.mark.parametrize("first, second", [("noj1", "yotoq1"), ("yotoq1", "noj1"), ("noj1", "noj1"), ("yotoq1", "yotoq1")])
def test_rule_a_one_hard_post_per_week_both_directions(first, second):
    pool = [cadet(i) for i in range(1, 8)]
    records = [rec(SUNDAY, 1, first)] + [rec(SUNDAY, i, "oquv") for i in range(2, 8)]
    hist = History(records, before=TUESDAY)
    pick = pick_cadets(pool, second, 3, TUESDAY, hist, rotation=ROT)
    assert 1 not in pick.chosen and len(pick.chosen) == 3
    assert pick.ranking[-1].cadet_id == 1 and pick.ranking[-1].hard_conflict
    # o'chiq bo'lsa — eski tartib: QOIDA A yo'q (haftadagi soni teng, kursant 1 orqaga surilmaydi)
    old = pick_cadets(pool, second, 3, TUESDAY, hist)
    assert not any(s.hard_conflict for s in old.ranking)


@pytest.mark.parametrize("light", ["markaz", "noj2", "oquv", "patrul"])
def test_rule_a_light_posts_still_allowed(light):
    pool = [cadet(i) for i in range(1, 6)]
    # kursant 1 shu hafta 1-NO‘J da turgan, lekin jami eng kam — yengil joyga birinchi bo'lib tanlanadi
    older = SUNDAY - dt.timedelta(days=14)
    records = [rec(SUNDAY, 1, "noj1")] + [r for i in range(2, 6) for r in (rec(SUNDAY, i, "noj2" if light != "noj2" else "oquv"),
                                                                          rec(older, i, "markaz" if light != "markaz" else "patrul"))]
    pick = pick_cadets(pool, light, 1, TUESDAY, History(records, before=TUESDAY), rotation=ROT)
    assert pick.chosen == (1,) and not pick.ranking[0].hard_conflict


def test_rule_a_relaxed_only_when_group_lacks_people_with_warning():
    """Hammasi shu hafta og'ir joyda turgan — o'rin bo'sh qolmaydi (QOIDA A yumshatiladi)."""
    pool = [cadet(i) for i in range(1, 5)]
    records = [rec(SUNDAY, i, "noj1") for i in range(1, 5)]
    pick = pick_cadets(pool, "yotoq1", 3, TUESDAY, History(records, before=TUESDAY), rotation=ROT)
    assert len(pick.chosen) == 3 and all(s.hard_conflict for s in pick.ranking)


def test_rule_a_all_groups_violations_only_when_forced_with_warning(sim_new):
    _, results, history = sim_new
    forced = {(r.day, s.cadet_id) for r in results for code, p in r.picks.items() if code in HARD
              for s in p.ranking[: len(p.chosen)] if s.hard_conflict}
    warned = sum(1 for r in results for w in r.warnings if "ikkinchi marta og‘ir joyga" in w)
    assert warned == len(forced)
    weekly = defaultdict(list)
    for r in history:
        if r.post in HARD:
            weekly[(r.cadet_id, week_start(r.day))].append(r.day)
    second = {(d, cid) for (cid, _), ds in weekly.items() for d in sorted(ds)[1:]}
    assert second == forced  # har bir ikkinchi og'ir joy — majburiy va ogohlantirilgan
    groups = {c.id: c.group for c in load_roster()}
    by_group = Counter(groups[cid] for _, cid in forced)
    assert set(by_group) <= {"YuP-125", "BI-225", "BI-325"}  # katta guruhlarda (14–16 kishi) buzilish yo'q


def test_rule_b_generator_never_assigns_yesterday_including_penalty():
    pool = [cadet(i) for i in range(1, 6)]
    yesterday = TUESDAY - dt.timedelta(days=1)
    records = [rec(yesterday, 1, "noj1", kind=KIND_PENALTY), rec(yesterday, 2, "oquv")]
    hist = History(records, before=TUESDAY)
    for rotation in (ROT, None):
        pick = pick_cadets(pool, "patrul", 4, TUESDAY, hist, rotation=rotation)
        assert pick.excluded == {1: EXCLUDED_YESTERDAY, 2: EXCLUDED_YESTERDAY} and len(pick.chosen) == 3
    # jazo naryadi QOIDA A, tsikl va og'ir joylar soniga kirmaydi
    s = score_candidate(pool[0], "yotoq1", TUESDAY, History([rec(SUNDAY, 1, "noj1", kind=KIND_PENALTY)], before=TUESDAY), ROT)
    assert (s.week_hard, s.hard_count, s.cycle, s.hard_conflict) == (0, 0, (), False)


def test_nekmurodov_case_same_post_any_role_goes_back():
    """Nekmurodov F: oxirgi marta 1-NO‘J navbatchisi edi — keyingi safar 1-NO‘J yordamchisi bo'ldi (eski tartib)."""
    roster = load_roster()
    nek = next(c for c in roster if c.short_name == "Nekmurodov F")
    group = [c for c in roster if c.group == nek.group and c.gender == MALE]
    records = [rec(SUNDAY - dt.timedelta(days=14), nek.id, "noj1", role=ROLE_DUTY)]
    records += [rec(SUNDAY, c.id, "oquv") for c in group if c.id != nek.id]  # boshqalar shu hafta bir martadan turgan
    hist = History(records, before=TUESDAY)
    old = pick_cadets(group, "noj1", 4, TUESDAY, hist)
    new = pick_cadets(group, "noj1", 4, TUESDAY, hist, rotation=ROT)
    assert nek.id in old.chosen  # eski tartib: haftadagi soni birinchi — yana 1-NO‘J
    assert nek.id not in new.chosen and next(s for s in new.ranking if s.cadet_id == nek.id).last_is_same_post


def test_no_third_time_unless_unavoidable(sim_new):
    _, _, history = sim_new
    weekly = Counter((r.cadet_id, week_start(r.day)) for r in history)
    assert max(weekly.values()) <= 2
    pool = [cadet(i) for i in range(1, 4)]
    records = [rec(SUNDAY, 1, "oquv"), rec(SUNDAY + dt.timedelta(days=1), 1, "markaz")]  # kursant 1 — 2 marta
    pick = pick_cadets(pool, "patrul", 3, SUNDAY + dt.timedelta(days=3), History(records, before=SUNDAY + dt.timedelta(days=3)), rotation=ROT)
    assert pick.chosen[-1] == 1 and pick.ranking[-1].third_time  # boshqa iloj yo'q — faqat oxirida


def test_cycle_avoids_repeat_until_all_five_and_restarts():
    old_day = SUNDAY - dt.timedelta(days=60)
    a = [rec(old_day + dt.timedelta(days=2 * k), 1, p) for k, p in enumerate(CYCLE_POSTS)]  # 5 ta joyning hammasi
    assert current_cycle(a, CYCLE_POSTS) == ()  # tsikl qaytadan boshlandi
    b = [rec(old_day - dt.timedelta(days=5), 2, "noj2")] + [rec(old_day - dt.timedelta(days=4 - k), 2, "markaz") for k in range(4)]
    assert current_cycle(b, CYCLE_POSTS) == ("noj2",)
    hist = History(a + b, before=TUESDAY)
    pool = [cadet(1), cadet(2)]
    old = pick_cadets(pool, "noj2", 1, TUESDAY, hist)
    new = pick_cadets(pool, "noj2", 1, TUESDAY, hist, rotation=ROT)
    assert old.chosen == (2,)  # eski: shu joyda soni teng, uzoq dam olgan oldin
    assert new.chosen == (1,) and new.ranking[1].cycle_repeat  # yangi: 2-kursant tsiklida 2-NO‘J bor


def test_hard_count_balance_on_hard_posts(sim_new, sim_old):
    old_day = SUNDAY - dt.timedelta(days=40)
    a = [rec(old_day, 1, "yotoq1"), rec(old_day + dt.timedelta(days=3), 1, "yotoq1"), rec(old_day + dt.timedelta(days=6), 1, "oquv")]
    b = [rec(old_day + dt.timedelta(days=1), 2, "noj2"), rec(old_day + dt.timedelta(days=4), 2, "oquv"),
         rec(old_day + dt.timedelta(days=9), 2, "oquv")]
    hist = History(a + b, before=TUESDAY)
    new = pick_cadets([cadet(1), cadet(2)], "noj1", 1, TUESDAY, hist, rotation=ROT)
    assert new.chosen == (2,) and [s.hard_count for s in new.ranking] == [0, 2]
    # 12 hafta: og'ir joylar farqi guruh ichida ≤ 2 (eski tartibdan yomon emas)
    cadets, _, history = sim_new
    for g in {c.group for c in cadets if c.gender == MALE}:
        ids = [c.id for c in cadets if c.group == g and c.gender == MALE]
        hard = Counter(r.cadet_id for r in history if r.post in HARD)
        assert max(hard[i] for i in ids) - min(hard[i] for i in ids) <= 2, g


def test_twelve_weeks_no_empty_slots_no_consecutive_days_no_same_post_streaks(sim_new, sim_old):
    for _, results, history in (sim_new,):
        assert not any(r.empty_slots for r in results)
        days_of = defaultdict(set)
        for r in history:
            days_of[r.cadet_id].add(r.day)
        assert not any(d + dt.timedelta(days=1) in ds for ds in days_of.values() for d in ds)
    # shu joy ketma-ket: yangi tartibda faqat YuP-125 (11 kishi) da majburiy holat qolishi mumkin
    cadets, _, history = sim_new
    group = {c.id: c.group for c in cadets if c.gender == MALE}
    by_cadet = defaultdict(list)
    for r in history:
        if r.cadet_id in group:
            by_cadet[r.cadet_id].append(r.post)
    streaks = Counter(group[c] for c, posts in by_cadet.items() for x, y in pairwise(posts) if x == y)
    assert set(streaks) <= {"YuP-125"} and streaks["YuP-125"] <= 1


def test_smallest_group_yup125(sim_new):
    """YuP-125 (11 kishi): bo'sh o'rin yo'q, QOIDA B saqlanadi, QOIDA A yumshatilsa — ogohlantirish bor."""
    cadets, results, history = sim_new
    ids = {c.id for c in cadets if c.group == "YuP-125" and c.gender == MALE}
    assert len(ids) == 11
    days_with = [r for r in results if any("ikkinchi marta og‘ir joyga" in w for w in r.warnings)
                 and any(a.cadet_id in ids for a in r.assignments if a.slot.post in HARD)]
    assert days_with  # haqiqatan yumshatilgan holat bor va u ogohlantirilgan
    dates = defaultdict(list)
    for r in history:
        if r.cadet_id in ids:
            dates[r.cadet_id].append(r.day)
    assert all((b - a).days > 1 for ds in dates.values() for a, b in pairwise(ds))


# ---------------------------------------------------------------------------
# Baza, sahifalar, ruxsatlar
# ---------------------------------------------------------------------------


@pytest.fixture
def admin(db_cadets):
    for code, group in INITIAL_ROTATION.items():
        InitialRotation.objects.create(post=Post.objects.get(code=code), group=StudyGroup.objects.get(code=group))
    DutyWeek.objects.create(start_date=SUNDAY)
    Setting.load()
    return User.objects.create_user("kursboshligi", password=PASSWORD, role="admin")


def generate_and_approve(day, admin):
    draft = services.generate_draft(day, admin)
    return days.approve(day, admin, draft.version)


def as_rows(duty_day):
    return sorted((a.post_slot.post.code, a.post_slot.position, a.cadet_id) for a in duty_day.assignments.select_related("post_slot__post"))


def engine_rows(result):
    return sorted((a.slot.post, a.slot.position, a.cadet_id) for a in result.assignments)


@pytest.mark.django_db
@pytest.mark.parametrize("enabled", [False, True])
def test_db_generator_follows_setting(admin, enabled):
    Setting.objects.filter(pk=1).update(post_rotation_enabled=enabled)
    history = []
    for i in range(7):
        day = SUNDAY + dt.timedelta(days=i)
        db = generate_and_approve(day, admin)
        res = generate_day(day, load_roster(), history, initial_rotation=INITIAL_ROTATION, rotation=ROT if enabled else None)
        history.extend(res.to_records())
        assert as_rows(db) == engine_rows(res), day


@pytest.fixture
def no_rest_case(admin):
    """Yakshanba tasdiqlangan, dushanba qoralama; X — yakshanba turgan va dushanba o'z guruhi joyiga nomzod."""
    generate_and_approve(SUNDAY, admin)
    monday = SUNDAY + dt.timedelta(days=1)
    draft = services.generate_draft(monday, admin)
    sunday = DutyDay.objects.get(date=SUNDAY)
    x = sunday.assignments.filter(post_slot__post__code="noj1").select_related("cadet__group").first().cadet
    post = draft.group_posts.get(group=x.group).post
    slot = PostSlot.objects.filter(post=post).order_by("-position").first()
    return draft, x, slot


@pytest.mark.django_db
def test_admin_can_place_unrested_cadet_only_with_reason_and_confirm(client, admin, no_rest_case):
    draft, x, slot = no_rest_case
    cands = {c.cadet.pk: c for c in days.replacement_candidates(draft, slot)}
    assert cands[x.pk].no_rest and not cands[x.pk].excluded and not cands[x.pk].best
    with pytest.raises(ServiceError, match="dam olmagan — sababini yozing"):
        days.replace_cadet(draft.date, slot.pk, x.pk, "sababli", "", admin, draft.version)
    with pytest.raises(ServiceError, match="qo‘yishni tasdiqlang"):
        days.replace_cadet(draft.date, slot.pk, x.pk, "sababli", "", admin, draft.version, no_rest_reason="Buyruq")
    assert not Assignment.objects.filter(duty_day=draft, cadet=x).exists()
    client.force_login(admin)
    url = reverse("day_replace", args=[draft.date.isoformat(), slot.pk])
    html = client.get(url).content.decode()
    assert 'data-testid="no-rest-candidate"' in html and "Bu kursant kecha naryadda bo‘lgan, dam olmagan" in html
    resp = client.post(url, {"cadet": x.pk, "reason": "sababli", "note": "", "version": draft.version,
                             "no_rest_reason": "Kishi yetmadi", "no_rest_confirm": "1"})
    assert resp.status_code == 302
    a = Assignment.objects.get(duty_day=draft, cadet=x)
    assert a.no_rest and a.source == "manual"
    log = ChangeLog.objects.get(action=ChangeLog.Action.NO_REST)
    assert log.reason == "Kishi yetmadi" and log.new_cadet == x and log.day_date == draft.date and log.user == admin
    assert ChangeLog.objects.filter(action=ChangeLog.Action.REPLACE, day_date=draft.date).exists()
    page = client.get(reverse("day_detail", args=[draft.date.isoformat()])).content.decode()
    assert 'data-testid="no-rest-badge"' in page and 'data-testid="rotation-notices"' in page
    # sutka tasdiqlanadi (ogohlantirish to'sqinlik qilmaydi); kursant ko'rinishida belgi yo'q
    draft.refresh_from_db()
    days.approve(draft.date, admin, draft.version)
    client.force_login(User.objects.create_user("k1", password=PASSWORD, role="cadet", cadet=x))
    page = client.get(reverse("day_detail", args=[draft.date.isoformat()])).content.decode()
    assert "no-rest-badge" not in page and "rotation-notices" not in page


@pytest.mark.django_db
def test_swap_keeps_no_rest_mark(admin, no_rest_case):
    draft, x, slot = no_rest_case
    days.replace_cadet(draft.date, slot.pk, x.pk, "sababli", "", admin, draft.version, no_rest_reason="Buyruq", no_rest_confirm=True)
    draft.refresh_from_db()
    other = PostSlot.objects.filter(post=slot.post).exclude(role=slot.role).first()
    days.swap_slots(draft.date, slot.pk, other.pk, "sababli", "", admin, draft.version)
    assert Assignment.objects.get(duty_day=draft, cadet=x).no_rest


@pytest.mark.django_db
def test_unrested_cadet_blocked_when_setting_off(admin, no_rest_case):
    draft, x, slot = no_rest_case
    Setting.objects.filter(pk=1).update(post_rotation_enabled=False)
    cands = {c.cadet.pk: c for c in days.replacement_candidates(draft, slot)}
    assert "kecha" in cands[x.pk].excluded and not cands[x.pk].no_rest
    with pytest.raises(ServiceError, match="ketma-ket ikki sutka turmaydi"):
        days.replace_cadet(draft.date, slot.pk, x.pk, "sababli", "", admin, draft.version, no_rest_reason="x", no_rest_confirm=True)


@pytest.mark.django_db
def test_commander_and_cadet_get_403(client, admin, no_rest_case):
    draft, x, slot = no_rest_case
    boss = CadetModel.objects.get(duty_position=f"{x.group.code} guruh komandiri")
    users = [User.objects.create_user("komandir", password=PASSWORD, role="commander", cadet=boss),
             User.objects.create_user("kursant", password=PASSWORD, role="cadet", cadet=x)]
    url = reverse("day_replace", args=[draft.date.isoformat(), slot.pk])
    for user in users:
        client.force_login(user)
        assert client.get(url).status_code == 403
        resp = client.post(url, {"cadet": x.pk, "reason": "sababli", "version": draft.version,
                                 "no_rest_reason": "x", "no_rest_confirm": "1"})
        assert resp.status_code == 403
    assert not Assignment.objects.filter(duty_day=draft, cadet=x).exists()
    assert not ChangeLog.objects.filter(action=ChangeLog.Action.NO_REST).exists()


@pytest.mark.django_db
def test_replace_window_warns_but_allows_second_hard_post_and_same_post(client, admin):
    for i in range(2):
        generate_and_approve(SUNDAY + dt.timedelta(days=i), admin)
    draft = services.generate_draft(TUESDAY, admin)
    slot = PostSlot.objects.get(post__code="yotoq1", position=1)  # TUESDAY: yakshanbadagi 1-NO‘J guruhi
    cands = days.replacement_candidates(draft, slot)
    warned = [c for c in cands if c.hard_conflict]
    assert warned and all(not c.excluded for c in warned)  # taqiq emas — ogohlantirish
    client.force_login(admin)
    html = client.get(reverse("day_replace", args=[TUESDAY.isoformat(), slot.pk])).content.decode()
    assert 'data-testid="rotation-warning"' in html and "shu hafta og‘ir joyda turgan" in html
    days.replace_cadet(TUESDAY, slot.pk, warned[0].cadet.pk, "sababli", "", admin, draft.version)
    page = client.get(reverse("day_detail", args=[TUESDAY.isoformat()])).content.decode()
    assert "bu hafta ikkinchi marta og‘ir joyda" in page and "qo‘lda qo‘yilgan" in page


@pytest.mark.django_db
@pytest.mark.parametrize("lang, words", [
    ("uz", ("Og‘ir joylar", "Tanlash tushuntirishi", "Joylar navbati va og‘ir joylar")),
    ("ru", ("Тяжёлые посты", "Пояснение выбора", "Очерёдность постов и тяжёлые посты")),
    ("en", ("Hard posts", "Selection explanation", "Post rotation and hard posts")),
])
def test_pages_open_in_three_languages(client, admin, lang, words):
    for i in range(2):
        generate_and_approve(SUNDAY + dt.timedelta(days=i), admin)
    services.generate_draft(TUESDAY, admin)
    user = User.objects.create_superuser(f"boss-{lang}", password=PASSWORD, role="admin", language=lang)
    client.force_login(user)
    day = client.get(reverse("day_detail", args=[TUESDAY.isoformat()])).content.decode()
    assert words[1] in day and 'data-testid="explanations"' in day
    stats = client.get(reverse("stats")).content.decode()
    assert words[0] in stats and 'data-testid="hard-count"' in stats and 'data-testid="hard-spread"' in stats
    someone = DutyDay.objects.get(date=SUNDAY).assignments.filter(post_slot__post__code="noj1").first().cadet
    card = client.get(reverse("cadet_detail", args=[someone.pk])).content.decode()
    assert 'data-testid="post-rotation"' in card and 'data-testid="cycle"' in card and "✓" in card
    with translation.override(lang):
        setting = client.get(reverse("admin:naryad_setting_change", args=[1])).content.decode()
    assert words[2] in setting
    for html in (day, stats, card, setting):
        assert not exempt_words(html), lang


@pytest.mark.django_db
def test_cadet_card_rotation_only_for_admin(client, admin):
    generate_and_approve(SUNDAY, admin)
    a = DutyDay.objects.get(date=SUNDAY).assignments.select_related("cadet").filter(post_slot__post__code="noj1").first()
    client.force_login(admin)
    html = client.get(reverse("cadet_detail", args=[a.cadet.pk])).content.decode()
    assert re.search(r'data-testid="hard-count">Og‘ir joylar: <b>1</b> marta', html)
    assert "1-NO‘J ✓" in html and "2-NO‘J –" in html
    client.force_login(User.objects.create_user("k1", password=PASSWORD, role="cadet", cadet=a.cadet))
    assert 'data-testid="post-rotation"' not in client.get(reverse("cadet_detail", args=[a.cadet.pk])).content.decode()


@pytest.mark.django_db
def test_excel_has_hard_posts_and_position_only(client, admin):
    for i in range(3):
        generate_and_approve(SUNDAY + dt.timedelta(days=i), admin)
    client.force_login(admin)
    wb = load_workbook(io.BytesIO(client.get(reverse("stats_excel")).content))
    text = " ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    header = [c.value for c in wb["Guruh YuP-125"][3]]
    assert "Og‘ir joylar" in header and "Lavozim" in header
    assert "Kurs starshinasi" in text and "Og‘ir joylar farqi" in text
    names = " ".join(" ".join(n) for n in CadetModel.objects.values_list("last_name", "first_name", "patronymic")).lower()
    found = {w for w in re.findall(r"\w*(?:ozod|exempt|освобожд)\w*|naryadga chiqmaydi", text, re.IGNORECASE) if w.lower() not in names}
    assert not found  # kursant ismlari (Ozodbek) hisobga olinmaydi
    pdf = client.get(reverse("day_pdf", args=[SUNDAY.isoformat()]))
    assert pdf.status_code == 200 and b"chiqmaydi" not in pdf.content


@pytest.mark.django_db
def test_nine_position_holders_never_candidates(admin):
    holders = set(CadetModel.objects.filter(duty_exempt=True).values_list("pk", flat=True))
    assert len(holders) == 9 and not holders & {c.id for c in services.engine_cadets()}
    labels = sorted({c.position_label for c in CadetModel.objects.filter(duty_exempt=True)})
    assert labels == ["Guruh komandiri", "Kaptyorchik", "Kurs pisari", "Kurs starshinasi"]


@pytest.mark.django_db
def test_setting_toggle_logged_and_affects_only_future_days(client, admin):
    first = generate_and_approve(SUNDAY, admin)
    before = as_rows(first)
    boss = User.objects.create_superuser("boss", password=PASSWORD, role="admin")
    client.force_login(boss)
    url = reverse("admin:naryad_setting_change", args=[1])
    assert client.get(url).status_code == 200
    s = Setting.load()
    notes = list(PrepNote.objects.order_by("post__display_order"))
    data = {
        "recipient_position": s.recipient_position, "recipient_name": s.recipient_name,
        "course_head_position": s.course_head_position, "course_head_rank": s.course_head_rank,
        "course_head_name": s.course_head_name, "course_number": 2, "approval_lead_hours": 24, "penalty_per_post": 1,
        "commander_penalty_max": 3, "commander_order_max_age_days": 30, "commander_daily_limit": 5,
        "hard_posts": ["noj1", "yotoq1"],  # post_rotation_enabled belgilanmagan — o'chiriladi
        "prep_notes-TOTAL_FORMS": str(len(notes)), "prep_notes-INITIAL_FORMS": str(len(notes)),
        "prep_notes-MIN_NUM_FORMS": "0", "prep_notes-MAX_NUM_FORMS": "1000",
        **{f"prep_notes-{i}-id": n.pk for i, n in enumerate(notes)},
        **{f"prep_notes-{i}-setting": s.pk for i, n in enumerate(notes)},
        "initial_rotation-TOTAL_FORMS": "6", "initial_rotation-INITIAL_FORMS": "6",
        "initial_rotation-MIN_NUM_FORMS": "0", "initial_rotation-MAX_NUM_FORMS": "6",
    }
    for i, r in enumerate(InitialRotation.objects.order_by("pk")):
        data.update({f"initial_rotation-{i}-id": r.pk, f"initial_rotation-{i}-setting": 1,
                     f"initial_rotation-{i}-post": r.post_id, f"initial_rotation-{i}-group": r.group_id})
    resp = client.post(url, data)
    assert resp.status_code == 302, resp.content.decode()[:3000]
    assert not Setting.load().post_rotation_enabled and Setting.load().post_rotation() is None
    log = ChangeLog.objects.get(action=ChangeLog.Action.POST_ROTATION)
    assert "o‘chirildi" in log.details and log.user == boss
    assert as_rows(DutyDay.objects.get(date=SUNDAY)) == before  # tasdiqlangan sutka o'zgarmaydi
    data["post_rotation_enabled"] = "on"
    data["hard_posts"] = ["noj1"]
    assert client.post(url, data).status_code == 302
    assert Setting.load().post_rotation() == PostRotation(hard_posts=("noj1",))
    assert ChangeLog.objects.filter(action=ChangeLog.Action.POST_ROTATION).count() == 2
