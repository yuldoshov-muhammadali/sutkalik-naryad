"""Kursant akkauntlarini ommaviy yaratish (create_cadet_accounts buyrug'i, boshqaruv panelidagi amallar).

19-bosqich: guruh komandiri lavozimidagi kursantga rol avtomatik `commander` (guruhida faol komandir bo'lmasa).
"""

from __future__ import annotations

import re
import secrets
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from .models import Cadet, User, UserRole

PASSWORD_ALPHABET = "abcdefghjkmnpqrstuvwxyz23456789"  # adashtiradigan belgilarsiz (0/o, 1/l/i)
PASSWORD_LENGTH = 10


def latin_slug(text: str) -> str:
    """O'zbekcha (lotin) matnni login uchun: kichik harf, faqat a-z va 0-9."""
    text = re.sub(r"[ʻʼ'`´‘’]", "", text or "")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def base_username(cadet: Cadet) -> str:
    first = (cadet.first_name or "").split()
    return f"{latin_slug(cadet.last_name)}.{latin_slug(first[0] if first else '')}".strip(".")


def unique_username(base: str, taken: set[str]) -> str:
    candidate, n = base, 1
    while candidate.casefold() in taken:
        n += 1
        candidate = f"{base}{n}"
    taken.add(candidate.casefold())
    return candidate


def temp_password() -> str:
    while True:
        pw = "".join(secrets.choice(PASSWORD_ALPHABET) for _ in range(PASSWORD_LENGTH))
        if any(c.isdigit() for c in pw) and any(c.isalpha() for c in pw):
            return pw


@dataclass
class CreatedAccount:
    cadet: Cadet
    username: str
    password: str


@dataclass
class AccountsReport:
    created: list[CreatedAccount] = field(default_factory=list)
    skipped: int = 0  # akkaunti allaqachon bor
    file: Path | None = None


def role_for(cadet: Cadet) -> str:
    """Guruh komandiri — commander (guruhida boshqa faol komandir bo'lmasa), qolganlar — cadet."""
    if cadet.is_group_commander and not User.objects.filter(
        role=UserRole.COMMANDER, is_active=True, cadet__group_id=cadet.group_id
    ).exists():
        return UserRole.COMMANDER
    return UserRole.CADET


@transaction.atomic
def create_cadet_accounts(output_dir: Path | None, today=None, cadets=None) -> AccountsReport:
    """Akkaunti yo'q faol kursantlarga akkaunt. `output_dir=None` — fayl yozilmaydi (boshqaruv paneli Excel'ni o'zi beradi)."""
    report = AccountsReport()
    taken = {u.casefold() for u in User.objects.values_list("username", flat=True)}
    if cadets is None:
        cadets = Cadet.objects.filter(is_active=True)
    cadets = cadets.filter(is_active=True).select_related("group").order_by("group__display_order", "list_number", "pk")
    for cadet in cadets:
        if User.objects.filter(cadet=cadet).exists():
            report.skipped += 1
            continue
        username = unique_username(base_username(cadet) or f"kursant{cadet.pk}", taken)
        password = temp_password()
        user = User(
            username=username,
            first_name=cadet.first_name,
            last_name=cadet.last_name,
            role=role_for(cadet),
            cadet=cadet,
            must_change_password=True,
        )
        user.set_password(password)
        user.full_clean(exclude=["password"])
        user.save()
        report.created.append(CreatedAccount(cadet=cadet, username=username, password=password))
    if report.created and output_dir is not None:
        report.file = write_accounts_xlsx(report.created, Path(output_dir), today or timezone.localdate())
    return report


@transaction.atomic
def reset_passwords(users) -> list[CreatedAccount]:
    """Parolni tiklash: yangi vaqtinchalik parol, birinchi kirishda almashtirish majburiy (komandirga ham)."""
    out = []
    for user in users.select_related("cadet__group"):
        if user.cadet is None:
            continue
        password = temp_password()
        user.set_password(password)
        user.must_change_password = True
        user.save()
        out.append(CreatedAccount(cadet=user.cadet, username=user.username, password=password))
    return out


def write_accounts_xlsx(accounts: list[CreatedAccount], output_dir: Path, today) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"akkauntlar_{today:%Y-%m-%d}.xlsx"
    n = 1
    while path.exists():
        n += 1
        path = output_dir / f"akkauntlar_{today:%Y-%m-%d}_{n}.xlsx"
    accounts_workbook(accounts, today).save(path)
    return path


def accounts_xlsx_bytes(accounts: list[CreatedAccount], today) -> bytes:
    import io

    out = io.BytesIO()
    accounts_workbook(accounts, today).save(out)
    return out.getvalue()


def accounts_workbook(accounts: list[CreatedAccount], today):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "akkauntlar"
    ws.append([f"{settings.SITE_NAME} — kursant akkauntlari ({today:%d.%m.%Y})"])
    ws.append(["Parol vaqtinchalik: birinchi kirishda o‘zingizning parolingizga almashtirasiz. Varaqni kesib, egasiga bering."])
    ws.append([])
    header = ["№", "Guruh", "F.I.Sh.", "Login", "Vaqtinchalik parol", "Rol", "Imzo"]
    ws.append(header)
    thin = Side(style="thin", color="999999")
    for i, acc in enumerate(accounts, start=1):
        role = "Guruh komandiri" if acc.cadet.is_group_commander and User.objects.filter(
            username=acc.username, role=UserRole.COMMANDER).exists() else "Kursant"
        ws.append([i, acc.cadet.group.code, acc.cadet.full_name, acc.username, acc.password, role, ""])
    ws["A1"].font = Font(bold=True, size=13)
    for cell in ws[4]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            cell.alignment = Alignment(vertical="center")
        if row[0].row > 4:
            row[3].font = Font(name="Consolas", size=12)
            row[4].font = Font(name="Consolas", size=12, bold=True)
    for col, width in zip("ABCDEFG", [5, 10, 42, 26, 20, 16, 16]):
        ws.column_dimensions[col].width = width
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    return wb
