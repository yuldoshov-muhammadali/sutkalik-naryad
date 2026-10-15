"""Kursantlarni Excel'dan import qilish: oldindan ko'rish → tasdiqlash (6-bosqich).

Ustunlar: familiya, ism, otasining ismi, guruh, jins.
Bitta ham xatoli qator bo'lsa import tasdiqlanmaydi. Faylda yo'q, lekin bazada
bor kursantlar o'chirilmaydi — faqat ro'yxat qilib ko'rsatiladi.
"""

from __future__ import annotations

import io
import re
from itertools import islice
from dataclasses import asdict, dataclass, field

from django.db import transaction
from django.utils import translation
from django.utils.translation import gettext as _

from naryad.engine.names import normalize_apostrophes, short_name as auto_short_name

from .cadets import has_history, next_list_number
from .models import Cadet, ChangeLog, Gender, StudyGroup

COLUMNS = {
    "familiya": "last_name",
    "ism": "first_name",
    "otasining ismi": "patronymic",
    "guruh": "group",
    "jins": "gender",
}
GENDER_WORDS = {
    "o‘g‘il": Gender.MALE, "ogil": Gender.MALE, "o‘gil": Gender.MALE, "erkak": Gender.MALE, "m": Gender.MALE,
    "male": Gender.MALE, "yigit": Gender.MALE,
    "qiz": Gender.FEMALE, "ayol": Gender.FEMALE, "f": Gender.FEMALE, "female": Gender.FEMALE, "q": Gender.FEMALE,
}
MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000


def norm(text) -> str:
    return re.sub(r"\s+", " ", normalize_apostrophes(str(text or "")).strip())


def key(text) -> str:
    return norm(text).casefold()


@dataclass
class ImportRow:
    line: int
    last_name: str
    first_name: str
    patronymic: str
    group: str
    gender: str


@dataclass
class Preview:
    rows: list[dict] = field(default_factory=list)  # sessiyada saqlash uchun
    new: list[dict] = field(default_factory=list)
    unchanged: list[dict] = field(default_factory=list)
    # {"row":…, "cadet_id":…, "changes":[…] (o'zbekcha, jurnal uchun), "changes_display":[…] (joriy tilda), "history_warning":bool}
    updated: list[dict] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    missing: list[dict] = field(default_factory=list)  # bazada bor, faylda yo'q

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.rows)


class ImportFileError(ValueError):
    pass


def _add_change(changes: list[str], changes_display: list[str], build) -> None:
    """O'zgarish matni ikki marta: o'zbekcha (jurnalga yoziladi) va joriy tilda (ekranda ko'rsatiladi)."""
    changes_display.append(build())
    with translation.override("uz"):
        changes.append(build())


MAX_UNPACKED_BYTES = 30 * 1024 * 1024


def check_xlsx_container(data: bytes) -> None:
    """Haqiqiy .xlsx (ZIP) ekanini va ochilgandagi hajmi oqilona ekanini tekshiradi ("zip bomba"dan himoya)."""
    import zipfile

    if not data.startswith(b"PK\x03\x04"):
        raise ImportFileError(_("Fayl .xlsx (Excel) formatida emas."))
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            unpacked = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile as exc:
        raise ImportFileError(_("Faylni o‘qib bo‘lmadi — .xlsx (Excel) formatida ekanini tekshiring.")) from exc
    if "[Content_Types].xml" not in names or not any(n.startswith("xl/") for n in names):
        raise ImportFileError(_("Fayl .xlsx (Excel) formatida emas."))
    if unpacked > MAX_UNPACKED_BYTES:
        raise ImportFileError(_("Fayl ichidagi ma’lumot juda katta."))


def read_xlsx(data: bytes) -> list[ImportRow]:
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError(_("Fayl juda katta (2 MB dan oshmasin)."))
    check_xlsx_container(data)
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — har qanday buzilgan fayl
        raise ImportFileError(_("Faylni o‘qib bo‘lmadi — .xlsx (Excel) formatida ekanini tekshiring.")) from exc
    ws = wb.worksheets[0]
    raw = list(islice(ws.iter_rows(values_only=True, max_col=30), MAX_ROWS + 2))
    wb.close()
    if len(raw) > MAX_ROWS + 1:
        raise ImportFileError(_("Qatorlar juda ko‘p; ko‘pi bilan %(max)s ta.") % {"max": MAX_ROWS})
    if not raw:
        raise ImportFileError(_("Fayl bo‘sh."))
    header = [key(h) for h in raw[0]]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        # Ustun nomlari (familiya, ism, …) — fayldagi o'zbekcha sarlavhalar, tarjima qilinmaydi
        raise ImportFileError(
            _("Sarlavha qatorida ustunlar yetishmaydi: %(missing)s. Kerakli: %(required)s.")
            % {"missing": ", ".join(missing), "required": ", ".join(COLUMNS)}
        )
    idx = {COLUMNS[c]: header.index(c) for c in COLUMNS}
    rows = []
    for line, values in enumerate(raw[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        get = lambda f: norm(values[idx[f]]) if idx[f] < len(values) else ""  # noqa: E731
        rows.append(ImportRow(line, get("last_name"), get("first_name"), get("patronymic"), get("group"), get("gender")))
    if len(rows) > MAX_ROWS:
        raise ImportFileError(_("Qatorlar juda ko‘p (%(count)s); ko‘pi bilan %(max)s ta.") % {"count": len(rows), "max": MAX_ROWS})
    return rows


def build_preview(rows: list[ImportRow]) -> Preview:
    preview = Preview(rows=[asdict(r) for r in rows])
    groups = {key(g.code): g for g in StudyGroup.objects.all()}
    cadets = list(Cadet.objects.select_related("group"))
    by_full = {}
    for c in cadets:
        by_full.setdefault((key(c.last_name), key(c.first_name), key(c.patronymic)), []).append(c)
    by_name_group = {}
    for c in cadets:
        by_name_group.setdefault((c.group_id, key(c.last_name), key(c.first_name)), []).append(c)

    seen: dict[tuple, int] = {}
    matched_ids: set[int] = set()
    for r in rows:
        problems = []
        if not r.last_name:
            problems.append(_("familiya bo‘sh"))
        if not r.first_name:
            problems.append(_("ism bo‘sh"))
        group = groups.get(key(r.group))
        if group is None:
            problems.append(_("guruh topilmadi: «%(group)s»") % {"group": r.group} if r.group else _("guruh bo‘sh"))
        gender = GENDER_WORDS.get(key(r.gender).replace("'", "‘").replace("’", "‘"))
        if gender is None:
            problems.append(_("jins noto‘g‘ri: «%(gender)s» (o‘g‘il yoki qiz)") % {"gender": r.gender})
        full_key = (key(r.last_name), key(r.first_name), key(r.patronymic))
        if r.last_name and r.first_name and group:
            dup_key = (group.pk, *full_key)
            if dup_key in seen:
                problems.append(_("fayl ichida takrorlangan (%(line)s-qator bilan bir xil)") % {"line": seen[dup_key]})
            seen.setdefault(dup_key, r.line)
        if problems:
            preview.errors.extend(_("%(line)s-qator: %(problem)s") % {"line": r.line, "problem": p} for p in problems)
            continue

        item = {**asdict(r), "group": group.code, "gender": gender, "gender_label": str(Gender(gender).label)}
        exact = [c for c in by_full.get(full_key, []) if c.group_id == group.pk]
        same_person_other_group = [c for c in by_full.get(full_key, []) if c.group_id != group.pk]
        same_name_in_group = [
            c for c in by_name_group.get((group.pk, key(r.last_name), key(r.first_name)), []) if c not in exact
        ]
        target, changes, changes_display = None, [], []
        if exact:
            target = exact[0]
        elif len(same_person_other_group) == 1:
            target = same_person_other_group[0]
            _add_change(
                changes, changes_display,
                lambda: _("Guruh: %(old)s → %(new)s") % {"old": target.group.code, "new": group.code},
            )
        elif len(same_name_in_group) == 1 and not same_person_other_group:
            target = same_name_in_group[0]
            _add_change(
                changes, changes_display,
                lambda: _("Otasining ismi: %(old)s → %(new)s") % {"old": target.patronymic or "—", "new": r.patronymic or "—"},
            )
        elif len(same_person_other_group) > 1:
            preview.errors.append(
                _("%(line)s-qator: «%(name)s» bir nechta guruhda bor — qaysi biri ekanini aniqlab bo‘lmaydi")
                % {"line": r.line, "name": f"{r.last_name} {r.first_name} {r.patronymic}"}
            )
            continue

        if target is None:
            preview.new.append(item)
            continue
        if target.pk in matched_ids:
            preview.errors.append(
                _("%(line)s-qator: bazadagi “%(name)s” faylda ikkinchi marta uchradi") % {"line": r.line, "name": target.full_name}
            )
            continue
        matched_ids.add(target.pk)
        if target.gender != gender:
            _add_change(
                changes, changes_display,
                lambda: _("Jinsi: %(old)s → %(new)s") % {"old": target.get_gender_display(), "new": Gender(gender).label},
            )
        if not changes:
            preview.unchanged.append({**item, "cadet_id": target.pk})
            continue
        history_warning = has_history(target) and any(c.startswith(("Guruh", "Jinsi")) for c in changes)  # changes — o'zbekcha
        preview.updated.append(
            {**item, "cadet_id": target.pk, "changes": changes, "changes_display": changes_display,
             "history_warning": history_warning}
        )

    preview.missing = [
        {"id": c.pk, "name": c.full_name, "group": c.group.code, "active": c.is_active}
        for c in sorted(cadets, key=lambda c: (c.group.display_order, c.last_name))
        if c.pk not in matched_ids and c.is_active
    ]
    return preview


def preview_from_session(rows: list[dict]) -> Preview:
    return build_preview([ImportRow(**r) for r in rows])


@transaction.atomic
def apply_preview(preview: Preview, user, filename: str = "") -> tuple[int, int]:
    """Oldindan ko'rilgan importni bitta tranzaksiyada yozadi. (yangi, yangilangan) sonini qaytaradi."""
    if preview.errors:
        raise ImportFileError(_("Xatoli qatorlar bor — import tasdiqlanmaydi."))
    groups = {g.code: g for g in StudyGroup.objects.all()}
    source = f" ({filename})" if filename else ""
    for item in preview.new:
        group = groups[item["group"]]
        cadet = Cadet.objects.create(
            group=group,
            list_number=next_list_number(group),
            last_name=item["last_name"],
            first_name=item["first_name"],
            patronymic=item["patronymic"],
            gender=item["gender"],
            short_name=auto_short_name(item["last_name"], item["first_name"]),
        )
        with translation.override("uz"):  # jurnal matni doim o'zbekcha saqlanadi
            details = f"Yangi kursant: {cadet.full_name}, {group.code}, {cadet.get_gender_display()}{source}"
        ChangeLog.objects.create(action=ChangeLog.Action.CADET_IMPORT, cadet=cadet, user=user, details=details)
    for item in preview.updated:
        cadet = Cadet.objects.select_for_update().get(pk=item["cadet_id"])
        if cadet.group_id != groups[item["group"]].pk:
            cadet.group = groups[item["group"]]
            cadet.list_number = next_list_number(cadet.group)
        cadet.patronymic = item["patronymic"]
        cadet.gender = item["gender"]
        cadet.save()
        details = "; ".join(item["changes"]) + source
        if item["history_warning"]:
            details += ". Naryad tarixi bor: tarix o‘chmaydi, o‘zgarish faqat keyingi naryadlarga ta’sir qiladi."
        ChangeLog.objects.create(action=ChangeLog.Action.CADET_IMPORT, cadet=cadet, user=user, details=details)
    return len(preview.new), len(preview.updated)


def template_xlsx() -> bytes:
    """Bo'sh shablon: sarlavha, ochiladigan ro'yxatlar va yo'riqnoma."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "kursantlar"
    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for col, width in zip("ABCDE", [22, 18, 28, 12, 10]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    codes = [g.code for g in StudyGroup.objects.filter(is_active=True).order_by("display_order")]
    if codes:
        dv_group = DataValidation(type="list", formula1='"' + ",".join(codes) + '"', allow_blank=True)
        ws.add_data_validation(dv_group)
        dv_group.add("D2:D1000")
    dv_gender = DataValidation(type="list", formula1='"o‘g‘il,qiz"', allow_blank=True)
    ws.add_data_validation(dv_gender)
    dv_gender.add("E2:E1000")
    ws.append(["Sa’dullayev", "Behro‘z", "Bahodir o‘g‘li", codes[4] if len(codes) > 4 else "BI-525", "o‘g‘il"])
    info = wb.create_sheet("yo‘riqnoma")
    info.column_dimensions["A"].width = 100
    for line in [
        "Kursantlarni import qilish shabloni",
        "",
        "Har qator — bitta kursant. 2-qatordagi namunani o‘chirib, o‘z ro‘yxatingizni yozing.",
        "familiya, ism — majburiy. otasining ismi — masalan: Bahodir o‘g‘li, Oybek qizi, Davronovich.",
        "guruh — " + (", ".join(codes) or "YuP-125, BI-225 …") + ".",
        "jins — o‘g‘il yoki qiz.",
        "Apostroflar (' ` ‘ ’ ʻ ʼ) avtomatik bir xil ko‘rinishga keltiriladi: o‘, g‘ va tutuq belgisi ’.",
        "Faylda yo‘q, lekin bazada bor kursantlar o‘chirilmaydi.",
    ]:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
