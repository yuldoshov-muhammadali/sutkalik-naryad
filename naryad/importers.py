"""Boshlang'ich ma'lumotlar: ro'yxatni yuklash va qog'ozdagi o'tgan tarixni import qilish.

Buyruqlar (`load_roster`, `import_history`) shu moduldagi funksiyalarni chaqiradi.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from naryad.engine.allocator import week_start
from naryad.engine.names import normalize_apostrophes

from .models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    Gender,
    GroupPostAssignment,
    Post,
    PostSlot,
    SlotRole,
    StudyGroup,
)

ROSTER_PATH = Path(settings.BASE_DIR) / "data" / "roster.json"

# Guruhlar: kod -> (ta'lim yo'nalishi, til, tartib). CLAUDE.md 2-bo'lim.
GROUP_INFO = {
    "YuP-125": ("Yurisprudensiya", "uz", 1),
    "BI-225": ("Bojxona ishi", "uz", 2),
    "BI-325": ("Bojxona ishi", "uz", 3),
    "BI-425": ("Bojxona ishi", "uz", 4),
    "BI-525": ("Bojxona ishi", "uz", 5),
    "BI-625": ("Bojxona ishi", "ru", 6),
}
# Kutilgan tarkib: guruh -> (o'g'il, qiz). CLAUDE.md 2-bo'lim.
EXPECTED_COMPOSITION = {
    "YuP-125": (13, 2),
    "BI-225": (14, 3),
    "BI-325": (14, 2),
    "BI-425": (15, 2),
    "BI-525": (17, 0),
    "BI-625": (14, 4),
}
EXPECTED_TOTAL = (87, 13)
# Naryadga chiqadigan o'g'il bolalar (komandirlar va kurs lavozimlari ozod) — CLAUDE.md 2-bo'lim
EXPECTED_ON_DUTY = {"YuP-125": 11, "BI-225": 12, "BI-325": 12, "BI-425": 14, "BI-525": 16, "BI-625": 13}
EXPECTED_ON_DUTY_TOTAL = 78


def norm(text: object) -> str:
    """Taqqoslash uchun: apostroflar bir xil, kichik harf, ortiqcha bo'shliqsiz."""
    return re.sub(r"\s+", " ", normalize_apostrophes(str(text or "")).strip()).casefold()


# ---------------------------------------------------------------------------
# load_roster
# ---------------------------------------------------------------------------


@dataclass
class RosterReport:
    groups_created: int = 0
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    composition: dict[str, tuple[int, int]] = field(default_factory=dict)
    on_duty: dict[str, int] = field(default_factory=dict)  # naryadga chiqadigan o'g'il bolalar
    exempt: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def totals(self) -> tuple[int, int]:
        return (sum(b for b, _ in self.composition.values()), sum(g for _, g in self.composition.values()))


ROSTER_FIELDS = ("last_name", "first_name", "patronymic", "gender", "short_name", "duty_exempt", "duty_position")
EXEMPT_REASON = "data/roster.json: komandirlar va kurs lavozimlari sutkalik naryadga qo‘yilmaydi (load_roster)."


@transaction.atomic
def load_roster(path: Path = ROSTER_PATH) -> RosterReport:
    """roster.json dan guruhlar va kursantlarni yuklaydi (idempotent).

    Kursant (guruh, ro'yxatdagi raqam) bo'yicha topiladi; o'zgargan maydonlar yangilanadi.
    """
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    report = RosterReport()

    groups: dict[str, StudyGroup] = {}
    for code in sorted({row["group"] for row in data}, key=lambda c: GROUP_INFO.get(c, ("", "", 99))[2]):
        direction, language, order = GROUP_INFO.get(code, ("", "uz", None))
        if order is None:
            report.warnings.append(f"Noma’lum guruh: {code} (CLAUDE.md 2-bo‘limida yo‘q)")
            order = 100 + len(groups)
        group, created = StudyGroup.objects.get_or_create(
            code=code, defaults={"direction": direction, "language": language, "display_order": order}
        )
        report.groups_created += int(created)
        groups[code] = group

    for row in data:
        values = {
            "last_name": normalize_apostrophes(row["last_name"]),
            "first_name": normalize_apostrophes(row["first_name"]),
            "patronymic": normalize_apostrophes(row.get("patronymic", "")),
            "gender": row["gender"],
            "short_name": normalize_apostrophes(row.get("short_name", "")),
            "duty_exempt": bool(row.get("duty_exempt", False)),
            "duty_position": normalize_apostrophes(row.get("duty_position", "")) if row.get("duty_exempt") else "",
        }
        cadet = Cadet.objects.filter(group=groups[row["group"]], list_number=row["list_number"]).first()
        if cadet is None:
            cadet = Cadet.objects.create(group=groups[row["group"]], list_number=row["list_number"], **values)
            report.created += 1
            if cadet.duty_exempt:
                log_exempt_change(cadet)
            continue
        changed = [f for f in ROSTER_FIELDS if getattr(cadet, f) != values[f]]
        if changed:
            exempt_changed = cadet.duty_exempt != values["duty_exempt"]
            for f in changed:
                setattr(cadet, f, values[f])
            cadet.save()
            if exempt_changed:
                log_exempt_change(cadet)
            report.updated += 1
        else:
            report.unchanged += 1

    for code, group in sorted(groups.items(), key=lambda kv: kv[1].display_order):
        boys = Cadet.objects.filter(group=group, gender=Gender.MALE, is_active=True).count()
        girls = Cadet.objects.filter(group=group, gender=Gender.FEMALE, is_active=True).count()
        report.composition[code] = (boys, girls)
        expected = EXPECTED_COMPOSITION.get(code)
        if expected and expected != (boys, girls):
            report.warnings.append(f"{code}: kutilgan {expected[0]} o‘g‘il / {expected[1]} qiz, bazada {boys} / {girls}")
        report.on_duty[code] = Cadet.objects.filter(group=group, gender=Gender.MALE, is_active=True, duty_exempt=False).count()
        if code in EXPECTED_ON_DUTY and EXPECTED_ON_DUTY[code] != report.on_duty[code]:
            report.warnings.append(
                f"{code}: naryadga chiqadigan o‘g‘il bolalar kutilgan {EXPECTED_ON_DUTY[code]}, bazada {report.on_duty[code]}"
            )
    report.exempt = [
        f"{c.group.code} — {c.full_name} ({c.duty_position})"
        for c in Cadet.objects.filter(duty_exempt=True).select_related("group").order_by("group__display_order", "list_number")
    ]
    if report.totals != EXPECTED_TOTAL:
        report.warnings.append(
            f"Jami: kutilgan {EXPECTED_TOTAL[0]} o‘g‘il / {EXPECTED_TOTAL[1]} qiz, bazada {report.totals[0]} / {report.totals[1]}"
        )
    return report


def log_exempt_change(cadet: Cadet) -> None:
    ChangeLog.objects.create(
        action=ChangeLog.Action.CADET_EXEMPT if cadet.duty_exempt else ChangeLog.Action.CADET_UNEXEMPT,
        cadet=cadet,
        reason=EXEMPT_REASON,
        details=f"Lavozimi: {cadet.duty_position}" if cadet.duty_exempt else "",
    )


# ---------------------------------------------------------------------------
# import_history
# ---------------------------------------------------------------------------

COLUMNS = {
    "sana": "date",
    "joy": "post",
    "lavozim": "role",
    "kursant": "cadet",
    "guruh": "group",
}
DATE_FORMATS = ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y")


@dataclass
class Row:
    line: int  # fayldagi qator raqami (sarlavha = 1)
    date: dt.date
    post: Post
    role: str
    cadet: Cadet
    group: StudyGroup | None


@dataclass
class ImportReport:
    rows_read: int = 0
    days: list[dt.date] = field(default_factory=list)
    weeks_created: list[dt.date] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    dry_run: bool = False

    @property
    def ok(self) -> bool:
        return not self.errors


class ImportFailed(Exception):
    def __init__(self, report: ImportReport):
        super().__init__("\n".join(report.errors))
        self.report = report


def read_table(path: Path) -> list[tuple[int, dict[str, object]]]:
    """xlsx yoki csv faylni o'qiydi: [(qator raqami, {ustun: qiymat})]."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    elif path.suffix.lower() == ".csv":
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t") if text.strip() else csv.excel
        rows = [tuple(r) for r in csv.reader(text.splitlines(), dialect)]
    else:
        raise ValueError(f"Fayl turi qo‘llab-quvvatlanmaydi: {path.suffix} (faqat .xlsx yoki .csv)")
    if not rows:
        raise ValueError("Fayl bo‘sh.")

    header = [norm(h) for h in rows[0]]
    missing = [name for name in COLUMNS if name not in header]
    if missing:
        raise ValueError(
            "Sarlavha qatorida ustunlar yetishmaydi: " + ", ".join(missing) + ". Kerakli ustunlar: " + ", ".join(COLUMNS)
        )
    index = {COLUMNS[name]: header.index(name) for name in COLUMNS}
    out = []
    for line, values in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append((line, {key: values[i] if i < len(values) else None for key, i in index.items()}))
    return out


def parse_date(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _post_aliases() -> dict[str, Post]:
    aliases: dict[str, Post] = {}
    extra = {"yotoq1": ["1-yotoqxona", "kurs", "kurs bo‘yicha", "1-sonli kursantlar yotoqxonasi"], "patrul": ["patrul"]}
    for post in Post.objects.all():
        names = [post.code, post.name, post.section_title, post.name.replace(" bo‘yicha", "")] + extra.get(post.code, [])
        for name in names:
            aliases[norm(name)] = post
    return aliases


def parse_role(value: object) -> str | None:
    text = norm(value)
    if "yordamchi" in text:
        return SlotRole.ASSISTANT
    if "navbatchi" in text:
        return SlotRole.DUTY
    if "posbon" in text or "patrul" in text:
        return SlotRole.GUARD
    return None


class CadetMatcher:
    """Kursantni qisqa yoki to'liq F.I.Sh. bo'yicha topadi. Taxmin qilmaydi."""

    def __init__(self):
        self.cadets = list(Cadet.objects.select_related("group"))

    @staticmethod
    def keys(c: Cadet) -> set[str]:
        return {
            norm(c.short_name),
            norm(c.full_name),
            norm(f"{c.last_name} {c.first_name}"),
        }

    def find(self, name: object, group: StudyGroup | None) -> list[Cadet]:
        key = norm(name)
        pool = [c for c in self.cadets if group is None or c.group_id == group.pk]
        return [c for c in pool if key in self.keys(c)]


def parse_rows(raw: list[tuple[int, dict[str, object]]], errors: list[str]) -> list[Row]:
    posts = _post_aliases()
    groups = {norm(g.code): g for g in StudyGroup.objects.all()}
    matcher = CadetMatcher()
    rows: list[Row] = []
    for line, v in raw:
        problems = []
        day = parse_date(v["date"])
        if day is None:
            problems.append(f"sana noto‘g‘ri: «{v['date']}» (namuna: 2026-09-18 yoki 18.09.2026)")
        post = posts.get(norm(v["post"]))
        if post is None:
            problems.append(f"joy topilmadi: «{v['post']}»")
        role = parse_role(v["role"])
        if role is None:
            problems.append(f"lavozim noto‘g‘ri: «{v['role']}» (navbatchi / navbatchi yordamchisi / posbon)")
        group = None
        if str(v["group"] or "").strip():
            group = groups.get(norm(v["group"]))
            if group is None:
                problems.append(f"guruh topilmadi: «{v['group']}»")
        cadet = None
        if not str(v["cadet"] or "").strip():
            problems.append("kursant ko‘rsatilmagan")
        elif not (str(v["group"] or "").strip() and group is None):
            found = matcher.find(v["cadet"], group)
            if not found:
                where = f" ({group.code} guruhida)" if group else ""
                problems.append(f"kursant topilmadi: «{v['cadet']}»{where}")
            elif len(found) > 1:
                names = ", ".join(f"{c.full_name} ({c.group.code})" for c in found)
                problems.append(f"bir nechta kursant mos keldi: «{v['cadet']}» → {names}; to‘liq F.I.Sh. yoki guruhni yozing")
            else:
                cadet = found[0]
        if problems:
            errors.extend(f"{line}-qator: {p}" for p in problems)
        else:
            rows.append(Row(line=line, date=day, post=post, role=role, cadet=cadet, group=group))
    return rows


def validate_day(day: dt.date, rows: list[Row], slots_by_post: dict[int, list[PostSlot]], errors: list[str]):
    """Bir sutka qatorlarini tekshiradi; muvaffaqiyatli bo'lsa (o'rin, qator) juftliklarini qaytaradi."""
    d = f"{day:%d.%m.%Y}"
    total_slots = sum(len(s) for s in slots_by_post.values())
    before = len(errors)
    lines = ", ".join(str(r.line) for r in rows)
    if len(rows) != total_slots:
        errors.append(f"{d}: {len(rows)} ta qator bor, {total_slots} ta bo‘lishi kerak (qatorlar: {lines})")

    seen: dict[int, int] = {}
    for r in rows:
        if r.cadet.pk in seen:
            errors.append(f"{r.line}-qator: {r.cadet.short_name} {d} sutkasida ikkinchi marta yozilgan ({seen[r.cadet.pk]}-qatorda ham bor)")
        seen.setdefault(r.cadet.pk, r.line)
        if r.cadet.gender != r.post.gender:
            if r.post.gender == Gender.FEMALE:
                errors.append(f"{r.line}-qator: {r.cadet.short_name} o‘g‘il — Qizlar yotoqxonasiga qo‘yib bo‘lmaydi")
            else:
                errors.append(f"{r.line}-qator: {r.cadet.short_name} qiz — qizlar faqat Qizlar yotoqxonasida turadi")
        if r.group is not None and r.group.pk != r.cadet.group_id:
            errors.append(f"{r.line}-qator: {r.cadet.short_name} {r.cadet.group.code} guruhida, faylda {r.group.code} yozilgan")

    by_post: dict[int, list[Row]] = defaultdict(list)
    for r in rows:
        by_post[r.post.pk].append(r)

    group_of_post: dict[int, int] = {}
    for post_id, post_rows in by_post.items():
        post = post_rows[0].post
        if post.gender == Gender.MALE:
            groups = {r.cadet.group.code for r in post_rows}
            if len(groups) > 1:
                errors.append(
                    f"{d}, {post.name}: bir nechta guruhdan kursant ({', '.join(sorted(groups))}) — har joy bitta guruhdan bo‘lishi kerak"
                )
            else:
                group_of_post[post_id] = post_rows[0].cadet.group_id
    dup_groups = [g for g, n in Counter(group_of_post.values()).items() if n > 1]
    for g in dup_groups:
        code = StudyGroup.objects.get(pk=g).code
        errors.append(f"{d}: {code} guruhi bir sutkada bir nechta joyga yozilgan — har guruh bitta joy beradi")

    pairs: list[tuple[PostSlot, Row]] = []
    for post_id, slots in slots_by_post.items():
        post_rows = by_post.get(post_id, [])
        need = Counter(s.role for s in slots)
        have = Counter(r.role for r in post_rows)
        if need != have:
            post = slots[0].post
            fmt = lambda c: ", ".join(f"{SlotRole(k).label.lower()}: {v}" for k, v in sorted(c.items())) or "yo‘q"  # noqa: E731
            errors.append(f"{d}, {post.name}: lavozimlar mos emas — kerak ({fmt(need)}), faylda ({fmt(have)})")
            continue
        for role in need:
            role_slots = [s for s in slots if s.role == role]
            role_rows = [r for r in post_rows if r.role == role]
            pairs.extend(zip(role_slots, role_rows))
    return pairs if len(errors) == before else None


def import_history(path: Path, *, dry_run: bool = False, user=None) -> ImportReport:
    """Qog'ozdagi o'tgan sutkalarni tasdiqlangan holda import qiladi (hammasi yoki hech narsa)."""
    report = ImportReport(dry_run=dry_run)
    try:
        raw = read_table(path)
    except ValueError as exc:
        report.errors.append(str(exc))
        raise ImportFailed(report) from exc
    report.rows_read = len(raw)

    rows = parse_rows(raw, report.errors)
    by_day: dict[dt.date, list[Row]] = defaultdict(list)
    for r in rows:
        by_day[r.date].append(r)
    report.days = sorted(by_day)

    existing = set(DutyDay.objects.filter(date__in=report.days).values_list("date", flat=True))
    for day in sorted(existing):
        report.errors.append(f"{day:%d.%m.%Y}: bu sutka bazada allaqachon bor — qayta import qilinmaydi")

    slots_by_post: dict[int, list[PostSlot]] = defaultdict(list)
    for s in PostSlot.objects.select_related("post").order_by("post__display_order", "position"):
        slots_by_post[s.post_id].append(s)

    plan: dict[dt.date, list[tuple[PostSlot, Row]]] = {}
    for day in report.days:
        pairs = validate_day(day, by_day[day], slots_by_post, report.errors)
        if pairs is not None:
            plan[day] = pairs

    if report.errors:
        raise ImportFailed(report)

    report.warnings.extend(consecutive_day_warnings(plan))

    try:
        with transaction.atomic():
            _write(plan, report, user, Path(path).name)
            if dry_run or report.errors:
                raise _Rollback  # --dry-run yoki xato: hech narsa yozilmaydi
    except _Rollback:
        pass
    if report.errors:
        raise ImportFailed(report)
    return report


class _Rollback(Exception):
    pass


def consecutive_day_warnings(plan: dict[dt.date, list[tuple[PostSlot, Row]]]) -> list[str]:
    """Ketma-ket ikki sutka turgan kursantlar (fayl + bazadagi tasdiqlangan sutkalar).

    Qog'ozdagi haqiqiy ma'lumot bo'lgani uchun import to'xtatilmaydi — faqat ogohlantirish
    (CLAUDE.md 10-bo'lim, 21-band).
    """
    days: dict[int, set[dt.date]] = defaultdict(set)
    names: dict[int, str] = {}
    lines: dict[tuple[int, dt.date], int] = {}
    for day, pairs in plan.items():
        for _, row in pairs:
            days[row.cadet.pk].add(day)
            names[row.cadet.pk] = f"{row.cadet.short_name} ({row.cadet.group.code})"
            lines[(row.cadet.pk, day)] = row.line
    if plan:
        around = {d + dt.timedelta(days=k) for d in plan for k in (-1, 1)}
        for cadet_id, day in Assignment.objects.filter(
            cadet_id__in=list(days), duty_day__status=DutyDay.Status.APPROVED, duty_day__date__in=around
        ).values_list("cadet_id", "duty_day__date"):
            days[cadet_id].add(day)
    out = []
    for cadet_id, dates in days.items():
        for day in sorted(dates):
            prev = day - dt.timedelta(days=1)
            if prev in dates and ((cadet_id, day) in lines or (cadet_id, prev) in lines):
                line = lines.get((cadet_id, day)) or lines.get((cadet_id, prev))
                out.append(
                    f"{line}-qator: {names[cadet_id]} ketma-ket ikki sutka turgan — {prev:%d.%m.%Y} va {day:%d.%m.%Y} "
                    "(qog‘ozdagi ma’lumot, import to‘xtatilmadi)"
                )
    return sorted(out, key=lambda s: int(s.split("-", 1)[0]))


def _write(plan, report: ImportReport, user, filename: str) -> None:
    now = timezone.now()
    for day, pairs in sorted(plan.items()):
        start = week_start(day)
        week = DutyWeek.objects.filter(start_date=start).first()
        if week is None:
            week = DutyWeek(start_date=start, note="qog‘ozdan import paytida yaratildi", created_by=user)
            try:
                week.full_clean()
            except Exception as exc:  # ValidationError: ustma-ust hafta va h.k.
                messages = getattr(exc, "messages", [str(exc)])
                report.errors.append(f"{day:%d.%m.%Y}: hafta yaratib bo‘lmadi — {'; '.join(messages)}")
                continue
            week.save()
            report.weeks_created.append(start)
        duty_day = DutyDay(
            week=week,
            date=day,
            status=DutyDay.Status.APPROVED,
            approved_by=user,
            approved_at=now,
            warnings=[],
        )
        duty_day.full_clean()
        duty_day.save()
        group_posts = {}
        for slot, row in pairs:
            if slot.post.gender == Gender.MALE:
                group_posts[slot.post_id] = row.cadet.group_id
        GroupPostAssignment.objects.bulk_create(
            GroupPostAssignment(duty_day=duty_day, post_id=post_id, group_id=group_id)
            for post_id, group_id in group_posts.items()
        )
        Assignment.objects.bulk_create(
            Assignment(
                duty_day=duty_day,
                post_slot=slot,
                cadet=row.cadet,
                group_id=row.cadet.group_id,
                source=Assignment.Source.MANUAL,
            )
            for slot, row in pairs
        )
        ChangeLog.objects.create(
            duty_day=duty_day,
            action=ChangeLog.Action.IMPORT,
            reason=f"qog‘ozdan import ({filename})",
            user=user,
        )
