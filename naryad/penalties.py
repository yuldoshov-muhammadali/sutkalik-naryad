"""Zalyotchiklar — navbatdan tashqari jazo naryadi (15-bosqich, CLAUDE.md 10-bo'lim, 61-band).

- Admin zalyot kiritadi (kursant, sabab turi, izoh, tanbeh sanasi, 1–10 ta naryad); bir kursantda bitta faol
  zalyot — yangisi qo'shilsa soni mavjudiga qo'shiladi (jurnalda alohida yozuv).
- Faqat o'g'il, faol va naryadga chiqadigan (lavozimda bo'lmagan) kursant zalyotchik bo'ladi.
- Jazo naryadi “o'tildi” deb faqat sutka TASDIQLANGANDA hisoblanadi; tasdiq bekor qilinsa hisob qaytadi.
  Soni tugagach zalyot avtomatik “bajarildi”.
- Algoritmga faol zalyotchiklar `engine.Offender` sifatida beriladi (naryad/engine/allocator.py: place_offenders).
"""

from __future__ import annotations

import datetime as dt
from collections.abc import Iterable
from dataclasses import dataclass, field

from django import forms
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, Q
from django.utils import timezone, translation
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy as _lazy

from naryad.engine import models as em

from .cadets import INPUT_CLASS
from .models import Assignment, Cadet, ChangeLog, DutyDay, Gender, Penalty, PenaltyEntry
from .services import ServiceError


def offender_error(cadet: Cadet) -> str | None:
    """Standart qoidalar: faqat o'g'il, faol, lavozimda bo'lmagan kursant zalyotchik bo'ladi."""
    if cadet.gender != Gender.MALE:
        return _("Zalyotchik faqat o‘g‘il kursant bo‘ladi (jazo o‘rinlari — 1-NO‘J va 1-Yotoqxona o‘g‘illar joyi).")
    if not cadet.is_active:
        return _("%(name)s faol emas — zalyotchik qilib bo‘lmaydi.") % {"name": cadet.short_name}
    if cadet.duty_exempt:
        return _("%(name)s — %(position)s (lavozimda) — avval lavozimni olib tashlang.") % {
            "name": cadet.short_name, "position": cadet.duty_position,
        }
    return None


def with_served(qs):
    """`served_count` — tasdiqlangan sutkalardagi jazo naryadlari soni (N+1 bo'lmasin)."""
    return qs.annotate(
        served_count=Count("assignments", filter=Q(assignments__duty_day__status=DutyDay.Status.APPROVED))
    )


def active_penalty(cadet: Cadet) -> Penalty | None:
    return with_served(Penalty.objects.filter(cadet=cadet, status=Penalty.Status.ACTIVE)).first()


def _reason_uz(reason_type: str) -> str:
    with translation.override("uz"):  # jurnal matni doim o'zbekcha
        return str(Penalty.Reason(reason_type).label)


class DuplicateEntry(ServiceError):
    """Bir xil kursant + sana + sabab 1 daqiqa ichida qayta yuborildi (ikki marta bosish) — yangi yozuv yaratilmaydi."""


DEDUP_SECONDS = 60


def commander_error(cadet: Cadet, order_date: dt.date, count: int, user) -> str | None:
    """Guruh komandiri uchun qo'shimcha qoidalar (20-bosqich, CLAUDE.md 68-band). Serverda tekshiriladi."""
    from .models import Setting

    setting = Setting.load()
    own = getattr(user, "cadet", None)
    if own is None or cadet.group_id != own.group_id:
        raise PermissionDenied(_("Faqat o‘z guruhingiz kursantiga zalyot kiritasiz."))
    if count > setting.commander_penalty_max:
        return _("Komandir bir kiritishda ko‘pi bilan %(n)s ta jazo naryadi kirita oladi.") % {"n": setting.commander_penalty_max}
    today = timezone.localdate()
    if order_date < today - dt.timedelta(days=setting.commander_order_max_age_days):
        return _("Tanbeh sanasi %(n)s kundan eski bo‘lmasin.") % {"n": setting.commander_order_max_age_days}
    start = timezone.make_aware(dt.datetime.combine(today, dt.time.min))
    if PenaltyEntry.objects.filter(created_by=user, created_at__gte=start).count() >= setting.commander_daily_limit:
        return _("Bir kunda ko‘pi bilan %(n)s ta zalyot kiritasiz.") % {"n": setting.commander_daily_limit}
    return None


@transaction.atomic
def add(cadet: Cadet, reason_type: str, note: str, order_date: dt.date, count: int, user) -> Penalty:
    """Zalyot kiritadi yoki faol zalyotga qo'shadi. Admin yoki (o'z guruhi uchun) guruh komandiri."""
    cadet = Cadet.objects.select_for_update().select_related("group").get(pk=cadet.pk)
    error = offender_error(cadet)
    if error:
        raise ServiceError(error)
    note = (note or "").strip()
    if not note:
        raise ServiceError(_("Izoh majburiy."))
    if reason_type not in Penalty.Reason.values:
        raise ServiceError(_("Sabab turi tanlanmagan."))
    if not (1 <= int(count) <= Penalty.MAX_ADD):
        raise ServiceError(_("Jazo naryadlari soni 1 dan 10 gacha bo‘lishi kerak."))
    count = int(count)
    if order_date > timezone.localdate():
        raise ServiceError(_("Tanbeh sanasi kelajakda bo‘lmasin."))
    by_commander = getattr(user, "role", None) == "commander"
    if by_commander:
        error = commander_error(cadet, order_date, count, user)
        if error:
            raise ServiceError(error)
    recent = timezone.now() - dt.timedelta(seconds=DEDUP_SECONDS)
    if PenaltyEntry.objects.filter(cadet=cadet, order_date=order_date, reason_type=reason_type, created_at__gte=recent).exists():
        raise DuplicateEntry(_("Bu zalyot hozirgina kiritildi — qayta yaratilmadi."))
    penalty = Penalty.objects.select_for_update().filter(cadet=cadet, status=Penalty.Status.ACTIVE).first()
    if penalty is None:
        penalty = Penalty.objects.create(
            cadet=cadet, reason_type=reason_type, note=note, order_date=order_date, count=count, created_by=user,
            reviewed_at=None if by_commander else timezone.now(), reviewed_by=None if by_commander else user,
        )
        details = f"Yangi zalyot: {count} ta jazo naryadi (tanbeh {order_date:%d.%m.%Y})."
    else:
        penalty.count += count
        fields = ["count"]
        if by_commander:  # komandir qo'shgan — admin yana ko'rib chiqadi
            penalty.reviewed_at, penalty.reviewed_by = None, None
            fields += ["reviewed_at", "reviewed_by"]
        penalty.save(update_fields=fields)
        details = f"Faol zalyotga qo‘shildi: +{count} ta (jami {penalty.count} ta; tanbeh {order_date:%d.%m.%Y})."
    if by_commander:
        details += " Kiritgan: guruh komandiri."
    PenaltyEntry.objects.create(
        penalty=penalty, cadet=cadet, reason_type=reason_type, note=note, order_date=order_date, count=count,
        created_by=user, by_commander=by_commander,
    )
    ChangeLog.objects.create(
        action=ChangeLog.Action.PENALTY_ADD, cadet=cadet, user=user, reason=f"{_reason_uz(reason_type)}: {note}", details=details
    )
    return penalty


@transaction.atomic
def edit(penalty: Penalty, reason_type: str, note: str, order_date: dt.date, count: int, user) -> Penalty:
    """Admin tahriri (kursant o'zgarmaydi). Soni allaqachon o'tilganidan kam bo'lmaydi; jurnalga eski → yangi."""
    # PostgreSQL: FOR UPDATE va GROUP BY (annotate) birga bo'lmaydi — avval qulf, keyin o'tilganlar alohida
    Penalty.objects.select_for_update().filter(pk=penalty.pk).values_list("pk", flat=True).get()
    penalty = with_served(Penalty.objects.filter(pk=penalty.pk)).get()
    if not penalty.is_active:
        raise ServiceError(_("Faqat faol zalyot tahrirlanadi."))
    note = (note or "").strip()
    if not note:
        raise ServiceError(_("Izoh majburiy."))
    if reason_type not in Penalty.Reason.values:
        raise ServiceError(_("Sabab turi tanlanmagan."))
    count = int(count)
    if count < 1:
        raise ServiceError(_("Jazo naryadlari soni kamida 1 ta."))
    if count < penalty.served:
        raise ServiceError(_("Jazo naryadlari soni allaqachon o‘tilganidan (%(n)s ta) kam bo‘lmaydi.") % {"n": penalty.served})
    if order_date > timezone.localdate():
        raise ServiceError(_("Tanbeh sanasi kelajakda bo‘lmasin."))
    changes = []
    for name, new, label in (
        ("reason_type", reason_type, "sabab"), ("note", note, "izoh"), ("order_date", order_date, "tanbeh sanasi"), ("count", count, "soni"),
    ):
        old = getattr(penalty, name)
        if old != new:
            if name == "reason_type":
                changes.append(f"{label}: {_reason_uz(old)} → {_reason_uz(new)}")
            elif name == "order_date":
                changes.append(f"{label}: {old:%d.%m.%Y} → {new:%d.%m.%Y}")
            else:
                changes.append(f"{label}: {old} → {new}")
            setattr(penalty, name, new)
    if not changes:
        return penalty
    penalty.save(update_fields=["reason_type", "note", "order_date", "count"])
    ChangeLog.objects.create(action=ChangeLog.Action.PENALTY_EDIT, cadet=penalty.cadet, user=user, details="; ".join(changes))
    if penalty.served >= penalty.count:
        sync([penalty.pk], user)
    return penalty


@transaction.atomic
def review(penalties, user) -> int:
    """“Ko'rib chiqildi”: faqat “Yangi” belgisini olib tashlaydi, zalyotni o'zgartirmaydi."""
    n = 0
    for penalty in penalties.filter(reviewed_at__isnull=True).select_related("cadet"):
        penalty.reviewed_at, penalty.reviewed_by = timezone.now(), user
        penalty.save(update_fields=["reviewed_at", "reviewed_by"])
        ChangeLog.objects.create(action=ChangeLog.Action.PENALTY_REVIEW, cadet=penalty.cadet, user=user,
                                 details="Komandir kiritgan zalyot ko‘rib chiqildi.")
        n += 1
    return n


def unreviewed():
    return Penalty.objects.filter(reviewed_at__isnull=True).exclude(status=Penalty.Status.CANCELLED)


@transaction.atomic
def cancel(penalty: Penalty, reason: str, user) -> Penalty:
    penalty = Penalty.objects.select_for_update().get(pk=penalty.pk)
    if not penalty.is_active:
        raise ServiceError(_("Zalyot faol emas — bekor qilib bo‘lmaydi."))
    reason = (reason or "").strip()
    if not reason:
        raise ServiceError(_("Bekor qilish sababi majburiy."))
    penalty.status = Penalty.Status.CANCELLED
    penalty.cancel_reason = reason
    penalty.finished_at = timezone.now()
    penalty.save(update_fields=["status", "cancel_reason", "finished_at"])
    ChangeLog.objects.create(
        action=ChangeLog.Action.PENALTY_CANCEL, cadet=penalty.cadet, user=user, reason=reason,
        details=f"Zalyot bekor qilindi: {penalty.served} / {penalty.count} ta o‘tilgan edi.",
    )
    return penalty


def sync(penalty_ids: Iterable[int], user=None) -> None:
    """Tasdiqlash / tasdiqni bekor qilishdan keyin: soni tugasa — “bajarildi”, hisob qaytsa — yana “faol”."""
    for penalty in with_served(Penalty.objects.filter(pk__in=set(penalty_ids)).select_related("cadet")):
        if penalty.status == Penalty.Status.ACTIVE and penalty.served >= penalty.count:
            penalty.status = Penalty.Status.DONE
            penalty.finished_at = timezone.now()
            penalty.save(update_fields=["status", "finished_at"])
            ChangeLog.objects.create(
                action=ChangeLog.Action.PENALTY_DONE, cadet=penalty.cadet, user=user,
                details=f"Barcha {penalty.count} ta jazo naryadi o‘tildi — zalyot bajarildi.",
            )
        elif penalty.status == Penalty.Status.DONE and penalty.served < penalty.count:
            other = Penalty.objects.filter(cadet=penalty.cadet, status=Penalty.Status.ACTIVE).exclude(pk=penalty.pk).first()
            if other is not None:  # shu orada yangi faol zalyot ochilgan — qolgan naryad unga qo'shiladi
                other.count += penalty.count - penalty.served
                other.save(update_fields=["count"])
                penalty.count = penalty.served
                penalty.save(update_fields=["count"])
                details = f"Tasdiq bekor qilindi: qolgan jazo naryadi faol zalyotga qo‘shildi (jami {other.count} ta)."
            else:
                penalty.status = Penalty.Status.ACTIVE
                penalty.finished_at = None
                penalty.save(update_fields=["status", "finished_at"])
                details = f"Tasdiq bekor qilindi: {penalty.served} / {penalty.count} ta o‘tilgan — zalyot yana faol."
            ChangeLog.objects.create(action=ChangeLog.Action.PENALTY_RESUME, cadet=penalty.cadet, user=user, details=details)


def day_penalty_ids(duty_day: DutyDay) -> list[int]:
    return list(duty_day.assignments.filter(kind=Assignment.Kind.PENALTY).values_list("penalty_id", flat=True))


def engine_offenders(day: dt.date) -> list[em.Offender]:
    """Faol zalyotchiklar algoritm uchun: qolgan soni va oxirgi jazo joyi `day` dan oldingi tasdiqlanganlar bo'yicha."""
    penalties = list(Penalty.objects.filter(status=Penalty.Status.ACTIVE).order_by("pk"))
    if not penalties:
        return []
    approved = Assignment.objects.filter(
        kind=Assignment.Kind.PENALTY, duty_day__status=DutyDay.Status.APPROVED, duty_day__date__lt=day
    )
    served: dict[int, int] = {}
    for pid in approved.filter(penalty__in=penalties).values_list("penalty_id", flat=True):
        served[pid] = served.get(pid, 0) + 1
    last_post: dict[int, str] = {}
    cadet_ids = [p.cadet_id for p in penalties]
    for cid, code in (
        approved.filter(cadet_id__in=cadet_ids).order_by("duty_day__date").values_list("cadet_id", "post_slot__post__code")
    ):
        last_post[cid] = code  # oxirgisi qoladi
    return [
        em.Offender(cadet_id=p.cadet_id, since=p.order_date, remaining=p.count - served.get(p.pk, 0), last_post=last_post.get(p.cadet_id))
        for p in penalties
        if p.count - served.get(p.pk, 0) > 0
    ]


def active_offender_ids() -> set[int]:
    return set(Penalty.objects.filter(status=Penalty.Status.ACTIVE).values_list("cadet_id", flat=True))


# ---------------------------------------------------------------------------
# Forma (Zalyotchiklar sahifasi)
# ---------------------------------------------------------------------------


class PenaltyForm(forms.Form):
    cadet = forms.ModelChoiceField(queryset=Cadet.objects.none(), label=_lazy("Kursant"))
    reason_type = forms.ChoiceField(choices=Penalty.Reason.choices, label=_lazy("Sabab turi"))
    order_date = forms.DateField(
        label=_lazy("Buyruq/tanbeh sanasi"), widget=forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d")
    )
    count = forms.IntegerField(label=_lazy("Jazo naryadlari soni"), min_value=1, max_value=Penalty.MAX_ADD, initial=1)
    note = forms.CharField(label=_lazy("Izoh"), max_length=500, widget=forms.TextInput())

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        # Standart qoidalar: faqat o'g'il, faol, lavozimda bo'lmagan kursantlar
        self.fields["cadet"].queryset = Cadet.objects.filter(
            gender=Gender.MALE, is_active=True, duty_exempt=False
        ).select_related("group").order_by("group__display_order", "last_name", "first_name")
        self.fields["cadet"].label_from_instance = lambda c: f"{c.group.code} · {c.full_name}"
        self.fields["cadet"].empty_label = _("— kursantni tanlang —")
        for f in self.fields.values():
            f.widget.attrs.setdefault("class", INPUT_CLASS)
        self.fields["count"].widget.attrs.update({"inputmode": "numeric"})


class CancelForm(forms.Form):
    reason = forms.CharField(label=_lazy("Bekor qilish sababi"), max_length=500)


# ---------------------------------------------------------------------------
# Zalyot hisobi (20-bosqich, CLAUDE.md 68-band)
# ---------------------------------------------------------------------------


@dataclass
class LedgerRow:
    key: object
    offenders: set = field(default_factory=set)
    entries: int = 0
    added: int = 0
    served: int = 0
    remaining: int = 0


def ledger(date_from: dt.date | None = None, date_to: dt.date | None = None, group=None) -> dict:
    """Hisob: kiritilgan — tanbeh sanasi davrga tushgan kiritishlar; o'tildi — davrdagi TASDIQLANGAN jazo naryadlari;
    qoldi — hozirgi faol zalyotlarning qolgani (davrga bog'liq emas). `group` — faqat shu guruh (komandir uchun)."""
    entries = PenaltyEntry.objects.select_related("cadet__group", "created_by")
    served = Assignment.objects.filter(kind=Assignment.Kind.PENALTY, duty_day__status=DutyDay.Status.APPROVED)
    if date_from:
        entries = entries.filter(order_date__gte=date_from)
        served = served.filter(duty_day__date__gte=date_from)
    if date_to:
        entries = entries.filter(order_date__lte=date_to)
        served = served.filter(duty_day__date__lte=date_to)
    active = with_served(Penalty.objects.filter(status=Penalty.Status.ACTIVE)).select_related("cadet__group")
    if group is not None:
        entries, served, active = entries.filter(cadet__group=group), served.filter(cadet__group=group), active.filter(cadet__group=group)

    groups: dict = {}
    cadets: dict = {}
    users: dict = {}

    def g(cadet):
        return groups.setdefault(cadet.group_id, LedgerRow(cadet.group))

    def c(cadet):
        return cadets.setdefault(cadet.pk, LedgerRow(cadet))

    for e in entries:
        for row in (g(e.cadet), c(e.cadet)):
            row.offenders.add(e.cadet_id)
            row.entries += 1
            row.added += e.count
        u = users.setdefault(e.created_by_id, LedgerRow(e.created_by))
        u.entries += 1
        u.added += e.count
        u.offenders.add(e.cadet_id)
    for a in served.select_related("cadet__group"):
        for row in (g(a.cadet), c(a.cadet)):
            row.served += 1
    for p in active:
        for row in (g(p.cadet), c(p.cadet)):
            row.remaining += p.remaining
            row.offenders.add(p.cadet_id)
    group_rows = sorted(groups.values(), key=lambda r: r.key.display_order)
    cadet_rows = sorted(cadets.values(), key=lambda r: (r.key.group.display_order, r.key.last_name, r.key.first_name))
    user_rows = sorted(users.values(), key=lambda r: -r.added)
    totals = LedgerRow("total", offenders={x for r in group_rows for x in r.offenders},
                       entries=sum(r.entries for r in group_rows), added=sum(r.added for r in group_rows),
                       served=sum(r.served for r in group_rows), remaining=sum(r.remaining for r in group_rows))
    return {"groups": group_rows, "cadets": cadet_rows, "users": user_rows, "totals": totals}


def user_label(user) -> str:
    if user is None:
        return "—"
    name = user.get_full_name() or user.username
    role = {"admin": "admin", "kadir": "kadir", "commander": "guruh komandiri"}.get(user.role, user.role)
    return f"{name} ({role})"


def ledger_xlsx(data: dict, date_from=None, date_to=None) -> bytes:
    """Excel (ustun nomlari o'zbekcha): Guruhlar, Kursantlar, Kim kiritdi."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    period = f"{date_from:%d.%m.%Y} – {date_to:%d.%m.%Y}" if date_from and date_to else (
        f"{date_from:%d.%m.%Y} dan" if date_from else (f"{date_to:%d.%m.%Y} gacha" if date_to else "butun tarix"))

    def sheet(ws, title, header, rows):
        ws.title = title
        ws.append([f"Zalyot hisobi — {period}"])
        ws["A1"].font = Font(bold=True)
        ws.append(header)
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, col in enumerate("ABCDEFG"):
            ws.column_dimensions[col].width = 40 if i == 0 else 16

    sheet(wb.active, "Guruhlar", ["Guruh", "Zalyotchilar soni", "Kiritilgan jazo naryadlari", "O‘tildi", "Qoldi"],
          [[r.key.code, len(r.offenders), r.added, r.served, r.remaining] for r in data["groups"]]
          + [["Jami", len(data["totals"].offenders), data["totals"].added, data["totals"].served, data["totals"].remaining]])
    sheet(wb.create_sheet(), "Kursantlar", ["F.I.Sh.", "Guruh", "Zalyotlar (kiritishlar)", "Kiritilgan jazo naryadlari", "O‘tildi", "Qoldi"],
          [[r.key.full_name, r.key.group.code, r.entries, r.added, r.served, r.remaining] for r in data["cadets"]])
    sheet(wb.create_sheet(), "Kim kiritdi", ["Kim kiritdi", "Zalyotlar (kiritishlar)", "Jazo naryadlari", "Kursantlar"],
          [[user_label(r.key), r.entries, r.added, len(r.offenders)] for r in data["users"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class EditForm(forms.Form):
    reason_type = forms.ChoiceField(choices=Penalty.Reason.choices, label=_lazy("Sabab turi"))
    order_date = forms.DateField(label=_lazy("Buyruq/tanbeh sanasi"), widget=forms.DateInput(attrs={"type": "date"}, format="%Y-%m-%d"))
    count = forms.IntegerField(label=_lazy("Jazo naryadlari soni (jami)"), min_value=1, max_value=99)
    note = forms.CharField(label=_lazy("Izoh"), max_length=500)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for f in self.fields.values():
            f.widget.attrs.setdefault("class", INPUT_CLASS)


class CommanderPenaltyForm(PenaltyForm):
    """Komandir formasi: faqat o'z guruhining o'g'il, faol, naryadga chiqadigan kursantlari; soni Sozlamalardagi chegaragacha."""

    def __init__(self, group, *args, **kwargs):
        from .models import Setting

        super().__init__(*args, **kwargs)
        self.fields["cadet"].queryset = self.fields["cadet"].queryset.filter(group=group)
        self.fields["cadet"].label_from_instance = lambda c: c.full_name
        limit = Setting.load().commander_penalty_max
        self.fields["count"] = forms.IntegerField(label=_lazy("Jazo naryadlari soni"), min_value=1, max_value=limit, initial=1)
        self.fields["count"].widget.attrs.update({"class": INPUT_CLASS, "inputmode": "numeric"})
        self.fields["order_date"].widget.attrs["max"] = f"{timezone.localdate():%Y-%m-%d}"
