"""Naryad tarixi, statistika va adolat ko'rsatkichi (9-bosqich).

Hammasi so'rov paytida faqat tasdiqlangan sutkalardagi Assignment dan hisoblanadi
(CLAUDE.md 8-bo'lim: statistika alohida saqlanmaydi). Adolat har guruh kesimida
va qizlar uchun alohida (CLAUDE.md 4-B.5). Adolat, jadvallar va guruhlar kesimi faqat ODDIY
naryadlar bo'yicha; navbatdan tashqari (jazo) naryadlari alohida ustunda (15-bosqich, 61-band).
"""

from __future__ import annotations

import datetime as dt
from collections import Counter, defaultdict
from dataclasses import dataclass, field

from django.db.models import Q
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy as _lazy

from naryad.engine.allocator import week_start
from naryad.engine.config import ROTATION_CHAIN

from .dashboard import current_duty_date
from .models import (
    Assignment,
    Cadet,
    ChangeLog,
    DutyDay,
    DutyWeek,
    Gender,
    GroupPostAssignment,
    Penalty,
    Post,
    Setting,
    SlotRole,
    StudyGroup,
)

ROLES = [SlotRole.DUTY, SlotRole.ASSISTANT, SlotRole.GUARD]
ROLE_SHORT = {SlotRole.DUTY: _lazy("Navbatchi"), SlotRole.ASSISTANT: _lazy("Yordamchi"), SlotRole.GUARD: _lazy("Posbon")}
REASONS = [("kasal", _lazy("Kasal")), ("tatil", _lazy("Ta’til")), ("sababli", _lazy("Sababli")), ("boshqa", _lazy("Boshqa"))]
# Eski yozuvlar (reason_code yo'q): ChangeLog.reason bazada doim o'zbekcha — tarjimasiz solishtiriladi
REASON_PREFIXES = [("kasal", "Kasal"), ("tatil", "Ta’til"), ("sababli", "Sababli")]
GIRLS = "qizlar"
POST_SHORT = {
    "markaz": _lazy("Markaz"),
    "noj1": _lazy("1-NO‘J"),
    "noj2": _lazy("2-NO‘J"),
    "yotoq1": _lazy("1-Yotoqxona"),
    "oquv": _lazy("O‘quv binosi"),
    "patrul": _lazy("Patrul"),
}


# ---------------------------------------------------------------------------
# Davr
# ---------------------------------------------------------------------------


@dataclass
class Period:
    mode: str = "all"  # all | range | week
    start: dt.date | None = None
    end: dt.date | None = None
    week: DutyWeek | None = None

    @property
    def label(self) -> str:
        if self.mode == "week" and self.week:
            return _("2-kurs haftasi: %(start)s – %(end)s") % {
                "start": f"{self.week.start_date:%d.%m.%Y}", "end": f"{self.week.end_date:%d.%m.%Y}",
            }
        if self.mode == "range":
            a = f"{self.start:%d.%m.%Y}" if self.start else _("boshidan")
            b = f"{self.end:%d.%m.%Y}" if self.end else _("hozirgacha")
            return f"{a} – {b}"
        return _("Butun tarix")

    def filter(self, qs, field_name: str = "duty_day__date"):
        if self.start:
            qs = qs.filter(**{f"{field_name}__gte": self.start})
        if self.end:
            qs = qs.filter(**{f"{field_name}__lte": self.end})
        return qs


def parse_date(value) -> dt.date | None:
    try:
        return dt.date.fromisoformat(value) if value else None
    except (TypeError, ValueError):
        return None


def period_from_params(params) -> Period:
    mode = params.get("davr", "all")
    if mode == "week":
        start = parse_date(params.get("hafta"))
        week = DutyWeek.objects.filter(start_date=start).first() if start else None
        if week:
            return Period("week", week.start_date, week.end_date, week)
    if mode == "range":
        start, end = parse_date(params.get("dan")), parse_date(params.get("gacha"))
        if start and end and end < start:
            start, end = end, start
        return Period("range", start, end)
    return Period()


def approved_assignments(period: Period):
    """Oddiy naryadlar (jazo naryadlari adolat hisobiga kirmaydi)."""
    qs = Assignment.objects.filter(duty_day__status=DutyDay.Status.APPROVED, kind=Assignment.Kind.NORMAL)
    return period.filter(qs).select_related("duty_day", "post_slot__post", "cadet__group", "group")


def reference_week(period: Period) -> dt.date | None:
    """“Joriy haftadagi soni” uchun hafta: tanlangan hafta yoki hozirgi/oxirgi 2-kurs haftasi."""
    if period.week:
        return period.week.start_date
    today = current_duty_date()
    week = DutyWeek.objects.filter(start_date__lte=today).order_by("-start_date").first()
    return week.start_date if week else None


# ---------------------------------------------------------------------------
# Kursant × joy / lavozim jadvali
# ---------------------------------------------------------------------------


@dataclass
class CadetRow:
    cadet: Cadet
    by_post: dict[str, int] = field(default_factory=dict)
    by_role: dict[str, int] = field(default_factory=dict)
    total: int = 0
    this_week: int = 0
    last_date: dt.date | None = None
    dates_posts: list[tuple[dt.date, str]] = field(default_factory=list)
    dates_roles: list[tuple[dt.date, str]] = field(default_factory=list)
    penalty: int = 0  # navbatdan tashqari (jazo) naryadlari — yuqoridagi sonlarga kirmaydi
    hard: int = 0  # og'ir joylarda (21-bosqich, Sozlamalardagi ro'yxat) — faqat oddiy naryadlar


def hard_posts() -> set[str]:
    """Og'ir joylar (Sozlamalar; standart 1-NO‘J va 1-Yotoqxona). Qizlar yotoqxonasi kirmaydi."""
    return set(Setting.load().hard_posts or ()) & set(ROTATION_CHAIN)


def pool_cadets(pool: str):
    """Hovuz: guruh kodi (o'g'il bolalar) yoki 'qizlar' (13 qiz)."""
    qs = Cadet.objects.select_related("group")
    if pool == GIRLS:
        return qs.filter(gender=Gender.FEMALE)
    return qs.filter(gender=Gender.MALE, group__code=pool)


def pool_of(cadet: Cadet) -> str:
    return GIRLS if cadet.gender == Gender.FEMALE else cadet.group.code


def cadet_rows(pool: str, period: Period) -> list[CadetRow]:
    return build_rows(list(pool_cadets(pool).order_by("last_name", "first_name")), period)


def rows_by_pool(period: Period) -> dict[str, list[CadetRow]]:
    """Barcha hovuzlar bitta so'rovda (bosh sahifa, statistika, Excel — N+1 bo'lmasin)."""
    rows = build_rows(list(Cadet.objects.select_related("group").order_by("last_name", "first_name")), period)
    out: dict[str, list[CadetRow]] = {code: [] for code, _title in pools()}
    for r in rows:
        out.setdefault(pool_of(r.cadet), []).append(r)
    return out


def build_rows(cadets: list[Cadet], period: Period) -> list[CadetRow]:
    rows = {c.pk: CadetRow(cadet=c) for c in cadets}
    ref_week = reference_week(period)
    hard = hard_posts()
    qs = Assignment.objects.filter(duty_day__status=DutyDay.Status.APPROVED, cadet_id__in=rows.keys())
    for a in period.filter(qs).select_related("duty_day", "post_slot__post").order_by("duty_day__date"):
        row = rows[a.cadet_id]
        if a.kind == Assignment.Kind.PENALTY:
            row.penalty += 1
            continue
        post = a.post_slot.post.code
        row.by_post[post] = row.by_post.get(post, 0) + 1
        row.by_role[a.post_slot.role] = row.by_role.get(a.post_slot.role, 0) + 1
        row.total += 1
        row.hard += 1 if post in hard else 0
        if ref_week and week_start(a.duty_day.date) == ref_week:
            row.this_week += 1
        row.last_date = a.duty_day.date
        row.dates_posts.append((a.duty_day.date, post))
        row.dates_roles.append((a.duty_day.date, a.post_slot.role))
    # faol emas va bu davrda naryadi yo'q kursantlar jadvalga kirmaydi
    return [r for r in rows.values() if r.cadet.is_active or r.total or r.penalty]


def heat(value: int, maximum: int, minimum: int = 0) -> str:
    """Issiqlik xaritasi: ko'p — to'q, kam — och (inline style, Tailwind'siz).
    `minimum` berilsa (jami ustuni), rang min…max oralig'ida — kichik farq ham ko'rinadi."""
    if not value or not maximum:
        return ""
    span = maximum - minimum
    alpha = 0.12 + 0.78 * (max(0, value - minimum) / span if span > 0 else 0.5)
    color = "#fff" if alpha > 0.55 else "#0f172a"
    return f"background-color: rgba(33, 59, 120, {alpha:.2f}); color: {color};"


# ---------------------------------------------------------------------------
# Adolat
# ---------------------------------------------------------------------------


def fairness_status(diff: int | None) -> str:
    if diff is None:
        return "none"
    if diff <= 1:
        return "green"
    if diff == 2:
        return "yellow"
    return "red"


STATUS_LABELS = {
    "green": _lazy("Adolatli (farq ≤ 1)"), "yellow": _lazy("Diqqat (farq = 2)"),
    "red": _lazy("Adolatsiz (farq ≥ 3)"), "none": _lazy("Ma’lumot yo‘q"),
}


@dataclass
class Fairness:
    pool: str
    title: str
    members: int
    total: int
    min: int | None
    max: int | None
    diff: int | None
    mean: float | None
    within_one_pct: float | None
    repeats: list[dict]  # ketma-ket ikki naryadi bir joyda
    status: str
    repeats_applicable: bool = True  # qizlarda joy bitta — tekshirilmaydi
    duty_repeats: list[dict] = field(default_factory=list)  # qizlar: ketma-ket ikki naryadda navbatchi
    hard_diff: int | None = None  # og'ir joylar farqi (max − min) guruh ichida (21-bosqich)

    @property
    def status_label(self) -> str:
        return str(STATUS_LABELS[self.status])  # joriy tilda (Excel — override("uz") ichida)


def fairness_from_rows(pool: str, title: str, rows: list[CadetRow]) -> Fairness:
    # Adolat faqat naryadga chiqadiganlar orasida: faol emaslar va lavozimdagilar hisobga olinmaydi
    active = [r for r in rows if r.cadet.on_duty]
    totals = [r.total for r in active]
    applicable = pool != GIRLS
    repeats = []
    for r in rows if applicable else []:
        for (d1, p1), (d2, p2) in zip(r.dates_posts, r.dates_posts[1:]):
            if p1 == p2:
                repeats.append({"cadet": r.cadet, "post": p1, "post_name": str(POST_SHORT.get(p1, p1)), "first": d1, "second": d2})
    duty_repeats = consecutive_duty(rows) if pool == GIRLS else []
    if not totals or sum(totals) == 0:
        return Fairness(pool, title, len(active), 0, None, None, None, None, None, repeats, "none", applicable, duty_repeats)
    mean = sum(totals) / len(totals)
    lo, hi = min(totals), max(totals)
    within = sum(1 for t in totals if abs(t - mean) <= 1)
    hard = [r.hard for r in active]
    return Fairness(
        pool, title, len(active), sum(totals), lo, hi, hi - lo, round(mean, 2),
        round(100 * within / len(totals), 1), repeats, fairness_status(hi - lo), applicable, duty_repeats,
        hard_diff=max(hard) - min(hard) if applicable else None,
    )


def consecutive_duty(rows: list[CadetRow]) -> list[dict]:
    """Qizlar: kursantning ketma-ket ikki naryadida ham navbatchi bo'lgan holatlar (CLAUDE.md 10-bo'lim, 50-band)."""
    out = []
    for r in rows:
        for (d1, r1), (d2, r2) in zip(r.dates_roles, r.dates_roles[1:]):
            if r1 == r2 == SlotRole.DUTY:
                out.append({"cadet": r.cadet, "first": d1, "second": d2})
    return out


def pools() -> list[tuple[str, str]]:
    groups = [(g.code, g.code) for g in StudyGroup.objects.filter(is_active=True).order_by("display_order")]
    return groups + [(GIRLS, _("Qizlar (13)"))]


def fairness_all(period: Period, by_pool: dict[str, list[CadetRow]] | None = None) -> list[Fairness]:
    by_pool = by_pool if by_pool is not None else rows_by_pool(period)
    return [fairness_from_rows(code, title, by_pool.get(code, [])) for code, title in pools()]


# ---------------------------------------------------------------------------
# Guruhlar kesimi
# ---------------------------------------------------------------------------


@dataclass
class GroupSection:
    group: StudyGroup
    by_post: dict[str, int]
    days: int
    total: int
    boys: int
    per_boy: float | None


def group_sections(period: Period) -> list[GroupSection]:
    gpa = period.filter(GroupPostAssignment.objects.filter(duty_day__status=DutyDay.Status.APPROVED))
    counts = Counter((g.group_id, g.post.code) for g in gpa.select_related("post"))
    totals = Counter(
        a["group_id"] for a in approved_assignments(period).filter(post_slot__post__gender=Gender.MALE).values("group_id")
    )
    boys_by_group = Counter(
        # bir o'g'ilga o'rtacha — faqat naryadga chiqadiganlarga bo'linadi (lavozimdagilarsiz)
        Cadet.objects.filter(gender=Gender.MALE, is_active=True, duty_exempt=False).values_list("group_id", flat=True)
    )
    out = []
    for g in StudyGroup.objects.filter(is_active=True).order_by("display_order"):
        by_post = {code: counts.get((g.pk, code), 0) for code in ROTATION_CHAIN}
        boys = boys_by_group.get(g.pk, 0)
        total = totals.get(g.pk, 0)
        out.append(GroupSection(g, by_post, sum(by_post.values()), total, boys, round(total / boys, 2) if boys else None))
    return out


# ---------------------------------------------------------------------------
# Qo'lda tuzatishlar nazorati
# ---------------------------------------------------------------------------


@dataclass
class ManualRow:
    cadet: Cadet
    pool: str
    placed: int = 0
    removed: int = 0
    pool_mean: float = 0.0
    alert: bool = False


def reason_code_of(entry: ChangeLog) -> str:
    if entry.reason_code:
        return entry.reason_code
    for code, label in REASON_PREFIXES:
        if entry.reason.startswith(label):
            return code
    return "boshqa"


def effective_replacements(period: Period) -> list[ChangeLog]:
    """Hisobga olinadigan almashtirishlar: faqat hozir tasdiqlangan sutkalarniki va shu sutka
    oxirgi marta (qayta) tuzilgandan keyingilari — qayta tuzishda yo'qolgan tuzatishlar sanalmaydi."""
    approved = set(period.filter(DutyDay.objects.filter(status=DutyDay.Status.APPROVED), "date").values_list("date", flat=True))
    last_generated: dict[dt.date, dt.datetime] = {}
    for day, created in ChangeLog.objects.filter(
        action__in=[ChangeLog.Action.GENERATE, ChangeLog.Action.IMPORT], day_date__in=approved
    ).values_list("day_date", "created_at"):
        if day not in last_generated or created > last_generated[day]:
            last_generated[day] = created
    entries = ChangeLog.objects.filter(action=ChangeLog.Action.REPLACE, day_date__in=approved).select_related(
        "user", "old_cadet__group", "new_cadet__group"
    )
    return [e for e in entries if e.day_date not in last_generated or e.created_at >= last_generated[e.day_date]]


def manual_control(period: Period) -> dict:
    entries = effective_replacements(period)
    by_reason = Counter()
    by_user = Counter()
    placed, removed = Counter(), Counter()
    for e in entries:
        by_reason[reason_code_of(e)] += 1
        by_user[(e.user.get_full_name() or e.user.username) if e.user else _("tizim")] += 1
        if e.new_cadet_id:
            placed[e.new_cadet_id] += 1
        if e.old_cadet_id:
            removed[e.old_cadet_id] += 1

    rows: list[ManualRow] = []
    members_by_pool: dict[str, list[Cadet]] = defaultdict(list)
    for c in Cadet.objects.select_related("group").filter(
        Q(is_active=True, duty_exempt=False) | Q(pk__in=list(placed) + list(removed))
    ):
        members_by_pool[pool_of(c)].append(c)
    for pool, _title in pools():
        members = members_by_pool.get(pool, [])
        if not members:
            continue
        mean = sum(placed[c.pk] for c in members) / len(members)
        for c in members:
            if placed[c.pk] or removed[c.pk]:
                row = ManualRow(c, pool, placed[c.pk], removed[c.pk], round(mean, 2))
                row.alert = is_manual_outlier(row.placed, mean)
                rows.append(row)
    rows.sort(key=lambda r: (-r.placed, -r.removed, r.cadet.last_name))
    return {
        "total": sum(by_reason.values()),
        "by_reason": [(str(label), by_reason.get(code, 0)) for code, label in REASONS],
        "by_user": sorted(by_user.items(), key=lambda kv: -kv[1]),
        "rows": rows,
        "alerts": [r for r in rows if r.alert],
    }


def is_manual_outlier(placed: int, pool_mean: float) -> bool:
    """Qo'lda qo'yilish guruh o'rtachasidan sezilarli ko'p: kamida 2 marta va o'rtachaning 2 baravaridan ko'p yoki teng."""
    return placed >= 2 and placed >= 2 * pool_mean


# ---------------------------------------------------------------------------
# Kursant uchun (o'z kartochkasi) — boshqalarning ismlarisiz
# ---------------------------------------------------------------------------


def cadet_comparison(cadet: Cadet) -> dict:
    pool = GIRLS if cadet.gender == Gender.FEMALE else cadet.group.code
    period = Period()
    rows = cadet_rows(pool, period)
    active = [r for r in rows if r.cadet.on_duty]  # guruh o'rtachasi — naryadga chiqadiganlar bo'yicha
    mean = sum(r.total for r in active) / len(active) if active else 0
    own = next((r.total for r in rows if r.cadet.pk == cadet.pk), 0)
    return {
        "exempt": cadet.duty_exempt,
        "own": own,
        "mean": round(mean, 1),
        "pool_title": _("qizlar") if pool == GIRLS else _("%(group)s guruhi") % {"group": pool},
        "members": len(active),
        "diff": round(own - mean, 1),
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def excel_export(period: Period) -> bytes:
    """Rasmiy hisobot: ustun nomlari va qiymatlar doim o'zbekcha (CLAUDE.md 59-band)."""
    from django.utils import translation

    with translation.override("uz"):
        return _excel_export(period)


def _excel_export(period: Period) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    posts = list(Post.objects.filter(gender=Gender.MALE).order_by("display_order"))

    def sheet(title: str, header: list[str], rows: list[list]):
        ws = wb.create_sheet(title[:31])
        ws.append([f"Davr: {period.label}"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(header)
        for cell in ws[3]:
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "B4"
        ws.column_dimensions["A"].width = 30
        return ws

    by_pool = rows_by_pool(period)
    for code, _label in pools():
        rows = by_pool.get(code, [])
        if code == GIRLS:
            header = ["F.I.Sh.", "Guruh", "Navbatchi", "Posbon", "Jami", "Joriy haftada", "Oxirgi naryad"]
            data = [[r.cadet.full_name, r.cadet.group.code, r.by_role.get(SlotRole.DUTY, 0), r.by_role.get(SlotRole.GUARD, 0),
                     r.total, r.this_week, r.last_date] for r in rows]
            sheet("Qizlar", header, data)
        else:
            header = (["F.I.Sh."] + [p.name for p in posts] + [str(ROLE_SHORT[r]) for r in ROLES]
                      + ["Jami", "Og‘ir joylar", "Joriy haftada", "Oxirgi naryad", "Lavozim", "Jazo naryadlari (navbatdan tashqari)"])
            data = [[r.cadet.full_name] + [r.by_post.get(p.code, 0) for p in posts] + [r.by_role.get(x, 0) for x in ROLES]
                    + [r.total, r.hard, r.this_week, r.last_date, r.cadet.position_label_uz, r.penalty]
                    for r in rows]
            sheet(f"Guruh {code}", header, data)

    fair = fairness_all(period, by_pool)
    sheet("Adolat", ["Hovuz", "Kursant", "Jami naryad", "Min", "Max", "Farq", "O‘rtacha", "±1 ichida, %", "Ketma-ket bir joy",
                    "Ketma-ket 2 marta navbatchi (qizlar)", "Holat", "Og‘ir joylar farqi"],
          [[f.title, f.members, f.total, f.min, f.max, f.diff, f.mean, f.within_one_pct,
            len(f.repeats) if f.repeats_applicable else "—", len(f.duty_repeats) if f.pool == GIRLS else "—", f.status_label,
            "—" if f.hard_diff is None else f.hard_diff]
           for f in fair])
    repeats = [[f.title, x["cadet"].full_name, x["post_name"], x["first"], x["second"]] for f in fair for x in f.repeats]
    repeats += [[f.title, x["cadet"].full_name, "Navbatchi (ketma-ket 2 marta)", x["first"], x["second"]]
                for f in fair for x in f.duty_repeats]
    sheet("Ketma-ket bir joy", ["Hovuz", "Kursant", "Joy / lavozim", "1-naryad", "2-naryad"], repeats)

    sections = group_sections(period)
    sheet("Guruhlar kesimi", ["Guruh"] + [p.name for p in posts] + ["Sutkalar", "Jami o‘rin", "Naryadga chiqadigan o‘g‘il", "1 o‘g‘ilga o‘rtacha"],
          [[s.group.code] + [s.by_post[p.code] for p in posts] + [s.days, s.total, s.boys, s.per_boy] for s in sections])

    from .penalties import with_served

    sheet("Zalyotchiklar", ["F.I.Sh.", "Guruh", "Sabab", "Izoh", "Tanbeh sanasi", "Jami", "O‘tildi", "Qoldi", "Holat"],
          [[p.cadet.full_name, p.cadet.group.code, p.get_reason_type_display(), p.note, p.order_date, p.count, p.served,
            p.remaining if p.is_active else 0, p.get_status_display()]
           for p in with_served(Penalty.objects.select_related("cadet__group")).order_by("status", "order_date", "pk")])

    manual = manual_control(period)
    ws = sheet("Qo‘lda tuzatishlar", ["Sabab", "Soni"], [list(x) for x in manual["by_reason"]])
    ws.append([])
    ws.append(["Kim tomonidan", "Soni"])
    for user, n in manual["by_user"]:
        ws.append([user, n])
    sheet("Qo‘lda — kursantlar", ["F.I.Sh.", "Guruh", "Qo‘lda qo‘yilgan", "Qo‘lda olingan", "Hovuz o‘rtachasi (qo‘yilish)", "Ogohlantirish"],
          [[r.cadet.full_name, r.cadet.group.code, r.placed, r.removed, r.pool_mean, "SEZILARLI KO‘P" if r.alert else ""] for r in manual["rows"]])

    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                if isinstance(cell.value, dt.date):
                    cell.number_format = "dd.mm.yyyy"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_filename(period: Period) -> str:
    if period.mode == "week" and period.week:
        return f"statistika_hafta_{period.week.start_date.isoformat()}.xlsx"
    if period.mode == "range":
        return f"statistika_{(period.start or dt.date.min).isoformat()}_{(period.end or current_duty_date()).isoformat()}.xlsx"
    return "statistika_butun_tarix.xlsx"
